"""Répartir la participation entre les secteurs, au prorata des venues.

## Le problème

Deux cotisations, deux natures très différentes :

- l'**adhésion** couvre l'assurance. C'est la part légale, elle ne se
  découpe pas ;
- la **participation** finance les secteurs. C'est elle qu'on répartit.

Jusqu'ici, on demandait à l'accueil « quel secteur vous fait venir ? » et
la participation entière était attribuée à ce secteur-là. Simple, mais
faux dès que la personne circule : quelqu'un qui vient trente fois dans
l'année, dont vingt en Numérique et dix ailleurs, faisait pourtant vivre
six secteurs avec une participation qui n'en créditait qu'un.

D'où le prorata : **chaque venue pèse une part égale**, et la
participation se découpe selon là où la personne est réellement allée.

## Ce qui rend ce calcul possible

``SessionActivite.secteur`` est stocké **séance par séance**, et pas au
niveau de l'atelier. Le modèle le dit lui-même : « animer une séance pour
un autre secteur : elle compte alors pour ce secteur-là, pas pour le
sien ». Autrement dit, la question difficile — à qui profite une séance
d'un atelier intersecteur — est déjà tranchée à la saisie. Ici, on compte.

## Les règles, et pourquoi

**Seules les venues réelles comptent** (``present`` et ``retard``). Une
absence excusée documente le fait que la personne N'EST PAS venue : le
secteur ne l'a pas accueillie ce jour-là, il ne touche pas la part. C'est
aussi le calcul le plus facile à défendre devant un financeur.

**Zéro venue sur l'année : tout va au secteur qui a fait venir la
personne.** Elle a bien payé, l'argent doit atterrir quelque part, et
c'est ce secteur qui a fait le travail d'accueil et d'inscription. On
retombe alors exactement sur l'ancienne règle — le prorata ne fait jamais
disparaître d'argent des vues.

**Les séances annulées ou supprimées ne comptent pas.** Une séance qui
n'a pas eu lieu n'a accueilli personne.

**La somme des parts fait exactement le montant réparti.** Pas
« à peu près » : au centime. Un tableau dont les colonnes ne totalisent
pas le montant encaissé est un tableau que la comptabilité renvoie, et
elle a raison. Voir ``decouper``.

## Deux montants, et ils ne servent pas à la même chose

- le **dû** : ce qui a été facturé. C'est le budget théorique ;
- l'**encaissé** : ce qui est réellement rentré. C'est ce qu'un référent
  peut engager.

Les deux sont calculés et affichés. En cas de doute, c'est l'encaissé qui
fait foi — on n'engage pas de l'argent qui n'est pas arrivé.
"""
from __future__ import annotations

from datetime import date

from app.extensions import db
from app.models import (
    Cotisation,
    Paiement,
    Participant,
    PresenceActivite,
    SessionActivite,
)

#: Les statuts de présence qui valent « la personne est venue ».
#: Une absence excusée n'en fait pas partie : elle atteste le contraire.
VENUES_REELLES = ("present", "retard")

#: Étiquette des personnes dont aucune venue n'a été trouvée sur l'année.
SANS_SECTEUR = "Sans secteur"


# ---------------------------------------------------------------------------
# L'arrondi : la partie qui doit être irréprochable
# ---------------------------------------------------------------------------

def decouper(montant: float, poids: dict[str, int]) -> dict[str, float]:
    """Découpe ``montant`` selon ``poids``, en bouclant au centime près.

    L'arrondi naïf (arrondir chaque part séparément) ne tombe presque
    jamais juste : 10 € en trois parts égales donne 3,33 × 3 = 9,99 €, et
    il manque un centime que personne ne sait expliquer. Sur une centaine
    de participations, l'écart devient visible et décrédibilise tout le
    tableau.

    On applique donc la **méthode du plus fort reste** : chaque part est
    d'abord tronquée au centime inférieur, puis les centimes restants sont
    donnés un par un aux parts dont la décimale sacrifiée était la plus
    grande. La somme est alors exacte par construction.

    En cas d'égalité parfaite des restes, on départage par le nom du
    secteur : le calcul doit rendre le même résultat à chaque exécution,
    sinon deux exports du même jour ne se ressemblent pas.
    """
    total_poids = sum(max(0, int(p)) for p in poids.values())
    if total_poids <= 0 or not poids:
        return {}

    centimes_a_repartir = int(round(float(montant or 0) * 100))
    if centimes_a_repartir == 0:
        return {cle: 0.0 for cle in poids}

    parts: dict[str, int] = {}
    restes: list[tuple[int, str]] = []
    for cle, p in poids.items():
        exact = centimes_a_repartir * max(0, int(p))
        entier, reste = divmod(exact, total_poids)
        parts[cle] = entier
        restes.append((reste, cle))

    # Les centimes non attribués par la troncature, aux plus forts restes.
    manquants = centimes_a_repartir - sum(parts.values())
    restes.sort(key=lambda item: (-item[0], item[1]))
    for _, cle in restes[:manquants]:
        parts[cle] += 1

    return {cle: round(centimes / 100, 2) for cle, centimes in parts.items()}


# ---------------------------------------------------------------------------
# Les venues
# ---------------------------------------------------------------------------

def bornes_annee_scolaire(annee: int) -> tuple[date, date]:
    """Du 1er septembre au 31 août — la convention de toute l'application."""
    return date(annee, 9, 1), date(annee + 1, 8, 31)


def venues_par_secteur(annee: int, participant_ids=None,
                       a_la_date: date | None = None) -> dict[int, dict[str, int]]:
    """``{participant_id: {secteur: nombre de venues}}`` pour l'année.

    ``a_la_date`` arrête le comptage à cette date incluse. C'est ce qui
    donne son sens à un arrêté : « au 31 décembre » doit compter les
    venues jusqu'au 31 décembre, et pas celles de février qu'on ne
    connaissait pas encore.

    Une seule requête, groupée en base : compter les venues de trois cents
    personnes ne doit pas coûter trois cents allers-retours.
    """
    debut, fin = bornes_annee_scolaire(annee)
    if a_la_date is not None and a_la_date < fin:
        fin = a_la_date
    return venues_entre(debut, fin, participant_ids)


def venues_entre(debut: date, fin: date, participant_ids=None) -> dict[int, dict[str, int]]:
    """Les venues comptées entre deux dates, quelles qu'elles soient.

    Le même comptage sert l'année scolaire et une période libre : c'est
    volontaire. Deux comptages pour une même notion finiraient par ne plus
    dire la même chose, et personne ne saurait lequel croire.
    """
    jour = db.func.coalesce(SessionActivite.date_session, SessionActivite.rdv_date)

    requete = (
        db.session.query(
            PresenceActivite.participant_id,
            SessionActivite.secteur,
            db.func.count(PresenceActivite.id),
        )
        .join(SessionActivite, SessionActivite.id == PresenceActivite.session_id)
        .filter(SessionActivite.is_deleted.is_(False))
        # Une séance annulée n'a accueilli personne.
        .filter(db.func.lower(db.func.coalesce(SessionActivite.statut, "")) != "annulee")
        .filter(jour.isnot(None))
        .filter(jour >= debut, jour <= fin)
        # Seules les venues réelles : une absence excusée atteste que la
        # personne n'est pas venue.
        .filter(PresenceActivite.presence_type.in_(VENUES_REELLES))
        .group_by(PresenceActivite.participant_id, SessionActivite.secteur)
    )
    if participant_ids is not None:
        ids = [int(i) for i in participant_ids if i]
        if not ids:
            return {}
        requete = requete.filter(PresenceActivite.participant_id.in_(ids))

    resultat: dict[int, dict[str, int]] = {}
    for pid, secteur, nb in requete.all():
        if not pid:
            continue
        nom = (secteur or "").strip()
        if not nom:
            continue
        compte = resultat.setdefault(int(pid), {})
        compte[nom] = compte.get(nom, 0) + int(nb or 0)
    return resultat


# ---------------------------------------------------------------------------
# « Numerique » et « Numérique » sont le même secteur
# ---------------------------------------------------------------------------

def table_de_canonisation() -> dict[str, str]:
    """Clé sans accents ni casse -> libellé officiel du secteur.

    Les secteurs sont des chaînes libres dans ``SessionActivite.secteur`` et
    ``Participant.created_secteur``. Une majuscule ou un accent de travers
    créerait un septième secteur fantôme dans le tableau, avec sa part
    d'argent dedans — exactement le doublon qu'on a déjà combattu sur les
    villes et les quartiers. On ramène donc chaque écriture à son libellé
    officiel.
    """
    from app.secteurs import get_secteur_labels
    from app.services.recherche_texte import sans_accent

    table: dict[str, str] = {}
    try:
        libelles = get_secteur_labels(active_only=False)
    except Exception:  # noqa: BLE001 — un référentiel indisponible ne casse pas un calcul
        libelles = []
    for libelle in libelles:
        cle = (sans_accent(libelle) or "").strip()
        if cle:
            table.setdefault(cle, libelle)
    return table


def canoniser(secteur: str | None, table: dict[str, str] | None = None) -> str:
    """Le libellé officiel, ou l'écriture d'origine si le secteur est inconnu.

    On ne jette JAMAIS un secteur qu'on ne reconnaît pas : il porte de
    l'argent. Mieux vaut le voir apparaître avec son orthographe bancale —
    c'est visible, donc corrigeable — que le voir disparaître.
    """
    from app.services.recherche_texte import sans_accent

    brut = (secteur or "").strip()
    if not brut:
        return ""
    if table is None:
        table = table_de_canonisation()
    return table.get((sans_accent(brut) or "").strip(), brut)


# ---------------------------------------------------------------------------
# La répartition
# ---------------------------------------------------------------------------

def _participations_avec_reglement(annee: int, a_la_date: date | None = None):
    """Les participations de l'année, avec ce qui a été encaissé dessus.

    ``a_la_date`` arrête la photo : on ne retient que les participations
    déjà enregistrées à cette date, et que les versements déjà reçus. Un
    arrêté au 31 décembre qui compterait un chèque de mars ne serait pas
    un arrêté.

    Un seul aller-retour, sous-requête de somme comme l'écran des impayés :
    interroger les paiements ligne à ligne ferait une requête par personne.
    """
    versements = db.session.query(
        Paiement.cotisation_id.label("cid"),
        db.func.coalesce(db.func.sum(Paiement.montant), 0.0).label("regle"),
    )
    if a_la_date is not None:
        versements = versements.filter(Paiement.date_paiement <= a_la_date)
    regle_subq = versements.group_by(Paiement.cotisation_id).subquery()

    requete = (
        db.session.query(Cotisation, db.func.coalesce(regle_subq.c.regle, 0.0))
        .outerjoin(regle_subq, regle_subq.c.cid == Cotisation.id)
        .filter(
            Cotisation.annee_scolaire == annee,
            Cotisation.type_cotisation == "participation",
        )
    )
    if a_la_date is not None:
        requete = requete.filter(Cotisation.date_reference <= a_la_date)
    return requete.all()


def periode_couverte(annee: int, a_la_date: date | None = None) -> tuple[date, date]:
    """Les deux dates que le tableau couvre réellement.

    « 2025-2026 » ne veut rien dire pour une comptabilité : elle a besoin
    d'un du-tel-jour-au-tel-jour. Et la borne haute n'est pas toujours le
    31 août — sur l'année en cours, rien n'existe après aujourd'hui, et
    annoncer une période qui va jusqu'en août laisserait croire que les
    mois à venir sont déjà comptés.
    """
    debut, fin = bornes_annee_scolaire(annee)
    butoir = a_la_date or date.today()
    return debut, min(fin, butoir) if butoir < fin else fin


def _titulaires(lignes):
    """Qui porte chaque participation, et quelles fiches il faut charger.

    Une participation est individuelle par construction — même sur un
    bulletin familial, chaque personne paye la sienne. Mais une reprise de
    données peut en avoir rattaché une à un foyer : on mutualise alors les
    venues de ses membres plutôt que de perdre la somme.
    """
    ids_directs = {c.participant_id for c, _, _ in lignes if c.participant_id}
    foyers = {c.foyer_id for c, _, _ in lignes if c.foyer_id and not c.participant_id}

    membres_par_foyer: dict[int, list[Participant]] = {}
    if foyers:
        for membre in Participant.query.filter(
                Participant.foyer_id.in_(list(foyers))).all():
            membres_par_foyer.setdefault(membre.foyer_id, []).append(membre)

    ids = set(ids_directs)
    for membres in membres_par_foyer.values():
        ids.update(m.id for m in membres)

    fiches = {
        p.id: p for p in Participant.query.filter(Participant.id.in_(list(ids))).all()
    } if ids else {}
    return membres_par_foyer, fiches, ids


def _composer(lignes, venues, table) -> dict:
    """Le moteur, commun à l'année scolaire et à la période libre.

    ``lignes`` : des triplets ``(cotisation, du, encaisse)`` déjà bornés par
    l'appelant — c'est LUI qui décide ce qui entre dans le périmètre, ici on
    ne fait que découper et additionner.

    Un seul moteur pour les deux modes : deux calculs de la même chose
    finissent par diverger, et c'est la comptabilité qui découvre l'écart.
    """
    membres_par_foyer, fiches, _ = _titulaires(lignes)

    personnes: list[dict] = []
    secteurs: dict[str, dict] = {}
    total_du = total_regle = 0.0
    nb_repli = 0

    for cotisation, du, encaisse in lignes:
        if cotisation.participant_id:
            porteurs = [fiches.get(cotisation.participant_id)]
        else:
            porteurs = membres_par_foyer.get(cotisation.foyer_id, [])
        porteurs = [t for t in porteurs if t is not None]
        if not porteurs:
            continue
        principal = porteurs[0]

        # Les venues de la personne (ou, pour un foyer, celles de tous ses
        # membres additionnées), ramenées aux libellés officiels.
        compte: dict[str, int] = {}
        for porteur in porteurs:
            for secteur_brut, nb in venues.get(porteur.id, {}).items():
                nom = canoniser(secteur_brut, table)
                if nom:
                    compte[nom] = compte.get(nom, 0) + nb

        repli = not compte
        if repli:
            # Aucune venue sur la période : tout revient au secteur qui a
            # fait venir la personne — l'ancienne règle, gardée pour ce cas.
            orienteur = ""
            for porteur in porteurs:
                orienteur = canoniser(porteur.created_secteur, table)
                if orienteur:
                    break
            compte = {orienteur or SANS_SECTEUR: 1}
            nb_repli += 1

        du = round(float(du or 0), 2)
        encaisse = round(float(encaisse or 0), 2)
        parts_du = decouper(du, compte)
        parts_regle = decouper(encaisse, compte)

        detail = {}
        for nom, nb in sorted(compte.items(), key=lambda kv: (-kv[1], kv[0])):
            detail[nom] = {
                "venues": 0 if repli else nb,
                "du": parts_du.get(nom, 0.0),
                "regle": parts_regle.get(nom, 0.0),
            }
            case = secteurs.setdefault(
                nom, {"du": 0.0, "regle": 0.0, "venues": 0, "nb_personnes": 0}
            )
            case["du"] = round(case["du"] + detail[nom]["du"], 2)
            case["regle"] = round(case["regle"] + detail[nom]["regle"], 2)
            case["venues"] += detail[nom]["venues"]
            case["nb_personnes"] += 1

        total_du = round(total_du + du, 2)
        total_regle = round(total_regle + encaisse, 2)
        personnes.append({
            "participant": principal,
            "titulaires": porteurs,
            "cotisation": cotisation,
            "du": du,
            "regle": encaisse,
            "total_venues": 0 if repli else sum(compte.values()),
            "repli": repli,
            "secteur_orienteur": canoniser(principal.created_secteur, table),
            "parts": detail,
        })

    personnes.sort(key=lambda p: (
        (p["participant"].nom or "").lower(), (p["participant"].prenom or "").lower()
    ))
    return {
        "personnes": personnes,
        "secteurs": dict(sorted(secteurs.items(), key=lambda kv: (-kv[1]["regle"], kv[0]))),
        "totaux": {
            "du": total_du,
            "regle": total_regle,
            "nb_personnes": len(personnes),
            "nb_repli": nb_repli,
            "venues": sum(case["venues"] for case in secteurs.values()),
        },
    }


def repartition(annee: int, a_la_date: date | None = None) -> dict:
    """La répartition de l'ANNÉE SCOLAIRE : par personne et par secteur.

    Retourne ``{mode, annee_scolaire, a_la_date, periode, personnes,
    secteurs, totaux}``.

    ``a_la_date`` rejoue l'année telle qu'elle était connue ce jour-là —
    venues et versements compris. Sans elle, c'est l'état du jour, qui
    continuera de bouger jusqu'au 31 août.

    ``personnes`` détaille CHAQUE participation : combien de venues dans
    quel secteur, et quelle part de son dû et de son encaissé y revient.
    C'est ce détail qui permet de répondre à « pourquoi le Numérique a-t-il
    420 € ? » sans ouvrir trente fiches.

    ``secteurs`` en est la somme, et rien d'autre : les totaux ne sont
    jamais calculés séparément du détail, sinon les deux finissent par
    diverger et plus personne ne sait lequel croire.
    """
    from app.services.cotisations import libelle_annee_scolaire

    debut, fin = periode_couverte(annee, a_la_date)
    enveloppe = {
        "mode": "annee_scolaire",
        "annee_scolaire": annee,
        "a_la_date": a_la_date,
        "periode": {"debut": debut, "fin": fin},
        "libelle_annee": libelle_annee_scolaire(annee),
    }

    brutes = _participations_avec_reglement(annee, a_la_date)
    if not brutes:
        return {**enveloppe, "personnes": [], "secteurs": {},
                "totaux": {"du": 0.0, "regle": 0.0, "nb_personnes": 0,
                           "nb_repli": 0, "venues": 0}}

    lignes = [(cotisation, cotisation.montant_du, regle) for cotisation, regle in brutes]
    _, _, ids = _titulaires(lignes)
    venues = venues_entre(debut, fin, ids)
    return {**enveloppe, **_composer(lignes, venues, table_de_canonisation())}


# ---------------------------------------------------------------------------
# La période libre : quand l'exercice comptable ne suit pas l'année scolaire
# ---------------------------------------------------------------------------

def _participations_de_la_periode(debut: date, fin: date):
    """Les participations qui touchent la fenêtre, dû et encaissé bornés.

    Ici, on ne raisonne plus par année scolaire mais par DATES, parce que
    l'exercice comptable peut suivre l'année civile. Deux notions,
    volontairement distinctes, qu'on ne mélange jamais dans la même colonne :

    - **l'encaissé** est la somme des versements REÇUS dans la fenêtre.
      C'est l'argent réellement entré, celui d'une comptabilité de
      trésorerie ;
    - **le dû** est le montant des participations ENREGISTRÉES dans la
      fenêtre (date de référence du tarif). C'est ce qui a été facturé sur
      la période.

    Une participation peut n'avoir que l'un des deux : facturée en
    septembre et payée en janvier, elle apporte son dû à un exercice et son
    encaissement à l'autre. C'est le comportement voulu — et c'est
    précisément pour ça qu'on ne pro-rate pas une cotisation annuelle au
    temps, ce qui ne voudrait rien dire.
    """
    versements = (
        db.session.query(
            Paiement.cotisation_id,
            db.func.coalesce(db.func.sum(Paiement.montant), 0.0),
        )
        .join(Cotisation, Cotisation.id == Paiement.cotisation_id)
        .filter(Cotisation.type_cotisation == "participation")
        .filter(Paiement.date_paiement >= debut, Paiement.date_paiement <= fin)
        .group_by(Paiement.cotisation_id)
        .all()
    )
    encaisse = {cid: round(float(montant or 0), 2) for cid, montant in versements}

    facturees = (
        Cotisation.query
        .filter(Cotisation.type_cotisation == "participation")
        .filter(Cotisation.date_reference >= debut, Cotisation.date_reference <= fin)
        .all()
    )
    du = {c.id: round(float(c.montant_du or 0), 2) for c in facturees}

    identifiants = set(encaisse) | set(du)
    if not identifiants:
        return []
    cotisations = Cotisation.query.filter(
        Cotisation.id.in_(list(identifiants))).all()
    return [(c, du.get(c.id, 0.0), encaisse.get(c.id, 0.0)) for c in cotisations]


def repartition_periode(debut: date, fin: date) -> dict:
    """La répartition sur une fenêtre de dates libre.

    Existe parce que les cotisations vivent en année SCOLAIRE tandis qu'un
    exercice comptable peut suivre l'année CIVILE. Sans ça, la comptabilité
    devrait recomposer son 1er janvier - 31 décembre à la main depuis deux
    années scolaires, et personne ne le ferait deux fois.

    Les venues qui découpent les montants sont celles de la MÊME fenêtre :
    « sur l'année civile, cette personne est venue tant de fois ici et tant
    de fois là, ses versements de l'année s'y répartissent ». Aucune venue
    dans la fenêtre : on retombe sur le secteur orienteur, comme partout
    ailleurs.

    Même moteur de découpe et même arrondi que l'année scolaire : les
    colonnes bouclent au centime ici aussi.
    """
    if fin < debut:
        debut, fin = fin, debut

    enveloppe = {
        "mode": "periode",
        "annee_scolaire": None,
        "a_la_date": fin,
        "periode": {"debut": debut, "fin": fin},
        "libelle_annee": f"du {debut:%d/%m/%Y} au {fin:%d/%m/%Y}",
    }

    lignes = _participations_de_la_periode(debut, fin)
    if not lignes:
        return {**enveloppe, "personnes": [], "secteurs": {},
                "totaux": {"du": 0.0, "regle": 0.0, "nb_personnes": 0,
                           "nb_repli": 0, "venues": 0}}

    _, _, ids = _titulaires(lignes)
    venues = venues_entre(debut, fin, ids)
    return {**enveloppe, **_composer(lignes, venues, table_de_canonisation())}


# ---------------------------------------------------------------------------
# Figer : l'arrêté
# ---------------------------------------------------------------------------

def arreter(annee: int, date_arrete: date, *, libelle: str | None = None,
            note: str | None = None, user_id: int | None = None):
    """Photographie la répartition à une date, et l'écrit. Elle ne bougera plus.

    Retourne ``(arrete, message)``. L'arrêté vaut ``None`` si le geste a
    été refusé, et le message dit pourquoi — il est destiné à l'écran.

    Trois refus, tous pour la même raison de fond : un arrêté est une pièce
    qu'on ressortira en réunion, il ne doit pas pouvoir mentir.

    - **une date antérieure à la rentrée** : rien ne s'est encore passé,
      la photo serait vide et son titre laisserait croire le contraire ;
    - **une date future** : on photographierait un état qui n'existe pas
      encore. Un « arrêté au 31 août » signé en mars annonce cinq mois
      qui n'ont pas eu lieu ;
    - **un arrêté déjà pris à cette date** : deux pièces différentes
      portant la même date et le même intitulé, c'est la garantie qu'un
      jour quelqu'un cite la mauvaise.
    """
    from app.models import RepartitionArretee, RepartitionArreteeLigne

    debut, fin = bornes_annee_scolaire(annee)
    if date_arrete < debut:
        return None, (
            f"Le {date_arrete:%d/%m/%Y} précède la rentrée du "
            f"{debut:%d/%m/%Y} : il n'y a rien à arrêter."
        )
    if date_arrete > date.today():
        return None, (
            "On ne peut pas arrêter une répartition à une date future : "
            "l'année n'est pas encore allée jusque-là."
        )
    if date_arrete > fin:
        date_arrete = fin

    if RepartitionArretee.query.filter_by(
            annee_scolaire=annee, date_arrete=date_arrete).first():
        return None, (
            f"Un arrêté existe déjà au {date_arrete:%d/%m/%Y} pour "
            f"{libelle_annee(annee)}. Supprimez-le d'abord si vous voulez le refaire."
        )

    vue = repartition(annee, a_la_date=date_arrete)

    arrete = RepartitionArretee(
        annee_scolaire=annee,
        date_arrete=date_arrete,
        libelle=(libelle or "").strip() or None,
        note=(note or "").strip() or None,
        total_du=vue["totaux"]["du"],
        total_regle=vue["totaux"]["regle"],
        cree_par_user_id=user_id,
    )
    db.session.add(arrete)
    db.session.flush()

    for personne in vue["personnes"]:
        fiche = personne["participant"]
        nom = f"{fiche.nom or ''} {fiche.prenom or ''}".strip()
        for secteur, part in personne["parts"].items():
            db.session.add(RepartitionArreteeLigne(
                arrete_id=arrete.id,
                secteur=secteur,
                participant_id=fiche.id,
                participant_nom=nom[:240] or None,
                venues=part["venues"],
                montant_du=part["du"],
                montant_regle=part["regle"],
                repli=personne["repli"],
            ))

    db.session.commit()
    return arrete, f"Répartition arrêtée au {date_arrete:%d/%m/%Y}."


def libelle_annee(annee: int) -> str:
    from app.services.cotisations import libelle_annee_scolaire

    return libelle_annee_scolaire(annee)


def arretes(annee: int | None = None) -> list:
    """Les arrêtés existants, du plus récent au plus ancien."""
    from app.models import RepartitionArretee

    requete = RepartitionArretee.query
    if annee is not None:
        requete = requete.filter_by(annee_scolaire=annee)
    return requete.order_by(
        RepartitionArretee.annee_scolaire.desc(),
        RepartitionArretee.date_arrete.desc(),
    ).all()


def dernier_arrete(annee: int):
    """Le dernier arrêté de l'année, s'il y en a un."""
    liste = arretes(annee)
    return liste[0] if liste else None


# ---------------------------------------------------------------------------
# Une seule personne : ce que l'accueil a sous les yeux
# ---------------------------------------------------------------------------

def repartition_personne(participant, annee: int) -> dict | None:
    """Comment la participation de CETTE personne se répartit, cette année.

    Pensé pour la fiche participant : à l'accueil, la question arrive de
    face — « et mes 20 €, ils vont où ? ». Y répondre en ouvrant un
    tableau de bord global serait absurde.

    Rend ``None`` s'il n'y a pas de participation enregistrée : il n'y a
    alors rien à répartir, et afficher un bloc vide ferait croire à un
    bug.

    Le calcul est celui de l'année entière, borné à une personne — jamais
    une seconde formule. Deux calculs de la même chose finissent par
    diverger, et c'est l'accueil qui se fait contredire par le tableau de
    la direction.
    """
    fiche_id = getattr(participant, "id", None)
    if not fiche_id:
        return None

    lignes = _participations_avec_reglement(annee)
    concernee = None
    for cotisation, regle in lignes:
        if cotisation.participant_id == fiche_id:
            concernee = (cotisation, regle)
            break
        if cotisation.foyer_id and not cotisation.participant_id:
            if getattr(participant, "foyer_id", None) == cotisation.foyer_id:
                concernee = (cotisation, regle)
                break
    if concernee is None:
        return None

    vue = repartition(annee)
    for personne in vue["personnes"]:
        if personne["participant"].id == fiche_id:
            return personne
        if any(t.id == fiche_id for t in personne["titulaires"]):
            return personne
    return None


# ---------------------------------------------------------------------------
# L'export : un classeur qui se défend tout seul
# ---------------------------------------------------------------------------

def export_xlsx(vue: dict, *, arrete=None):
    """Le classeur de la répartition. Deux onglets et un contrôle.

    Un tableur circule par courriel, détaché de l'écran qui l'a produit.
    Trois mois plus tard, personne ne saura dire s'il portait un chiffre
    provisoire ou un arrêté. **L'en-tête le dit donc en toutes lettres**,
    sur chaque onglet.

    L'onglet « Contrôle » est là pour la comptabilité : il rapproche la
    somme des colonnes de la somme des participations et affiche l'écart.
    C'est la première chose qu'elle vérifiera — autant la lui donner
    plutôt que de la laisser la refaire à la main.
    """
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    if arrete is not None:
        nature = f"ARRÊTÉ FIGÉ au {arrete.date_arrete:%d/%m/%Y} — ces montants ne bougeront plus"
        if arrete.libelle:
            nature = f"{arrete.libelle} — {nature}"
    elif vue.get("mode") == "periode":
        bornes_libres = vue["periode"]
        nature = (
            f"PÉRIODE LIBRE du {bornes_libres['debut']:%d/%m/%Y} au "
            f"{bornes_libres['fin']:%d/%m/%Y} — encaissements reçus et venues "
            "constatées dans cette fenêtre. Une saisie rétroactive (présence "
            "pointée après coup, règlement enregistré en retard) peut encore "
            "la modifier : pour un chiffre définitif, figer un arrêté."
        )
    else:
        nature = (
            f"RÉPARTITION PROVISOIRE au {date.today():%d/%m/%Y} — "
            "elle bougera à chaque séance pointée. Pour engager un budget, "
            "utiliser un arrêté figé."
        )

    # « 2025-2026 » ne veut rien dire pour une comptabilité : elle a besoin
    # d'un du-tel-jour-au-tel-jour. Et le classeur circulera loin de l'écran
    # qui l'a produit — il doit porter ses bornes lui-même.
    bornes = vue.get("periode") or {}
    if bornes.get("debut") and bornes.get("fin") and vue.get("mode") == "periode":
        couverture = (
            f"Période couverte : du {bornes['debut']:%d/%m/%Y} au "
            f"{bornes['fin']:%d/%m/%Y} (fenêtre choisie, indépendante de "
            "l'année scolaire)."
        )
    elif bornes.get("debut") and bornes.get("fin"):
        couverture = (
            f"Période couverte : du {bornes['debut']:%d/%m/%Y} au {bornes['fin']:%d/%m/%Y} "
            "(année scolaire de septembre à août — à ne pas confondre avec "
            "l'exercice comptable, qui peut suivre l'année civile)."
        )
    else:
        couverture = "Période couverte : année scolaire de septembre à août."

    gras = Font(bold=True)
    wb = Workbook()

    # --- Onglet 1 : par secteur -------------------------------------------
    feuille = wb.active
    feuille.title = "Par secteur"
    titre = vue.get("libelle_annee") or ""
    feuille.append([f"Répartition de la participation — {titre}"])
    feuille["A1"].font = Font(bold=True, size=13)
    feuille.append([nature])
    feuille.append([couverture])
    feuille.append([
        "Au prorata des venues réelles (présent, retard). Les absences excusées "
        "ne comptent pas. Qui n'est venu nulle part est rattaché au secteur "
        "qui l'a fait venir."
    ])
    feuille.append([])

    if vue.get("mode") == "periode":
        # « Dû » et « Encaissé » ne comptent pas la même chose sur une
        # fenêtre libre : l'un est ce qui a été facturé dans la période,
        # l'autre ce qui y est rentré. Une participation facturée en
        # septembre et payée en janvier apporte l'un à un exercice et
        # l'autre au suivant. Les intituler pareil serait un piège.
        entetes = ["Secteur", "Venues", "Personnes", "Part de l'encaissé",
                   "Facturé dans la période (€)", "Encaissé dans la période (€)"]
    else:
        entetes = ["Secteur", "Venues", "Personnes", "Part de l'encaissé",
                   "Dû (€)", "Encaissé (€)"]
    feuille.append(entetes)
    ligne_entetes = feuille.max_row
    for cellule in feuille[ligne_entetes]:
        cellule.font = gras

    total_regle = vue["totaux"]["regle"] or 0.0
    for nom, case in vue["secteurs"].items():
        feuille.append([
            nom, case["venues"], case["nb_personnes"],
            round(case["regle"] / total_regle, 4) if total_regle else 0,
            case["du"], case["regle"],
        ])
    # Les index se déduisent de l'écriture réelle : une ligne d'en-tête
    # ajoutée plus haut décalerait silencieusement tous les formats.
    premiere = ligne_entetes + 1
    for ligne in range(premiere, feuille.max_row + 1):
        feuille.cell(row=ligne, column=4).number_format = "0.0%"
        feuille.cell(row=ligne, column=5).number_format = "#,##0.00"
        feuille.cell(row=ligne, column=6).number_format = "#,##0.00"

    feuille.append([])
    feuille.append(["TOTAL", vue["totaux"]["venues"], vue["totaux"]["nb_personnes"],
                    1 if total_regle else 0, vue["totaux"]["du"], vue["totaux"]["regle"]])
    ligne_total = feuille.max_row
    for cellule in feuille[ligne_total]:
        cellule.font = gras
    feuille.cell(row=ligne_total, column=4).number_format = "0.0%"
    feuille.cell(row=ligne_total, column=5).number_format = "#,##0.00"
    feuille.cell(row=ligne_total, column=6).number_format = "#,##0.00"

    feuille.freeze_panes = f"A{ligne_entetes + 1}"
    for index, largeur in enumerate((38, 10, 12, 18, 14, 14), start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur
    feuille["A2"].alignment = Alignment(wrap_text=False)

    # --- Onglet 2 : le détail qui justifie chaque part ---------------------
    detail = wb.create_sheet("Détail par personne")
    detail.append([f"Détail de la répartition — {vue['libelle_annee']}"])
    detail["A1"].font = Font(bold=True, size=13)
    detail.append([nature])
    detail.append([couverture])
    colonnes = ["Personne", "Secteur", "Venues dans ce secteur",
                "Total venues de la personne", "Dû réparti (€)",
                "Encaissé réparti (€)", "Sans venue (repli)"]
    detail.append(colonnes)
    for cellule in detail[4]:
        cellule.font = gras

    nb_lignes = 0
    for personne in vue["personnes"]:
        fiche = personne.get("participant")
        nom = personne.get("nom") or (
            f"{fiche.nom or ''} {fiche.prenom or ''}".strip() if fiche else "")
        for secteur, part in personne["parts"].items():
            detail.append([
                nom, secteur, part["venues"], personne["total_venues"],
                part["du"], part["regle"],
                "oui" if personne["repli"] else "",
            ])
            nb_lignes += 1
    for ligne in range(5, 5 + nb_lignes):
        detail.cell(row=ligne, column=5).number_format = "#,##0.00"
        detail.cell(row=ligne, column=6).number_format = "#,##0.00"

    detail.freeze_panes = "A5"
    if nb_lignes:
        detail.auto_filter.ref = f"A4:{get_column_letter(len(colonnes))}{4 + nb_lignes}"
    for index, largeur in enumerate((32, 38, 22, 26, 16, 18, 16), start=1):
        detail.column_dimensions[get_column_letter(index)].width = largeur

    # --- Onglet 3 : le contrôle que la comptabilité ferait à la main -------
    controle = wb.create_sheet("Contrôle")
    controle.append(["Contrôle de bouclage"])
    controle["A1"].font = Font(bold=True, size=13)
    controle.append([nature])
    controle.append([couverture])
    controle.append(["Vérification", "Dû (€)", "Encaissé (€)"])
    for cellule in controle[4]:
        cellule.font = gras

    somme_secteurs_du = round(sum(c["du"] for c in vue["secteurs"].values()), 2)
    somme_secteurs_regle = round(sum(c["regle"] for c in vue["secteurs"].values()), 2)
    controle.append(["Somme des parts par secteur", somme_secteurs_du, somme_secteurs_regle])
    controle.append(["Somme des participations", vue["totaux"]["du"], vue["totaux"]["regle"]])
    controle.append(["Écart",
                     round(somme_secteurs_du - vue["totaux"]["du"], 2),
                     round(somme_secteurs_regle - vue["totaux"]["regle"], 2)])
    for ligne in range(5, 8):
        for colonne in (2, 3):
            controle.cell(row=ligne, column=colonne).number_format = "#,##0.00"
    for cellule in controle[7]:
        cellule.font = gras
    controle.append([])
    controle.append([
        "Un écart non nul signalerait un arrondi perdu. La répartition "
        "distribue les centimes restants au plus fort reste : l'écart doit "
        "toujours être de 0,00 €."
    ])
    for index, largeur in enumerate((44, 16, 18), start=1):
        controle.column_dimensions[get_column_letter(index)].width = largeur

    sortie = BytesIO()
    wb.save(sortie)
    sortie.seek(0)
    return sortie
