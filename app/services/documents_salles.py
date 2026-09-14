"""Contrats, conventions, états des lieux, factures, attestations.

Les documents sont construits DIRECTEMENT à partir des données, sans
gabarit binaire dans le dépôt : un contrat se régénère toujours à
l'identique, et il n'y a pas de fichier .docx à maintenir en parallèle du
code. Ce qui doit rester modifiable sans développeur — les conditions
générales, l'identité du bailleur — vit en base et s'édite depuis
l'application.

La conversion en PDF réutilise LibreOffice tel que le module d'émargement
l'a déjà câblé : même binaire, mêmes réglages, mêmes déboires évités.
"""
from __future__ import annotations

import os
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from app.models import STATUTS_RESERVATION_LABELS, Reservation

#: Conditions générales proposées par défaut. Volontairement sobres et
#: génériques : chaque structure doit les relire et les adapter, c'est un
#: point de départ, pas un modèle juridique validé.
CONDITIONS_GENERALES_DEFAUT = """\
1. Objet — La présente mise à disposition est consentie à titre précaire et \
révocable. Elle ne confère au bénéficiaire aucun droit au maintien dans les \
lieux, ni aucun droit à la propriété commerciale.

2. Destination — Les locaux sont mis à disposition pour l'usage décrit au \
présent document, à l'exclusion de tout autre. Toute sous-location ou mise à \
disposition à un tiers est interdite.

3. Assurance — Le bénéficiaire déclare être assuré en responsabilité civile \
pour l'activité exercée et s'engage à fournir une attestation en cours de \
validité avant la remise des clés.

4. État des lieux — Un état des lieux contradictoire est établi à l'entrée et \
à la sortie. Les dégradations constatées sont à la charge du bénéficiaire.

5. Capacité et sécurité — L'effectif maximal fixé par la commission de \
sécurité ne peut être dépassé. Les issues de secours doivent rester libres en \
permanence.

6. Restitution — Les locaux sont restitués propres, le mobilier remis en place \
et les déchets évacués. À défaut, les frais de remise en état sont retenus sur \
la caution.

7. Annulation — Toute annulation doit être signalée au plus tôt. \
Les modalités de remboursement éventuel sont précisées aux conditions \
particulières.

8. Résiliation — La structure peut mettre fin à la mise à disposition sans \
préavis en cas de manquement aux présentes conditions ou de nécessité de \
service dûment motivée.
"""

GRIS = RGBColor(0x55, 0x5B, 0x6A)


# ---------------------------------------------------------------------------
# Mise en page
# ---------------------------------------------------------------------------

def _document(paysage: bool = False) -> Document:
    doc = Document()
    section = doc.sections[0]
    if paysage:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for marge in ("left_margin", "right_margin"):
        setattr(section, marge, Cm(2.0))
    section.top_margin = section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    return doc


def _titre(doc, texte: str, sous_titre: str = "") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texte.upper())
    r.bold = True
    r.font.size = Pt(16)
    if sous_titre:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(sous_titre)
        r2.font.size = Pt(10)
        r2.font.color.rgb = GRIS


def _section(doc, texte: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(texte)
    r.bold = True
    r.font.size = Pt(11)


def _ligne(doc, libelle: str, valeur, gras: bool = False) -> None:
    """« Libellé : valeur », en sautant les valeurs vides."""
    if valeur in (None, "", 0) and valeur != 0:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.add_run(f"{libelle} : ").bold = False
    r = p.add_run(str(valeur))
    r.bold = gras


def _paragraphe(doc, texte: str, taille: float = 10.5, italique: bool = False) -> None:
    for bloc in (texte or "").split("\n\n"):
        bloc = bloc.strip()
        if not bloc:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(bloc)
        r.font.size = Pt(taille)
        r.italic = italique


def _tableau(doc, entetes: list[str], lignes: list[list[str]], largeurs=None):
    table = doc.add_table(rows=1, cols=len(entetes))
    table.style = "Table Grid"
    for i, entete in enumerate(entetes):
        cellule = table.rows[0].cells[i]
        cellule.text = ""
        r = cellule.paragraphs[0].add_run(entete)
        r.bold = True
        r.font.size = Pt(9.5)
    for ligne in lignes:
        cellules = table.add_row().cells
        for i, valeur in enumerate(ligne):
            cellules[i].text = ""
            r = cellules[i].paragraphs[0].add_run(str(valeur))
            r.font.size = Pt(9.5)
    if largeurs:
        for ligne in table.rows:
            for i, largeur in enumerate(largeurs):
                if i < len(ligne.cells):
                    ligne.cells[i].width = Cm(largeur)
    return table


def _signatures(doc, gauche: str, droite: str) -> None:
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Fait à …………………………………, le ……… / ……… / ………")
    r.font.size = Pt(10)
    doc.add_paragraph()
    table = doc.add_table(rows=2, cols=2)
    for i, titre in enumerate((gauche, droite)):
        cellule = table.rows[0].cells[i]
        cellule.text = ""
        r = cellule.paragraphs[0].add_run(titre)
        r.bold = True
        r.font.size = Pt(9.5)
        # Place réservée pour la signature manuscrite.
        table.rows[1].cells[i].text = "\n\n\n"


def _parties(doc, site, preneur, nom_structure_defaut: str = "") -> None:
    """Le bloc « entre les soussignés », des deux côtés."""
    bailleur = site.identite_bailleur(nom_structure_defaut) if site else {}

    _section(doc, "Entre les soussignés")
    _ligne(doc, "La structure", bailleur.get("nom") or "—", gras=True)
    adresse = ", ".join(filter(None, [
        bailleur.get("adresse"),
        " ".join(filter(None, [bailleur.get("code_postal"), bailleur.get("ville")])),
    ]))
    _ligne(doc, "Adresse", adresse)
    _ligne(doc, "SIRET", bailleur.get("siret"))
    if bailleur.get("representant"):
        qualite = bailleur.get("qualite")
        _ligne(doc, "Représentée par", f"{bailleur['representant']}"
               + (f", {qualite}" if qualite else ""))
    _paragraphe(doc, "ci-après désignée « la structure »,", italique=True)

    _section(doc, "Et")
    _ligne(doc, "Le bénéficiaire", preneur.nom if preneur else "—", gras=True)
    if preneur:
        _ligne(doc, "Adresse", preneur.coordonnees)
        _ligne(doc, "SIRET", preneur.siret)
        _ligne(doc, "Représenté par", preneur.representant or preneur.contact_nom)
        _ligne(doc, "Téléphone", preneur.telephone)
        _ligne(doc, "Courriel", preneur.email)
    _paragraphe(doc, "ci-après désigné « le bénéficiaire ».", italique=True)


def _dates_reservation(doc, reservation: Reservation) -> None:
    occupations = sorted(
        [o for o in reservation.occupations or [] if o.statut != "annule"],
        key=lambda o: (o.date_jour, o.minute_debut),
    )
    if not occupations:
        _paragraphe(doc, "Aucune date retenue.")
        return
    if len(occupations) == 1:
        occ = occupations[0]
        _ligne(doc, "Date", f"{occ.date_jour.strftime('%d/%m/%Y')}, de {occ.plage}", gras=True)
        return

    _ligne(doc, "Période",
           f"du {occupations[0].date_jour.strftime('%d/%m/%Y')} "
           f"au {occupations[-1].date_jour.strftime('%d/%m/%Y')}", gras=True)
    _ligne(doc, "Nombre de dates", len(occupations))
    _tableau(
        doc, ["Date", "Horaire"],
        [[o.date_jour.strftime("%a %d/%m/%Y"), o.plage] for o in occupations],
        largeurs=[5.5, 4.0],
    )


# ---------------------------------------------------------------------------
# Les documents
# ---------------------------------------------------------------------------

def contrat(reservation: Reservation, nom_structure: str = "") -> Document:
    """Contrat de mise à disposition, ou CONVENTION quand c'est gratuit.

    La distinction n'est pas cosmétique : une occupation gratuite se
    formalise par une convention, et un document intitulé « contrat de
    location » sur une mise à disposition sans contrepartie ferait
    mauvaise figure devant un financeur comme devant un contrôle.
    """
    convention = reservation.gratuite
    doc = _document()
    site = reservation.espace.site if reservation.espace else None

    _titre(
        doc,
        "Convention de mise à disposition" if convention else "Contrat de mise à disposition",
        f"Référence {reservation.reference}"
        + (f" — {STATUTS_RESERVATION_LABELS.get(reservation.statut, '')}" if reservation.statut != "confirmee" else ""),
    )
    doc.add_paragraph()

    _parties(doc, site, reservation.preneur, nom_structure)

    _section(doc, "Article 1 — Objet")
    _paragraphe(doc, (
        "La structure met à disposition du bénéficiaire, à titre précaire et "
        "révocable, le local désigné ci-après, pour l'usage et aux dates indiqués."
    ))
    _ligne(doc, "Local", reservation.espace.chemin if reservation.espace else "—", gras=True)
    if reservation.espace:
        _ligne(doc, "Capacité maximale autorisée", reservation.espace.capacite_reglementaire)
        _ligne(doc, "Accessibilité PMR", "oui" if reservation.espace.pmr else "non")
    _ligne(doc, "Usage prévu", reservation.titre, gras=True)
    _ligne(doc, "Effectif annoncé", reservation.effectif)

    _section(doc, "Article 2 — Dates et horaires")
    _dates_reservation(doc, reservation)

    _section(doc, "Article 3 — Conditions financières")
    if convention:
        _paragraphe(doc, "La mise à disposition est consentie à titre gratuit.")
        if reservation.motif_gratuite:
            _ligne(doc, "Motif", reservation.motif_gratuite)
        if reservation.montant_calcule:
            _paragraphe(doc, (
                f"À titre indicatif, la valeur de cette mise à disposition au barème "
                f"en vigueur s'élève à {reservation.montant_calcule:.2f} €. Elle constitue "
                "une contribution volontaire en nature de la structure."
            ), taille=9.5, italique=True)
    else:
        lignes = [
            [l.get("libelle", ""), l.get("detail", ""), f"{l.get('montant', 0):.2f} €"]
            for l in reservation.detail
        ]
        if lignes:
            _tableau(doc, ["Désignation", "Détail", "Montant"], lignes, largeurs=[5.0, 8.5, 2.5])
        _ligne(doc, "Montant total", f"{reservation.montant_du:.2f} €", gras=True)
        if reservation.montant_manuel is not None:
            _paragraphe(doc, (
                f"Montant appliqué par dérogation au barème "
                f"({reservation.montant_calcule:.2f} €). Motif : "
                f"{reservation.motif_montant_manuel or '—'}."
            ), taille=9.5, italique=True)
        if site:
            _paragraphe(doc, site.mention_tva_affichee, taille=9.5, italique=True)

    if reservation.caution_montant:
        _ligne(doc, "Dépôt de garantie", f"{reservation.caution_montant:.2f} €")
    if reservation.acompte_montant:
        _ligne(doc, "Acompte à verser", f"{reservation.acompte_montant:.2f} €")

    _section(doc, "Article 4 — Assurance")
    preneur = reservation.preneur
    if preneur and preneur.assurance_rc_fin:
        _paragraphe(doc, (
            "Le bénéficiaire justifie d'une assurance en responsabilité civile "
            f"valable jusqu'au {preneur.assurance_rc_fin.strftime('%d/%m/%Y')}"
            + (f" (référence : {preneur.assurance_reference})" if preneur.assurance_reference else "")
            + "."
        ))
    else:
        _paragraphe(doc, (
            "ATTENTION : aucune attestation d'assurance en responsabilité civile "
            "n'est enregistrée. Elle doit être fournie avant toute remise de clés."
        ))

    if reservation.referent or reservation.espace and reservation.espace.detenteur_cle:
        _section(doc, "Article 5 — Accès aux locaux")
        _ligne(doc, "Référent ouverture / fermeture", reservation.referent)
        if reservation.espace and reservation.espace.detenteur_cle:
            _ligne(doc, "Détenteur des clés", reservation.espace.detenteur_cle)

    if reservation.conditions_particulieres:
        _section(doc, "Conditions particulières")
        _paragraphe(doc, reservation.conditions_particulieres)

    if reservation.espace and reservation.espace.habilitation_requise:
        _section(doc, "Prérequis")
        _paragraphe(doc, reservation.espace.habilitation_requise)

    doc.add_page_break()
    _section(doc, "Conditions générales")
    _paragraphe(
        doc,
        (site.conditions_generales if site and site.conditions_generales else CONDITIONS_GENERALES_DEFAUT),
        taille=9.5,
    )

    _signatures(doc, "Pour la structure", "Le bénéficiaire\n(lu et approuvé)")
    return doc


def etat_des_lieux(reservation: Reservation, sortie: bool = False) -> Document:
    """Le formulaire à remplir sur place, à deux, stylo à la main.

    Pensé pour le papier : des lignes vides plutôt que des champs pré-remplis,
    parce qu'un état des lieux se constate, il ne se génère pas.
    """
    doc = _document()
    _titre(
        doc,
        f"État des lieux — {'sortie' if sortie else 'entrée'}",
        f"{reservation.reference} · {reservation.preneur.nom if reservation.preneur else ''}",
    )
    doc.add_paragraph()

    _ligne(doc, "Local", reservation.espace.chemin if reservation.espace else "—", gras=True)
    _ligne(doc, "Bénéficiaire", reservation.preneur.nom if reservation.preneur else "—")
    _ligne(doc, "Usage", reservation.titre)
    _ligne(doc, "Date de l'état des lieux", "……… / ……… / ………         Heure : ………h………")

    _section(doc, "Constat")
    rubriques = [
        "Sols et revêtements", "Murs, plafonds, peintures", "Menuiseries, portes, fenêtres",
        "Éclairage et électricité", "Mobilier (tables, chaises)", "Matériel mis à disposition",
        "Sanitaires", "Propreté générale", "Clés et badges remis",
    ]
    _tableau(
        doc,
        ["Élément", "Bon", "Moyen", "Mauvais", "Observations"],
        [[r, "☐", "☐", "☐", ""] for r in rubriques],
        largeurs=[5.0, 1.3, 1.6, 1.9, 6.5],
    )

    _section(doc, "Dégradations constatées")
    _paragraphe(doc, "\n".join(["…" * 60] * 4), taille=10)

    if sortie:
        _section(doc, "Suites")
        _ligne(doc, "Dépôt de garantie",
               f"{reservation.caution_montant:.2f} €" if reservation.caution_montant else "—")
        _paragraphe(doc, (
            "☐ Restitution intégrale du dépôt de garantie\n"
            "☐ Retenue partielle — montant : ………………… € — motif : …………………………………………………\n"
            "☐ Retenue totale — motif : …………………………………………………………………………………………"
        ), taille=10)

    _signatures(doc, "Pour la structure", "Le bénéficiaire")
    return doc


def facture(reservation: Reservation, numero: str, nom_structure: str = "") -> Document:
    """Facture d'une mise à disposition payante."""
    doc = _document()
    site = reservation.espace.site if reservation.espace else None

    _titre(doc, "Facture", f"N° {numero} — {date.today().strftime('%d/%m/%Y')}")
    doc.add_paragraph()

    bailleur = site.identite_bailleur(nom_structure) if site else {}
    _section(doc, "Émetteur")
    _ligne(doc, "Structure", bailleur.get("nom") or "—", gras=True)
    _ligne(doc, "Adresse", ", ".join(filter(None, [
        bailleur.get("adresse"),
        " ".join(filter(None, [bailleur.get("code_postal"), bailleur.get("ville")])),
    ])))
    _ligne(doc, "SIRET", bailleur.get("siret"))

    _section(doc, "Destinataire")
    preneur = reservation.preneur
    _ligne(doc, "Bénéficiaire", preneur.nom if preneur else "—", gras=True)
    if preneur:
        _ligne(doc, "Adresse", preneur.coordonnees)
        _ligne(doc, "SIRET", preneur.siret)

    _section(doc, "Objet")
    _ligne(doc, "Référence de la mise à disposition", reservation.reference)
    _ligne(doc, "Local", reservation.espace.chemin if reservation.espace else "—")
    _ligne(doc, "Usage", reservation.titre)

    _section(doc, "Détail")
    lignes = [
        [l.get("libelle", ""), l.get("detail", ""), f"{l.get('montant', 0):.2f} €"]
        for l in reservation.detail
    ]
    if lignes:
        _tableau(doc, ["Désignation", "Détail", "Montant"], lignes, largeurs=[5.0, 8.5, 2.5])

    doc.add_paragraph()
    _ligne(doc, "TOTAL À RÉGLER", f"{reservation.montant_du:.2f} €", gras=True)
    if reservation.acompte_regle_le and reservation.acompte_montant:
        _ligne(doc, f"Acompte reçu le {reservation.acompte_regle_le.strftime('%d/%m/%Y')}",
               f"{reservation.acompte_montant:.2f} €")
        _ligne(doc, "RESTE À RÉGLER", f"{reservation.reste_du:.2f} €", gras=True)
    if site:
        _paragraphe(doc, site.mention_tva_affichee, taille=9.5, italique=True)
    if reservation.caution_montant:
        _paragraphe(doc, (
            f"Un dépôt de garantie de {reservation.caution_montant:.2f} € est demandé "
            "séparément ; il n'est pas encaissé au titre de la présente facture et sera "
            "restitué après état des lieux de sortie."
        ), taille=9.5, italique=True)
    return doc


def attestation(reservation: Reservation, nom_structure: str = "") -> Document:
    """Attestation d'occupation, que les associations demandent pour leurs
    propres dossiers de subvention."""
    doc = _document()
    site = reservation.espace.site if reservation.espace else None
    bailleur = site.identite_bailleur(nom_structure) if site else {}

    _titre(doc, "Attestation de mise à disposition de locaux")
    doc.add_paragraph()

    representant = bailleur.get("representant") or ""
    qualite = bailleur.get("qualite") or ""
    _paragraphe(doc, (
        f"Je soussigné{'e' if qualite.lower().startswith(('la ', 'présidente', 'directrice')) else ''} "
        f"{representant or '…………………………………'}"
        + (f", {qualite}" if qualite else "")
        + f" de {bailleur.get('nom') or '…………………………………'}, atteste que :"
    ))

    preneur = reservation.preneur
    _ligne(doc, "Bénéficiaire", preneur.nom if preneur else "—", gras=True)
    if preneur:
        _ligne(doc, "Adresse", preneur.coordonnees)

    occupations = sorted(
        [o for o in reservation.occupations or [] if o.statut != "annule"],
        key=lambda o: (o.date_jour, o.minute_debut),
    )
    if occupations:
        total_heures = sum(o.duree_minutes for o in occupations) / 60.0
        _paragraphe(doc, (
            f"a bénéficié de la mise à disposition du local « "
            f"{reservation.espace.nom if reservation.espace else '—'} » pour l'activité "
            f"« {reservation.titre} », sur {len(occupations)} date(s) entre le "
            f"{occupations[0].date_jour.strftime('%d/%m/%Y')} et le "
            f"{occupations[-1].date_jour.strftime('%d/%m/%Y')}, "
            f"soit un total de {total_heures:.1f} heures d'occupation."
        ))
        if reservation.gratuite:
            _paragraphe(doc, (
                "Cette mise à disposition a été consentie à titre gratuit"
                + (f" ({reservation.motif_gratuite})" if reservation.motif_gratuite else "")
                + (f". Sa valorisation au barème en vigueur s'élève à "
                   f"{reservation.montant_calcule:.2f} €."
                   if reservation.montant_calcule else ".")
            ))

    _paragraphe(doc, "Cette attestation est délivrée pour servir et valoir ce que de droit.")
    _signatures(doc, "Pour la structure", "")
    return doc


# ---------------------------------------------------------------------------
# Écriture et conversion
# ---------------------------------------------------------------------------

def ecrire(doc: Document, dossier: str, nom_fichier: str, *, pdf: bool = True) -> tuple[str, str | None]:
    """Écrit le .docx et tente le PDF. Retourne ``(docx, pdf ou None)``.

    L'absence de LibreOffice n'est pas une erreur : on rend le .docx, qui
    s'ouvre partout et reste modifiable avant signature.
    """
    os.makedirs(dossier, exist_ok=True)
    chemin_docx = os.path.join(dossier, nom_fichier if nom_fichier.endswith(".docx") else f"{nom_fichier}.docx")
    doc.save(chemin_docx)

    if not pdf:
        return chemin_docx, None
    try:
        from flask import current_app

        from app.activite.services.docx_utils import _try_docx_to_pdf

        return chemin_docx, _try_docx_to_pdf(current_app, chemin_docx)
    except Exception:  # noqa: BLE001 - le docx suffit, le PDF est un confort
        return chemin_docx, None


def nom_fichier(reservation: Reservation, genre: str) -> str:
    """Nom parlant et triable : « contrat_LOC-2026-0007_Gym-Volontaire »."""
    preneur = (reservation.preneur.nom if reservation.preneur else "") or "preneur"
    propre = "".join(c if c.isalnum() or c in "-_" else "-" for c in preneur).strip("-")[:40]
    return f"{genre}_{reservation.reference}_{propre}"
