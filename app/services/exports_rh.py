"""Exports RH de l'espace salarié : classeur global, relevé individuel, état de frais.

Trois sorties, pour une paie mensuelle :

- le **classeur Excel global** (direction, finance) : synthèse par salarié,
  totaux par secteur, détail des heures sup, des récupérations et des frais
  km, et en option un onglet par salarié ;
- le **relevé individuel** (page imprimable) : solde au début de la
  période, chaque mouvement, solde à la fin ;
- l'**état de frais kilométriques** mensuel (page imprimable).

Règles :
- le secteur d'une ligne est celui FIGÉ à sa saisie (colonne ``secteur``) :
  un changement de secteur ne réimpute pas le passé ;
- un solde « au » jour J compte les heures sup datées ≤ J et les
  récupérations ACCEPTÉES datées ≤ J (tous secteurs : c'est le compteur de
  la personne) ;
- jamais de coût horaire ni de salaire dans ces exports : ils circulent
  (cabinet comptable, impression) et le verrou des salaires ne doit pas se
  contourner par un fichier Excel.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from io import BytesIO

from app.extensions import db
from app.services.espace_salarie import STATUTS_RECUP, format_euros, format_minutes, libelle_lien


@dataclass
class Filtres:
    du: date
    au: date
    secteurs: list[str] = field(default_factory=list)
    salarie_ids: list[int] = field(default_factory=list)
    recup: bool = True               # heures sup et récupérations
    km: bool = True                  # frais kilométriques
    onglet_par_salarie: bool = False

    @property
    def libelle_periode(self) -> str:
        return f"du {self.du.strftime('%d/%m/%Y')} au {self.au.strftime('%d/%m/%Y')}"


# ---------------------------------------------------------------------------
# Périodes
# ---------------------------------------------------------------------------

def mois_courant(aujourd_hui: date | None = None) -> str:
    return (aujourd_hui or date.today()).strftime("%Y-%m")


def bornes_mois(mois: str | None) -> tuple[date, date]:
    """« 2026-10 » → (1er octobre, 31 octobre). Mois illisible : mois en cours."""
    try:
        annee, numero = (int(x) for x in (mois or "").split("-")[:2])
        debut = date(annee, numero, 1)
    except (ValueError, TypeError):
        debut = date.today().replace(day=1)
    suivant = (debut.replace(day=28) + timedelta(days=4)).replace(day=1)
    return debut, suivant - timedelta(days=1)


def periode_demandee(args) -> tuple[date, date, str]:
    """Période d'une requête : ``du``/``au`` explicites, sinon le mois ``mois``."""
    from app.services.espace_salarie import parse_date
    du, au = parse_date(args.get("du")), parse_date(args.get("au"))
    mois = (args.get("mois") or "").strip() or mois_courant()
    if du and au and du <= au:
        return du, au, ""
    du, au = bornes_mois(mois)
    return du, au, mois


# ---------------------------------------------------------------------------
# Lectures
# ---------------------------------------------------------------------------

def solde_au(salarie_id: int, jour: date) -> int:
    from app.models import DemandeRecuperation, HeureSupplementaire
    credit = db.session.query(db.func.coalesce(db.func.sum(HeureSupplementaire.minutes), 0)).filter(
        HeureSupplementaire.salarie_id == salarie_id, HeureSupplementaire.date_travail <= jour).scalar()
    debit = db.session.query(db.func.coalesce(db.func.sum(DemandeRecuperation.minutes), 0)).filter(
        DemandeRecuperation.salarie_id == salarie_id, DemandeRecuperation.statut == "acceptee",
        DemandeRecuperation.date_recuperation <= jour).scalar()
    return int(credit or 0) - int(debit or 0)


def _filtrer(q, modele, f: Filtres, colonne_date):
    q = q.filter(colonne_date >= f.du, colonne_date <= f.au)
    if f.secteurs:
        q = q.filter(modele.secteur.in_(f.secteurs))
    if f.salarie_ids:
        q = q.filter(modele.salarie_id.in_(f.salarie_ids))
    return q


def heures(f: Filtres) -> list:
    from app.models import HeureSupplementaire as H
    return _filtrer(H.query, H, f, H.date_travail).order_by(H.date_travail.asc(), H.id.asc()).all()


def recuperations(f: Filtres) -> list:
    """Demandes de la période (hors brouillons et annulées), quel que soit le statut."""
    from app.models import DemandeRecuperation as D
    q = _filtrer(D.query, D, f, D.date_recuperation).filter(D.statut.notin_(("brouillon", "annulee")))
    return q.order_by(D.date_recuperation.asc(), D.id.asc()).all()


def frais_km(f: Filtres) -> list:
    from app.models import FraisKilometrique as K
    return _filtrer(K.query, K, f, K.date_trajet).order_by(K.date_trajet.asc(), K.id.asc()).all()


def salaries_concernes(f: Filtres, lignes: list) -> list:
    """Salariés choisis ; sinon ceux présents sur la période (et du secteur
    choisi), plus ceux qui ont des lignes dans le filtre."""
    from app.models import Salarie
    if f.salarie_ids:
        return Salarie.query.filter(Salarie.id.in_(f.salarie_ids)).order_by(Salarie.nom, Salarie.prenom).all()
    ids = {ligne.salarie_id for ligne in lignes}
    presents = []
    for s in Salarie.query.order_by(Salarie.nom, Salarie.prenom).all():
        present = not ((s.date_entree and s.date_entree > f.au) or (s.date_sortie and s.date_sortie < f.du))
        if s.id in ids or (present and (not f.secteurs or s.secteur in f.secteurs)):
            presents.append(s)
    return presents


def synthese(f: Filtres) -> dict:
    """Une ligne par salarié, plus les totaux par secteur et les lignes de détail."""
    lignes_h = heures(f) if f.recup else []
    lignes_r = recuperations(f) if f.recup else []
    lignes_k = frais_km(f) if f.km else []
    salaries = salaries_concernes(f, lignes_h + lignes_r + lignes_k)
    veille = f.du - timedelta(days=1)

    par_salarie, par_secteur = [], {}
    for s in salaries:
        h = [x for x in lignes_h if x.salarie_id == s.id]
        r = [x for x in lignes_r if x.salarie_id == s.id]
        k = [x for x in lignes_k if x.salarie_id == s.id]
        ligne = {
            "salarie": s,
            "heures_sup": sum(x.minutes for x in h if not x.est_ajustement),
            "retraits": -sum(x.minutes for x in h if x.est_ajustement),
            "recup_prises": sum(x.minutes for x in r if x.statut == "acceptee"),
            "recup_en_attente": sum(x.minutes for x in r if x.statut in ("soumise", "transmise")),
            "solde_debut": solde_au(s.id, veille) if f.recup else None,
            "solde_fin": solde_au(s.id, f.au) if f.recup else None,
            "km_nb": len(k), "km": sum(x.distance_km for x in k),
            "km_montant": sum(x.montant_centimes for x in k),
            "km_a_imputer": sum(x.montant_centimes for x in k if x.depense_id is None),
        }
        par_salarie.append(ligne)

    def secteur_de(x):
        # Secteur figé à la saisie ; à défaut (ligne antérieure non reprise),
        # celui de la fiche.
        return (x.secteur or x.salarie.secteur or "Non affecté").strip() or "Non affecté"

    for x in lignes_h:
        d = par_secteur.setdefault(secteur_de(x), _total_vide())
        d["heures_sup" if not x.est_ajustement else "retraits"] += abs(x.minutes)
    for x in lignes_r:
        if x.statut == "acceptee":
            par_secteur.setdefault(secteur_de(x), _total_vide())["recup_prises"] += x.minutes
    for x in lignes_k:
        d = par_secteur.setdefault(secteur_de(x), _total_vide())
        d["km_nb"] += 1
        d["km"] += x.distance_km
        d["km_montant"] += x.montant_centimes
    return {"par_salarie": par_salarie, "par_secteur": dict(sorted(par_secteur.items())),
            "heures": lignes_h, "recuperations": lignes_r, "frais_km": lignes_k}


def _total_vide() -> dict:
    return {"heures_sup": 0, "retraits": 0, "recup_prises": 0, "km_nb": 0, "km": 0, "km_montant": 0}


def releve(salarie, du: date, au: date) -> dict:
    """Relevé individuel : mouvements datés et solde courant."""
    f = Filtres(du=du, au=au, salarie_ids=[salarie.id])
    debut = solde_au(salarie.id, du - timedelta(days=1))
    mouvements = []
    for h in heures(f):
        libelle = (f"Retrait par {'toi-même' if h.par_interesse else 'la direction'} : "
                   f"{h.commentaire_direction or ''}" if h.est_ajustement else f"Heures sup — {h.motif or ''}")
        lien = libelle_lien(h)
        mouvements.append({"date": h.date_travail, "libelle": libelle + (f" ({lien})" if lien else ""),
                           "minutes": h.minutes, "cle": (h.date_travail, 0, h.id)})
    autres = []
    for d in recuperations(f):
        if d.statut == "acceptee":
            mouvements.append({"date": d.date_recuperation,
                               "libelle": f"Récupération{(' — ' + d.motif) if d.motif else ''}",
                               "minutes": -d.minutes, "cle": (d.date_recuperation, 1, d.id)})
        else:
            autres.append(d)
    mouvements.sort(key=lambda m: m["cle"])
    courant = debut
    for m in mouvements:
        courant += m["minutes"]
        m["solde"] = courant
    return {"salarie": salarie, "du": du, "au": au, "solde_debut": debut, "solde_fin": courant,
            "mouvements": mouvements, "autres_demandes": autres, "frais_km": frais_km(f),
            "statuts": STATUTS_RECUP}


# ---------------------------------------------------------------------------
# Classeur Excel
# ---------------------------------------------------------------------------

def _heures_decimales(minutes) -> float | None:
    return None if minutes is None else round(int(minutes) / 60, 2)


def _euros(centimes) -> float:
    return round(int(centimes or 0) / 100, 2)


def _feuille(wb, titre: str, entetes: list[str], lignes: list[list], largeurs: list[int] | None = None,
             formats: dict[int, str] | None = None, intro: str | None = None):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(titre)
    depart = 1
    if intro:
        ws.cell(row=1, column=1, value=intro).font = Font(italic=True)
        depart = 3
    for col, entete in enumerate(entetes, start=1):
        c = ws.cell(row=depart, column=col, value=entete)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, ligne in enumerate(lignes, start=depart + 1):
        for col, valeur in enumerate(ligne, start=1):
            c = ws.cell(row=i, column=col, value=valeur)
            if formats and col in formats:
                c.number_format = formats[col]
    ws.freeze_panes = ws.cell(row=depart + 1, column=1)
    if lignes:
        ws.auto_filter.ref = f"A{depart}:{ws.cell(row=depart, column=len(entetes)).column_letter}{depart + len(lignes)}"
    for col, largeur in enumerate(largeurs or [], start=1):
        ws.column_dimensions[ws.cell(row=depart, column=col).column_letter].width = largeur
    return ws


FORMAT_H = '0.00" h"'
FORMAT_EUR = '#,##0.00" €"'


def _nom_onglet(nom: str, pris: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", nom).strip()[:28] or "Salarié"
    candidat, n = base, 2
    while candidat.lower() in pris:
        candidat = f"{base[:25]} {n}"
        n += 1
    pris.add(candidat.lower())
    return candidat


def classeur(f: Filtres) -> bytes:
    from openpyxl import Workbook
    from app.services.frais_km import format_taux, libelle_vehicule

    donnees = synthese(f)
    wb = Workbook()
    wb.remove(wb.active)
    intro = f"Export RH {f.libelle_periode}" + (f" — secteur(s) : {', '.join(f.secteurs)}" if f.secteurs else "")

    entetes = ["Salarié", "Secteur (fiche)"]
    if f.recup:
        entetes += ["Solde au début (h)", "Heures sup (h)", "Retraits (h)", "Récup prises (h)",
                    "Récup en attente (h)", "Solde à la fin (h)"]
    if f.km:
        entetes += ["Notes km", "Km", "Frais km (€)", "Dont non imputés (€)"]
    lignes, formats = [], {}
    for x in donnees["par_salarie"]:
        ligne = [x["salarie"].nom_complet, x["salarie"].secteur or ""]
        if f.recup:
            ligne += [_heures_decimales(x[k]) for k in ("solde_debut", "heures_sup", "retraits", "recup_prises",
                                                         "recup_en_attente", "solde_fin")]
        if f.km:
            ligne += [x["km_nb"], x["km"], _euros(x["km_montant"]), _euros(x["km_a_imputer"])]
        lignes.append(ligne)
    col = 3
    if f.recup:
        formats.update({c: FORMAT_H for c in range(col, col + 6)})
        col += 6
    if f.km:
        formats.update({col + 2: FORMAT_EUR, col + 3: FORMAT_EUR})
    _feuille(wb, "Synthèse", entetes, lignes, [28, 16] + [14] * (len(entetes) - 2), formats, intro)

    entetes_s, lignes_s, formats_s = ["Secteur"], [], {}
    if f.recup:
        entetes_s += ["Heures sup (h)", "Retraits (h)", "Récup prises (h)"]
    if f.km:
        entetes_s += ["Notes km", "Km", "Frais km (€)"]
    for secteur, t in donnees["par_secteur"].items():
        ligne = [secteur]
        if f.recup:
            ligne += [_heures_decimales(t["heures_sup"]), _heures_decimales(t["retraits"]),
                      _heures_decimales(t["recup_prises"])]
        if f.km:
            ligne += [t["km_nb"], t["km"], _euros(t["km_montant"])]
        lignes_s.append(ligne)
    if f.recup:
        formats_s.update({2: FORMAT_H, 3: FORMAT_H, 4: FORMAT_H})
    if f.km:
        formats_s[len(entetes_s)] = FORMAT_EUR
    _feuille(wb, "Par secteur", entetes_s, lignes_s, [24] + [14] * (len(entetes_s) - 1), formats_s,
             "Secteur figé au moment de la saisie de chaque ligne.")

    if f.recup:
        _feuille(wb, "Heures sup", ["Date", "Salarié", "Secteur", "Heures", "Nature", "Motif", "Activité liée"],
                 [[h.date_travail, h.salarie.nom_complet, h.secteur or "", _heures_decimales(h.minutes),
                   ("Retrait (par l'intéressé·e)" if h.par_interesse else "Retrait") if h.est_ajustement else "Déclaration",
                   (h.commentaire_direction if h.est_ajustement else h.motif) or "", libelle_lien(h) or ""]
                  for h in donnees["heures"]],
                 [12, 26, 16, 10, 22, 40, 36], {1: "DD/MM/YYYY", 4: FORMAT_H})
        _feuille(wb, "Récupérations", ["Date", "Salarié", "Secteur", "Heures", "Statut", "Notifiée",
                                        "Décidée par l'intéressé·e", "Motif", "Commentaire direction"],
                 [[d.date_recuperation, d.salarie.nom_complet, d.secteur or "", _heures_decimales(d.minutes),
                   STATUTS_RECUP.get(d.statut, {}).get("label", d.statut), "oui" if d.notifiee_le else "non",
                   "oui" if d.decision_par_interesse else "", d.motif or "", d.commentaire_direction or ""]
                  for d in donnees["recuperations"]],
                 [12, 26, 16, 10, 24, 10, 14, 30, 36], {1: "DD/MM/YYYY", 4: FORMAT_H})
    if f.km:
        _feuille(wb, "Frais km", ["Date", "Salarié", "Secteur", "Motif", "Véhicule", "CV", "Électrique", "Km",
                                   "Barème", "Taux €/km", "Montant (€)", "Dépense"],
                 [[k.date_trajet, k.salarie.nom_complet, k.secteur or "", k.motif,
                   libelle_vehicule(k.type_vehicule), k.puissance_fiscale, "oui" if k.electrique else "",
                   k.distance_km, k.annee_bareme, format_taux(k.taux_millieme), _euros(k.montant_centimes),
                   f"n° {k.depense_id}" if k.depense_id else "à imputer"]
                  for k in donnees["frais_km"]],
                 [12, 26, 16, 36, 16, 6, 10, 8, 8, 10, 12, 12], {1: "DD/MM/YYYY", 11: FORMAT_EUR})

    if f.onglet_par_salarie:
        pris = {ws.title.lower() for ws in wb.worksheets}
        for x in donnees["par_salarie"]:
            s = x["salarie"]
            r = releve(s, f.du, f.au) if f.recup else None
            lignes_ind = []
            if r:
                lignes_ind.append([f.du - timedelta(days=1), "Solde au début de la période", None,
                                   _heures_decimales(r["solde_debut"])])
                lignes_ind += [[m["date"], m["libelle"], _heures_decimales(m["minutes"]),
                                _heures_decimales(m["solde"])] for m in r["mouvements"]]
            if f.km:
                lignes_ind += [[k.date_trajet, f"Frais km — {k.motif} ({k.distance_km} km)", None, None,
                                _euros(k.montant_centimes)] for k in frais_km(Filtres(du=f.du, au=f.au, salarie_ids=[s.id]))]
            _feuille(wb, _nom_onglet(s.nom_complet, pris), ["Date", "Mouvement", "Heures", "Solde (h)", "Frais km (€)"],
                     lignes_ind, [12, 60, 10, 10, 12], {1: "DD/MM/YYYY", 3: FORMAT_H, 4: FORMAT_H, 5: FORMAT_EUR},
                     f"{s.nom_complet} — {f.libelle_periode}")

    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()

