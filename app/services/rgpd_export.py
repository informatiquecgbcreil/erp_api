"""Export « droit d'accès » RGPD (article 15).

Construit un classeur Excel rassemblant toutes les données qu'une
personne est en droit de réclamer : identité, droit à l'image,
orientations, inscriptions annuelles, présences aux activités, parcours
d'insertion, notes pédagogiques, pièces jointes et questionnaires.

La CNIL impose de fournir cette copie sous un mois : ce module permet
de répondre en un clic, sans compétence technique.
"""
from openpyxl import Workbook
from openpyxl.styles import Font

from app.models import (
    Evaluation,
    InscriptionAnnuelle,
    ObjectifSuivi,
    OrientationAccesDroit,
    Participant,
    PasseportNote,
    PasseportPieceJointe,
    PresenceActivite,
    QuestionnaireResponseGroup,
)

DROIT_IMAGE_LIBELLES = {
    "non_renseigne": "Non renseigné",
    "accepte": "Accepté",
    "refuse": "Refusé",
}


def _texte(value) -> str:
    if value is None:
        return ""
    return str(value)


def _ecrire_feuille(wb: Workbook, titre: str, entetes: list[str], lignes: list[list]):
    ws = wb.create_sheet(title=titre[:31])
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for ligne in lignes:
        ws.append([_texte(v) for v in ligne])
        for cell in ws[ws.max_row]:
            cell.data_type = "s"
    # Largeurs lisibles
    for idx, entete in enumerate(entetes, start=1):
        largeur = max([len(entete)] + [len(_texte(l[idx - 1])) for l in lignes] or [10])
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(largeur + 2, 12), 60)
    return ws


def construire_export_rgpd(participant: Participant) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    # --- Identité ---
    def _genre_lisible(fiche):
        """Le mot, pas le code : un export remis à la personne concernée
        doit se lire sans dictionnaire."""
        from app.services.genre import libelle_participant

        return libelle_participant(fiche)

    identite = [
        ("Nom", participant.nom),
        ("Prénom", participant.prenom),
        ("Date de naissance", participant.date_naissance),
        ("Année de naissance", participant.annee_naissance),
        ("Latitude", participant.latitude), ("Longitude", participant.longitude),
        ("Adresse géocodée", participant.geocode_query),
        ("Précision géographique", participant.geocode_precision),
        ("Score géographique", participant.geocode_score), ("Géocodage le", participant.geocoded_at),
        ("Code portail", participant.portail_code),
        ("Genre", _genre_lisible(participant)),
        ("Adresse", participant.adresse),
        ("Ville", participant.ville),
        ("Quartier", participant.quartier.nom if participant.quartier else ""),
        ("Email", participant.email),
        ("Téléphone", participant.telephone),
        ("Type de public", participant.type_public),
        ("Secteur de rattachement", participant.created_secteur),
        ("Fiche créée le", participant.created_at),
        ("Dernière modification", participant.updated_at),
        ("Pays d'origine", participant.pays_origine),
        ("Titre de séjour", participant.titre_sejour_type),
        ("Diplôme obtenu", participant.diplome_obtenu),
    ]
    _ecrire_feuille(wb, "Identité", ["Donnée", "Valeur"], [list(l) for l in identite])

    # --- Droit à l'image ---
    _ecrire_feuille(
        wb,
        "Droit à l'image",
        ["Statut", "Date du recueil", "Recueilli par"],
        [[
            DROIT_IMAGE_LIBELLES.get(participant.droit_image_statut, participant.droit_image_statut),
            participant.droit_image_date,
            participant.droit_image_recueilli_par,
        ]],
    )

    # --- Inscriptions annuelles (bulletins de rentrée) ---
    bulletins = (
        InscriptionAnnuelle.query.filter_by(participant_id=participant.id)
        .order_by(InscriptionAnnuelle.annee_scolaire.desc())
        .all()
    )
    _ecrire_feuille(
        wb,
        "Inscriptions annuelles",
        ["Année scolaire", "Date d'inscription", "Statut", "Adresse", "E-mail", "Téléphone",
         "Secteur qui fait venir", "Type d'inscription", "Foyer déclaré",
         "Ateliers souhaités", "Autres souhaits",
         "Bénévolat", "Bénévolat — pour quoi faire", "Bénévolat — disponibilités",
         "Règlement", "Montant dû", "Montant réglé", "Remarques"],
        [[
            b.libelle_annee,
            b.date_inscription,
            b.statut_label,
            b.adresse_complete,
            b.email,
            b.telephone,
            b.secteur_orienteur,
            b.type_inscription_label,
            " ; ".join(m.nom_complet for m in b.membres or []),
            " ; ".join(a.nom for a in b.ateliers or []),
            b.ateliers_libre,
            "Oui" if b.benevolat_souhaite else "Non",
            b.benevolat_mission,
            b.creneaux_benevolat_libelle if b.benevolat_souhaite else "",
            b.reglement_label,
            b.reglement_du,
            b.reglement_montant,
            b.commentaire,
        ] for b in bulletins],
    )

    # --- Orientations / accès aux droits ---
    orientations = (
        OrientationAccesDroit.query.filter_by(participant_id=participant.id)
        .order_by(OrientationAccesDroit.date_orientation.desc())
        .all()
    )
    _ecrire_feuille(
        wb,
        "Orientations",
        ["Date", "Domaine", "Demande", "Statut", "Urgence", "Suite prévue", "Note"],
        [[o.date_orientation, o.domaine, o.demande, o.statut, o.urgence, o.suite_prevue, o.note] for o in orientations],
    )

    # --- Présences aux activités ---
    presences = (
        PresenceActivite.query.filter_by(participant_id=participant.id)
        .order_by(PresenceActivite.id.desc())
        .all()
    )
    lignes_presences = []
    for pr in presences:
        session = pr.session if hasattr(pr, "session") else None
        atelier = getattr(session, "atelier", None) if session else None
        lignes_presences.append([
            getattr(session, "date_session", None) or getattr(session, "date", None),
            getattr(atelier, "nom", ""),
            getattr(session, "secteur", ""),
            pr.motif or pr.motif_autre or "",
        ])
    _ecrire_feuille(wb, "Présences activités", ["Date", "Atelier", "Secteur", "Motif"], lignes_presences)

    # --- Notes pédagogiques (passeport) ---
    notes = (
        PasseportNote.query.filter_by(participant_id=participant.id)
        .order_by(PasseportNote.created_at.desc())
        .all()
    )
    _ecrire_feuille(
        wb,
        "Notes pédagogiques",
        ["Date", "Catégorie", "Secteur", "Contenu"],
        [[n.created_at, n.categorie, n.secteur, n.contenu] for n in notes],
    )

    # --- Pièces jointes (passeport) ---
    pieces = (
        PasseportPieceJointe.query.filter_by(participant_id=participant.id)
        .order_by(PasseportPieceJointe.created_at.desc())
        .all()
    )
    _ecrire_feuille(
        wb,
        "Pièces jointes",
        ["Date", "Titre", "Nom du fichier", "Catégorie"],
        [[pj.created_at, pj.titre, pj.original_name, pj.categorie] for pj in pieces],
    )

    # --- Évaluations et objectifs suivis ---
    evaluations = Evaluation.query.filter_by(participant_id=participant.id).all()
    _ecrire_feuille(
        wb,
        "Évaluations",
        ["Identifiant", "Détail"],
        [[e.id, _resume_objet(e)] for e in evaluations],
    )
    objectifs = ObjectifSuivi.query.filter_by(participant_id=participant.id).all()
    _ecrire_feuille(
        wb,
        "Objectifs suivis",
        ["Identifiant", "Détail"],
        [[o.id, _resume_objet(o)] for o in objectifs],
    )

    # --- Questionnaires ---
    groupes = QuestionnaireResponseGroup.query.filter_by(participant_id=participant.id).all()
    _ecrire_feuille(
        wb,
        "Questionnaires",
        ["Identifiant", "Détail"],
        [[g.id, _resume_objet(g)] for g in groupes],
    )

    # --- Parcours insertion (tables dédiées) ---
    lignes_insertion = []
    for attr in ("insertion_profile", "insertion_parcours", "insertion_positionnements", "insertion_certifications"):
        objets = getattr(participant, attr, None)
        if objets is None:
            continue
        if not isinstance(objets, (list, tuple)):
            try:
                objets = list(objets)
            except TypeError:
                objets = [objets]
        for obj in objets:
            if obj is not None:
                lignes_insertion.append([attr, _resume_objet(obj)])
    _ecrire_feuille(wb, "Insertion", ["Rubrique", "Détail"], lignes_insertion)

    from app import models as m
    from app.extensions import db
    for title, model in [("Profil insertion", m.ParticipantInsertionProfile),
                         ("Parcours insertion", m.ParticipantInsertionParcours),
                         ("Positionnements", m.ParticipantInsertionPositionnement),
                         ("Certifications", m.ParticipantInsertionCertification),
                         ("Bénévolat", m.BenevoleHeures), ("Participation HART", m.HartEvaluation),
                         ("Cotisations", m.Cotisation), ("Inscriptions activités", m.InscriptionActivite),
                         ("Portail apprenant", m.PortailAttempt), ("Défis", m.DefiTransition)]:
        columns = [c.name for c in model.__table__.columns if c.name != "participant_id"]
        rows = model.query.filter_by(participant_id=participant.id).all()
        _ecrire_feuille(wb, title, columns, [[_lisible(row, name) for name in columns] for row in rows])
    payments = m.Paiement.query.join(m.Cotisation).filter(m.Cotisation.participant_id == participant.id).all()
    columns = [c.name for c in m.Paiement.__table__.columns]
    _ecrire_feuille(wb, "Règlements", columns, [[_lisible(row, name) for name in columns] for row in payments])

    # --- Audit 3.7 : ce qui manquait ---
    _feuilles_complementaires(wb, participant)
    return wb


def _feuilles_complementaires(wb: Workbook, participant: Participant) -> None:
    from app import models as m
    from app.extensions import db

    # Réponses aux questionnaires, question par question.
    lignes = []
    for groupe in QuestionnaireResponseGroup.query.filter_by(participant_id=participant.id).all():
        questionnaire = db.session.get(m.Questionnaire, groupe.questionnaire_id)
        for reponse in m.QuestionResponse.query.filter_by(response_group_id=groupe.id).all():
            question = db.session.get(m.Question, reponse.question_id)
            valeur = reponse.value_text if reponse.value_text is not None else (
                reponse.value_number if reponse.value_number is not None else reponse.value_json)
            lignes.append([groupe.created_at, getattr(questionnaire, "nom", ""),
                           getattr(question, "label", ""), valeur])
    _ecrire_feuille(wb, "Réponses questionnaires", ["Date", "Questionnaire", "Question", "Réponse"], lignes)

    # Foyer et ses cotisations (sans l'identité des autres membres).
    foyer = db.session.get(m.Foyer, participant.foyer_id) if participant.foyer_id else None
    lignes = []
    if foyer is not None:
        lignes.append(["Foyer", foyer.nom, "", ""])
        for cotisation in m.Cotisation.query.filter_by(foyer_id=foyer.id).all():
            lignes.append([f"Cotisation du foyer {cotisation.annee_scolaire}", cotisation.type_cotisation,
                           cotisation.montant_du, cotisation.date_reference])
    _ecrire_feuille(wb, "Foyer", ["Rubrique", "Libellé", "Montant", "Date"], lignes)

    # Présence sur le bulletin d'inscription d'un proche.
    lignes = []
    for membre in m.InscriptionAnnuelleMembre.query.filter_by(participant_id=participant.id).all():
        bulletin = db.session.get(m.InscriptionAnnuelle, membre.inscription_id)
        if bulletin is not None and bulletin.participant_id == participant.id:
            continue
        lignes.append([getattr(bulletin, "libelle_annee", ""), f"bulletin n° {membre.inscription_id}",
                       membre.nom, membre.prenom, membre.date_naissance, membre.lien_filiation])
    _ecrire_feuille(wb, "Bulletin d'un proche",
                    ["Année", "Bulletin", "Nom inscrit", "Prénom inscrit", "Date de naissance", "Lien"], lignes)

    # Sommes reçues de la personne.
    lignes = [[e.date_encaissement, e.montant, m.MODES_PAIEMENT_LABELS.get(e.mode, e.mode), e.libelle,
               "contre-passation" if e.origine_id else ""]
              for e in m.Encaissement.query.filter_by(participant_id=participant.id).all()]
    _ecrire_feuille(wb, "Encaissements", ["Date", "Montant", "Mode", "Libellé", "Nature"], lignes)

    # Destinataires des orientations.
    lignes = [[o.date_orientation, o.demande, o.partenaire.nom if o.partenaire else ""]
              for o in OrientationAccesDroit.query.filter_by(participant_id=participant.id).all()]
    _ecrire_feuille(wb, "Orientations - destinataires", ["Date", "Demande", "Partenaire destinataire"], lignes)

    # Matériel consommé lors des présences.
    lignes = [[c.created_at, c.materiel_nom_snapshot, c.quantite, c.duree_minutes_snapshot, c.kwh_snapshot]
              for c in m.PresenceMaterielConsommation.query.filter_by(participant_id=participant.id).all()]
    _ecrire_feuille(wb, "Matériel utilisé", ["Date", "Matériel", "Quantité", "Durée (min)", "kWh"], lignes)


def _lisible(row, nom):
    """Valeur d'une colonne ; pour une clé vers un référentiel, son libellé
    (« Carte de séjour » et non « titre_sejour_type_id = 1 »)."""
    valeur = getattr(row, nom, None)
    colonne = row.__table__.columns.get(nom)
    if valeur is None or colonne is None or not colonne.foreign_keys:
        return valeur
    from app.extensions import db
    from sqlalchemy import select
    cle = next(iter(colonne.foreign_keys))
    table = cle.column.table
    libelles = [c for c in ("libelle", "label", "nom", "titre", "code") if c in table.c]
    if not libelles or table.name in {"participant", "user"}:
        return valeur
    try:
        trouve = db.session.execute(select(table.c[libelles[0]]).where(cle.column == valeur)).scalar()
    except Exception:  # noqa: BLE001
        return valeur
    return f"{trouve} (n° {valeur})" if trouve else valeur


def fichiers_de_la_personne(participant: Participant) -> list[tuple[str, str]]:
    """(nom dans l'archive, chemin sur le disque) des fichiers qui concernent
    la personne : pièces du passeport et signatures. Seuls les fichiers
    présents dans les dossiers de l'application sont repris."""
    from pathlib import Path
    from flask import current_app
    racines = [Path(current_app.instance_path).resolve(), Path(current_app.config["APP_UPLOAD_DIR"]).resolve()]

    def autorise(chemin):
        try:
            cible = Path(chemin).resolve()
        except (OSError, ValueError):
            return False
        return cible.is_file() and any(cible.is_relative_to(r) for r in racines)

    fichiers = []
    for piece in PasseportPieceJointe.query.filter_by(participant_id=participant.id).all():
        if piece.file_path and autorise(piece.file_path):
            nom = Path(piece.original_name or piece.file_path).name
            fichiers.append((f"pieces-jointes/{piece.id}-{nom}", piece.file_path))
    for presence in PresenceActivite.query.filter_by(participant_id=participant.id).all():
        if presence.signature_path and autorise(presence.signature_path):
            fichiers.append((f"signatures/presence-{presence.id}{Path(presence.signature_path).suffix}",
                             presence.signature_path))
    return fichiers


def construire_archive_rgpd(participant: Participant):
    """ZIP : le classeur complet et les fichiers eux-mêmes (audit 3.7)."""
    import zipfile
    from io import BytesIO
    classeur = BytesIO()
    construire_export_rgpd(participant).save(classeur)
    sortie = BytesIO()
    with zipfile.ZipFile(sortie, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("donnees.xlsx", classeur.getvalue())
        fichiers = fichiers_de_la_personne(participant)
        for nom, chemin in fichiers:
            archive.write(chemin, nom)
        archive.writestr("LISEZMOI.txt", (
            "Copie des données personnelles (RGPD, article 15).\n"
            "donnees.xlsx : toutes les informations enregistrées, une feuille par rubrique.\n"
            f"{len(fichiers)} fichier(s) joint(s) : pièces du passeport et signatures d'émargement.\n"))
    sortie.seek(0)
    return sortie


def _resume_objet(obj) -> str:
    """Représentation lisible de toutes les colonnes simples d'une ligne."""
    morceaux = []
    for colonne in obj.__table__.columns:
        nom = colonne.name
        if nom in {"id", "participant_id"} or nom.endswith("_id"):
            continue
        valeur = getattr(obj, nom, None)
        if valeur not in (None, ""):
            morceaux.append(f"{nom}: {valeur}")
    return " | ".join(morceaux)
