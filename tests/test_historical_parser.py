"""Parser regression cases and an opt-in integration test of the real XLSX.

The private source workbook must never be copied into the repository. Set
HISTORICAL_XLSX_PATH to run the integration test against its original path.
"""
import hashlib
import os
from datetime import date, datetime

import pytest
from openpyxl import Workbook

from app.ateliers.historical_parser import parse_workbook


def _register(tmp_path, days=(5,), *, month=datetime(2026, 1, 1),
              weekdays=None, slots=None, rows=None, sheet_name="ATELIER",
              activity="ATELIER", explicit_months=None, total=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.cell(2, 1, "ETAT DE PRESENCES 2026")
    ws.cell(2, 9, month)
    headers = ["NOMS", "PRENOMS", "ANNEE NAISSANCE", "AGE", "SEXE", "QUARTIER", "ACTIVITE"]
    for col, label in enumerate(headers, 1):
        ws.cell(6, col, label)
    for offset, day in enumerate(days):
        col = 9 + offset
        if day is not None:
            ws.cell(6, col, day)
        if weekdays:
            ws.cell(4, col, weekdays[offset])
        ws.cell(5, col, slots[offset] if slots else "M")
    for offset, value in (explicit_months or {}).items():
        ws.cell(2, 9 + offset, value)
    total_col = 9 + len(days)
    ws.cell(6, total_col, "TOTAL")
    rows = rows or [("DUPONT", "Jean", 1985, [1] * len(days))]
    for row, (nom, prenom, birth, presences) in enumerate(rows, 7):
        for col, value in {1: nom, 2: prenom, 3: birth, 6: "ROUHER", 7: activity}.items():
            if value is not None:
                ws.cell(row, col, value)
        for offset, value in enumerate(presences):
            if value is not None:
                ws.cell(row, 9 + offset, value)
        ws.cell(row, total_col, sum(value == 1 for value in presences))
    footer = 7 + len(rows)
    ws.cell(footer, 1, "TOTAL")
    ws.cell(footer, total_col, total if total is not None else sum(sum(value == 1 for value in row[3]) for row in rows))
    path = tmp_path / "stats_2026.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_standard_sheet_retains_raw_values_and_cell_provenance(tmp_path):
    path = _register(tmp_path, weekdays=["L"])
    stage = parse_workbook(path)
    assert stage["counts"]["participants_source"] == 1
    assert stage["sessions"][0]["date_session"] == "2026-01-05"
    assert stage["people"][0]["raw"]["prenom"] == "Jean"
    assert stage["people"][0]["field_cells"]["prenom"] == "B7"
    assert stage["attendance"][0]["source_cell"] == "I7"
    assert stage["sheets"][0]["totals"]["status"] == "coherent"


def test_implicit_month_change_uses_order_and_checks_weekday(tmp_path):
    days = (12, 14, 19, 21, 26, 27, 2, 4, 9, 11)
    stage = parse_workbook(_register(tmp_path, days, weekdays=["L", "ME", "L", "ME", "L", "M", "L", "ME", "L", "ME"]))
    assert [s["date_session"] for s in stage["sessions"]][-4:] == ["2026-02-02", "2026-02-04", "2026-02-09", "2026-02-11"]
    assert stage["counts"]["sessions_review"] == 0
    assert any(a["code"] == "implicit_month_rollover" and not a["blocking"] for a in stage["anomalies"])


def test_explicit_month_and_year_change(tmp_path):
    stage = parse_workbook(_register(tmp_path, (26, 2), month="décembre 2026", explicit_months={1: "janvier 2027"}))
    assert [s["date_session"] for s in stage["sessions"]] == ["2026-12-26", "2027-01-02"]


def test_full_dates_and_unformatted_excel_serials(tmp_path):
    stage = parse_workbook(_register(tmp_path, (46027, date(2026, 1, 6))))
    assert [s["date_session"] for s in stage["sessions"]] == ["2026-01-05", "2026-01-06"]


def test_multiple_same_day_sessions_keep_slots_and_columns(tmp_path):
    stage = parse_workbook(_register(tmp_path, (5, 5, 5), slots=["M", "AM", "AM"]))
    assert len({s["key"] for s in stage["sessions"]}) == 3
    assert {s["date_session"] for s in stage["sessions"]} == {"2026-01-05"}
    assert [s["source_slot"] for s in stage["sessions"]] == ["M", "AM", "AM"]
    assert len(stage["attendance"]) == 3
    assert all("heure_debut" not in s for s in stage["sessions"])


@pytest.mark.parametrize("token", [1, True, "1", "x", "P", "présent", "OUI", "o"])
def test_documented_presence_tokens(tmp_path, token):
    stage = parse_workbook(_register(tmp_path, rows=[("DUPONT", "Jean", 1985, [token])]))
    assert len(stage["attendance"]) == 1


@pytest.mark.parametrize("token", [None, 0, False, "absent", "non"])
def test_absence_tokens(tmp_path, token):
    stage = parse_workbook(_register(tmp_path, rows=[("DUPONT", "Jean", 1985, [token])]))
    assert stage["attendance"] == []


def test_unknown_presence_is_blocking_and_does_not_guess_numeric_count(tmp_path):
    stage = parse_workbook(_register(tmp_path, rows=[("DUPONT", "Jean", 1985, [2])]))
    anomaly = next(a for a in stage["anomalies"] if a["code"] == "unknown_presence_token")
    assert anomaly["source_cell"] == "I7"
    assert anomaly["blocking"]
    assert stage["sheets"][0]["totals"]["status"] == "uninterpretable"


def test_complementary_and_duplicate_source_rows_are_never_discarded(tmp_path):
    stage = parse_workbook(_register(tmp_path, (5, 12, 19), rows=[
        ("DUPONT", "Jean", 1985, [1, None, 1]),
        ("dupont", "jean", None, [None, 1, None]),
        ("DUPONT", "Jean", 1985, [1, None, None]),
    ]))
    assert len(stage["people"]) == 3
    assert len(stage["attendance"]) == 4
    assert stage["people"][1]["raw"]["annee_naissance"] is None


def test_business_activity_name_overrides_truncated_tab(tmp_path):
    name = "ACCES AUX DROITS NUMERIQUE ATELIER INDIVIDUELLE"
    stage = parse_workbook(_register(tmp_path, sheet_name="ACCES AUX DROITS NUMERIQUE ATEL", activity=name))
    assert stage["activities"][0]["name"] == name
    assert stage["activities"][0]["source_sheet"] == "ACCES AUX DROITS NUMERIQUE ATEL"
    assert stage["activities"][0]["name_source"] == "ACTIVITE"


def test_named_and_automatically_detected_non_activity_sheets(tmp_path):
    path = _register(tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for name in ("TOTAL JANV A dec 2026", "ADULTES", "ADULTES SAVEURS", "Synthèse 2026", "Paramètres", "Vide"):
        sheet = wb.create_sheet(name)
        if name == "Paramètres":
            sheet.cell(1, 1, "option inconnue")
    wb.save(path)
    stage = parse_workbook(path)
    kinds = {s["name"]: s["classification"] for s in stage["sheets"]}
    assert kinds["TOTAL JANV A dec 2026"] == "summary"
    assert kinds["ADULTES"] == kinds["ADULTES SAVEURS"] == "technical"
    assert kinds["Synthèse 2026"] == "summary"
    assert kinds["Vide"] == "technical"
    assert kinds["Paramètres"] == "unrecognized"
    assert len(stage["activities"]) == 1


def test_missing_date_retains_presence_for_review(tmp_path):
    stage = parse_workbook(_register(tmp_path, (None,), weekdays=["S"]))
    assert stage["sessions"][0]["status"] == "REVIEW"
    assert stage["sessions"][0]["date_session"] is None
    assert len(stage["attendance"]) == 1


def test_weekday_contradiction_is_review_not_silent_month_repair(tmp_path):
    stage = parse_workbook(_register(tmp_path, (3,), weekdays=["M"]))
    session = stage["sessions"][0]
    assert session["status"] == "REVIEW"
    assert session["date_session"] is None
    assert session["candidate_date"] == "2026-01-03"
    assert "2026-02-03" in session["candidate_dates"]


def test_bad_april_day_does_not_shift_all_following_columns_to_may(tmp_path):
    stage = parse_workbook(_register(tmp_path, (31, 2, 7), month=datetime(2026, 4, 1), weekdays=["M", "J", "M"]))
    assert [s["date_session"] for s in stage["sessions"]] == [None, "2026-04-02", "2026-04-07"]
    assert stage["sessions"][0]["candidate_dates"] == ["2026-03-31"]


def test_missing_identity_preserves_all_source_attendance_for_resolution(tmp_path):
    stage = parse_workbook(_register(tmp_path, rows=[("DUPONT", None, 1985, [1])]))
    assert len(stage["people"]) == 1
    assert len(stage["attendance"]) == 1
    row = stage["sheets"][0]["review_rows"][0]
    assert row["raw"]["nom"] == "DUPONT"
    assert row["attendance"][0]["source_cell"] == "I7"
    assert stage["counts"]["presence_cells"] == 1


def test_excel_total_mismatch_exposes_both_numbers(tmp_path):
    stage = parse_workbook(_register(tmp_path, total=8))
    total = stage["sheets"][0]["totals"]
    assert (total["status"], total["calculated"], total["excel_total"]) == ("difference", 1, 8)


def test_parser_is_deterministic_and_never_modifies_input(tmp_path):
    path = _register(tmp_path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    first, second = parse_workbook(path), parse_workbook(path)
    assert first == second
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_real_historical_workbook_integration():
    path = os.environ.get("HISTORICAL_XLSX_PATH")
    if not path:
        pytest.skip("Set HISTORICAL_XLSX_PATH to the private STATS_2026_par_activite.xlsx")
    stage = parse_workbook(path)
    assert stage["counts"]["sheets"] == 69
    assert stage["counts"]["activities"] == 66
    assert stage["counts"]["ignored_sheets"] == 3
    assert stage["counts"]["sessions"] == 880
    assert stage["counts"]["participants_source"] == 1464
    assert stage["counts"]["presence_cells"] == 7031
    assert len(stage["attendance"]) == 7031
    assert sum(len(r.get("attendance", [])) for s in stage["sheets"] for r in s["review_rows"]) == 40
    assert all(s["totals"]["status"] == "coherent" for s in stage["sheets"] if s["classification"] == "activity")
    assert all(s["date_session"] is None for s in stage["sessions"] if s["status"] == "REVIEW")
