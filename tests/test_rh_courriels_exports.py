"""Espace salarié : e-mails au fil de l'eau, secteur figé, exports RH.

- les e-mails partent vers la bonne personne, jamais vers l'auteur de
  l'action ; un serveur absent ou en panne ne bloque rien (file + réessais) ;
  chacun peut couper ses e-mails ; aucune donnée sensible dans le corps ;
- le secteur d'une ligne est celui du moment de la saisie ;
- le classeur Excel, le relevé et l'état de frais donnent les bons chiffres,
  aux bonnes personnes, sans jamais un coût horaire.
"""
import datetime as dt
import io
import uuid

import pytest

from tests.conftest import ADMIN_EMAIL, signature_tracee


def _suf():
    return uuid.uuid4().hex[:6]


def _compte(app, nom, role, secteur=None):
    email = f"{uuid.uuid4().hex[:8]}@example.org"
    with app.app_context():
        from app.extensions import db
        from app.models import Role, User
        u = User(email=email, nom=nom, secteur_assigne=secteur)
        u.set_password("pw-test-123")
        u.roles.append(Role.query.filter_by(code=role).first())
        db.session.add(u)
        db.session.commit()
        return email, u.id


def _client(app, email):
    c = app.test_client()
    assert c.post("/", data={"email": email, "password": "pw-test-123"}).status_code == 302
    return c


def _salarie(app, role="animateur", secteur="Numérique"):
    suf = _suf()
    email, uid = _compte(app, f"Sal{suf} Test{suf}", role, secteur)
    with app.app_context():
        from app.extensions import db
        from app.models import Salarie
        aujourd_hui = dt.date.today()
        s = Salarie(nom=f"Test{suf}", prenom=f"Sal{suf}", secteur=secteur, user_id=uid,
                    date_entree=dt.date(aujourd_hui.year - 2, 1, 1), date_sortie=dt.date(aujourd_hui.year + 1, 12, 31))
        db.session.add(s)
        db.session.commit()
        sid = s.id
    return _client(app, email), sid, uid, email


@pytest.fixture(autouse=True)
def file_vide(app):
    """La base de test est partagée : on repart d'une file d'e-mails vide."""
    with app.app_context():
        from app.extensions import db
        from app.models import CourrielRh
        CourrielRh.query.delete()
        db.session.commit()
    yield


@pytest.fixture()
def smtp(app, monkeypatch):
    """Serveur de mail « configuré » ; les envois sont capturés."""
    from app.services import instance_settings, notifications
    boite = []
    monkeypatch.setattr(instance_settings, "resolve_mail_settings",
                        lambda config: {"host": "smtp.test", "sender": "erp@test", "port": 587,
                                        "use_tls": False, "username": "", "password": ""})
    monkeypatch.setattr(notifications, "_envoyer_texte", lambda to, subject, body: boite.append((to, subject, body)))
    return boite


def _pour(boite, email):
    return [m for m in boite if m[0] == email]


# ---------------------------------------------------------------------------
# E-mails
# ---------------------------------------------------------------------------

def test_le_circuit_previent_la_bonne_personne_a_chaque_etape(app, admin_client, smtp):
    client, sid, uid, email_sal = _salarie(app)
    email_a, _ = _compte(app, f"Assist {_suf()}", "assistant_direction")
    assistant = _client(app, email_a)

    client.post("/salarie/recuperations/demandes", data={
        "date_recuperation": dt.date.today().isoformat(), "duree": "2h", "signature_data": signature_tracee()})
    assert _pour(smtp, email_a), "l'assistant·e doit être prévenu·e de la demande signée"
    assert not _pour(smtp, email_sal), "personne n'est prévenu de sa propre action"
    with app.app_context():
        from app.models import DemandeRecuperation
        did = DemandeRecuperation.query.filter_by(salarie_id=sid).one().id

    assistant.post(f"/salarie/equipe/recuperations/{did}/transmettre", data={"signature_data": signature_tracee()})
    assert any("à décider" in m[1] for m in _pour(smtp, ADMIN_EMAIL))

    admin_client.post(f"/salarie/equipe/recuperations/{did}/decider", data={
        "decision": "refuser", "commentaire": "Motif très confidentiel", "signature_data": signature_tracee()})
    decision = _pour(smtp, email_sal)
    assert len(decision) == 1 and "refusée" in decision[0][1]
    assert "Motif très confidentiel" not in decision[0][2], "le commentaire se lit dans l'application"
    assert f"/salarie/recuperations/demandes/{did}" in decision[0][2]
    with app.app_context():
        from app.models import CourrielRh
        assert CourrielRh.query.filter_by(objet_id=did, envoye_le=None).count() == 0


def test_desinscription_et_retrait_d_heures(app, admin_client, smtp):
    client, sid, uid, email_sal = _salarie(app)
    client.post("/salarie/recuperations/heures", data={
        "date_travail": dt.date.today().isoformat(), "duree": "1h", "motif": "réunion"})
    with app.app_context():
        from app.models import HeureSupplementaire
        h1 = HeureSupplementaire.query.filter_by(salarie_id=sid).one().id
    admin_client.post(f"/salarie/equipe/heures/{h1}/retirer", data={"commentaire": "Déjà payée"})
    assert any("retirées" in m[1] for m in _pour(smtp, email_sal))

    assert client.post("/salarie/preferences/courriels", data={"actif": "0"}).status_code == 302
    client.post("/salarie/recuperations/heures", data={
        "date_travail": dt.date.today().isoformat(), "duree": "1h", "motif": "bis"})
    with app.app_context():
        from app.models import HeureSupplementaire
        h2 = HeureSupplementaire.query.filter_by(salarie_id=sid, est_ajustement=False).order_by(
            HeureSupplementaire.id.desc()).first().id
    avant = len(_pour(smtp, email_sal))
    admin_client.post(f"/salarie/equipe/heures/{h2}/retirer", data={"commentaire": "Doublon"})
    assert len(_pour(smtp, email_sal)) == avant, "désinscrit : plus d'e-mail"
    assert "Recevoir à nouveau" in client.get("/dashboard?espace=salarie").get_data(as_text=True)
    client.get("/dashboard?espace=activites")


def test_serveur_absent_ou_en_panne_ne_bloque_rien(app, admin_client, monkeypatch):
    from app.services import instance_settings, notifications
    client, sid, uid, email_sal = _salarie(app)
    # 1) Pas de serveur configuré : l'action passe, l'e-mail attend sans tentative.
    monkeypatch.setattr(instance_settings, "resolve_mail_settings",
                        lambda config: {"host": "", "sender": "", "port": 587, "use_tls": False,
                                        "username": "", "password": ""})
    client.post("/salarie/recuperations/demandes", data={
        "date_recuperation": dt.date.today().isoformat(), "duree": "1h", "signature_data": signature_tracee()})
    with app.app_context():
        from app.models import CourrielRh, DemandeRecuperation
        did = DemandeRecuperation.query.filter_by(salarie_id=sid).one().id
        assert DemandeRecuperation.query.get(did).statut == "soumise"
        en_file = CourrielRh.query.filter_by(objet_id=did, objet_type="demande_recuperation").all()
        assert en_file and all(c.envoye_le is None and c.tentatives == 0 for c in en_file)
    # 2) Serveur configuré mais en panne : l'action passe, la tentative est comptée.
    monkeypatch.setattr(instance_settings, "resolve_mail_settings",
                        lambda config: {"host": "smtp.test", "sender": "erp@test", "port": 587,
                                        "use_tls": False, "username": "", "password": ""})

    def panne(to, subject, body):
        raise OSError("serveur injoignable")
    monkeypatch.setattr(notifications, "_envoyer_texte", panne)
    r = admin_client.post(f"/salarie/equipe/recuperations/{did}/decider",
                          data={"decision": "accepter", "signature_data": signature_tracee()})
    assert r.status_code == 302
    with app.app_context():
        from app.models import CourrielRh, DemandeRecuperation
        assert DemandeRecuperation.query.get(did).statut == "acceptee"
        echecs = CourrielRh.query.filter_by(objet_id=did, objet_type="demande_recuperation", envoye_le=None).all()
        assert echecs and all(c.tentatives >= 1 and "injoignable" in c.derniere_erreur for c in echecs)
    page = admin_client.get("/admin/notifications").get_data(as_text=True)
    assert "E-mails de l" in page and "serveur injoignable" in page
    # 3) Le serveur revient : la relance vide la file.
    envoyes = []
    monkeypatch.setattr(notifications, "_envoyer_texte", lambda to, subject, body: envoyes.append(to))
    admin_client.post("/admin/notifications/courriels-rh/relancer")
    with app.app_context():
        from app.models import CourrielRh
        assert CourrielRh.query.filter_by(objet_id=did, objet_type="demande_recuperation", envoye_le=None).count() == 0
    assert email_sal in envoyes


def test_document_partage_previent_le_destinataire(app, admin_client, smtp):
    client, sid, uid, email_sal = _salarie(app)
    with app.app_context():
        from app.models import TypeDocumentRh
        type_id = TypeDocumentRh.query.filter_by(code="fiche_paie").one().id
    admin_client.post("/salarie/documents/deposer", content_type="multipart/form-data", data={
        "type_id": type_id, "salarie_id": sid, "fichier": (io.BytesIO(b"%PDF-1.4"), "paie.pdf")})
    mails = _pour(smtp, email_sal)
    assert len(mails) == 1 and "Fiche de paie" in mails[0][1]
    assert "paie.pdf" not in mails[0][2], "le document ne voyage jamais par e-mail"


def test_recapitulatifs_rh_et_frais_km(app):
    from app.services import notifications
    _, sid, _, _ = _salarie(app)
    with app.app_context():
        from app.extensions import db
        from app.models import DemandeRecuperation, FraisKilometrique
        vieille = dt.datetime.now() - dt.timedelta(days=10)
        db.session.add(DemandeRecuperation(salarie_id=sid, date_recuperation=dt.date.today(), minutes=60,
                                           statut="soumise", created_at=vieille, updated_at=vieille))
        db.session.add(FraisKilometrique(salarie_id=sid, date_trajet=dt.date.today(), annee=dt.date.today().year,
                                         annee_bareme=dt.date.today().year, type_vehicule="voiture",
                                         puissance_fiscale=5, distance_km=10, motif="x", montant_centimes=636))
        db.session.commit()
        lignes_rh = notifications._lignes_rh_a_traiter(3, dt.date.today())
        assert any("à transmettre" in ligne for ligne in lignes_rh)
        assert any("pas encore passée" in ligne for ligne in notifications._lignes_frais_km(None, dt.date.today()))
        assert {"rh_a_traiter", "frais_km"} <= set(notifications.TYPES_NOTIFICATION)


# ---------------------------------------------------------------------------
# Secteur figé et exports
# ---------------------------------------------------------------------------

def _jeu_export(app, secteur):
    """Un salarié, un mois passé : 5 h sup, 1 h retirée, 2 h récupérées, 2 trajets."""
    client, sid, uid, email = _salarie(app, secteur=secteur)
    debut = dt.date(dt.date.today().year - 1, 3, 1)
    with app.app_context():
        from app.extensions import db
        from app.models import DemandeRecuperation, FraisKilometrique, HeureSupplementaire
        h = HeureSupplementaire(salarie_id=sid, date_travail=debut + dt.timedelta(days=4), minutes=300,
                                motif="fête", secteur=secteur)
        db.session.add(h)
        db.session.flush()
        db.session.add_all([
            HeureSupplementaire(salarie_id=sid, date_travail=debut + dt.timedelta(days=4), minutes=-60,
                                motif="retrait", est_ajustement=True, origine_id=None, secteur=secteur),
            HeureSupplementaire(salarie_id=sid, date_travail=debut - dt.timedelta(days=20), minutes=90,
                                motif="avant", secteur=secteur),
            DemandeRecuperation(salarie_id=sid, date_recuperation=debut + dt.timedelta(days=10), minutes=120,
                                statut="acceptee", secteur=secteur),
            FraisKilometrique(salarie_id=sid, date_trajet=debut + dt.timedelta(days=2), annee=debut.year,
                              annee_bareme=debut.year, type_vehicule="voiture", puissance_fiscale=5,
                              distance_km=20, motif="trajet A", montant_centimes=1272, secteur=secteur),
            FraisKilometrique(salarie_id=sid, date_trajet=debut + dt.timedelta(days=12), annee=debut.year,
                              annee_bareme=debut.year, type_vehicule="voiture", puissance_fiscale=5,
                              distance_km=30, motif="trajet B", montant_centimes=1908, secteur=secteur),
        ])
        db.session.commit()
    return client, sid, debut.strftime("%Y-%m")


def test_le_secteur_est_fige_a_la_saisie(app, admin_client):
    secteur = f"Fige{_suf()}"
    client, sid, _, _ = _salarie(app, secteur=secteur)
    client.post("/salarie/recuperations/heures", data={
        "date_travail": dt.date.today().isoformat(), "duree": "1h", "motif": "avant changement"})
    with app.app_context():
        from app.extensions import db
        from app.models import HeureSupplementaire, Salarie
        db.session.get(Salarie, sid).secteur = "Autre secteur"
        db.session.commit()
        assert HeureSupplementaire.query.filter_by(salarie_id=sid).one().secteur == secteur


def test_classeur_excel_chiffres_et_onglets(app, admin_client):
    from openpyxl import load_workbook
    secteur = f"Exp{_suf()}"
    _, sid, mois = _jeu_export(app, secteur)
    r = admin_client.get(f"/salarie/equipe/exports/classeur.xlsx?mois={mois}&secteur={secteur}&par_salarie=1")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["Content-Type"]
    wb = load_workbook(io.BytesIO(r.data))
    assert {"Synthèse", "Par secteur", "Heures sup", "Récupérations", "Frais km"} <= set(wb.sheetnames)
    assert len(wb.sheetnames) == 6, "un onglet pour l'unique salarié du secteur"
    synthese = list(wb["Synthèse"].iter_rows(min_row=4, values_only=True))
    assert len(synthese) == 1
    ligne = synthese[0]
    # Solde début 1,5 h ; +5 h ; retrait 1 h ; récup 2 h ; solde fin 3,5 h ; 50 km ; 31,80 €.
    assert ligne[2:8] == (1.5, 5.0, 1.0, 2.0, 0.0, 3.5)
    assert ligne[8:12] == (2, 50, 31.8, 31.8)
    entetes = " ".join(str(c.value) for ws in wb.worksheets for c in ws[3] if c.value)
    assert "coût" not in entetes.lower() and "taux horaire" not in entetes.lower()
    page = admin_client.get(f"/salarie/equipe/exports?mois={mois}&secteur={secteur}").get_data(as_text=True)
    assert "3h30" in page and "31,80" in page


def test_droits_des_exports(app, admin_client):
    from openpyxl import load_workbook
    secteur = f"Drt{_suf()}"
    client, sid, mois = _jeu_export(app, secteur)
    email_f, _ = _compte(app, f"Compta {_suf()}", "finance")
    finance = _client(app, email_f)
    r = finance.get(f"/salarie/equipe/exports/classeur.xlsx?mois={mois}&secteur={secteur}")
    assert r.status_code == 200
    noms = set(load_workbook(io.BytesIO(r.data)).sheetnames)
    assert "Frais km" in noms and "Heures sup" not in noms, "la comptabilité ne voit que les frais"
    assert client.get("/salarie/equipe/exports").status_code == 403

    assert client.get(f"/salarie/releve/{sid}?mois={mois}").status_code == 200
    collegue, _, _, _ = _salarie(app)
    assert collegue.get(f"/salarie/releve/{sid}?mois={mois}").status_code == 403
    assert collegue.get(f"/salarie/frais-km/etat/{sid}?mois={mois}").status_code == 403
    assert finance.get(f"/salarie/frais-km/etat/{sid}?mois={mois}").status_code == 200
    assert finance.get(f"/salarie/releve/{sid}?mois={mois}").status_code == 403

    releve = admin_client.get(f"/salarie/releve/{sid}?mois={mois}").get_data(as_text=True)
    assert "1h30" in releve and "3h30" in releve  # solde début / fin
    etat = client.get(f"/salarie/frais-km/etat/{sid}?mois={mois}").get_data(as_text=True)
    assert "trajet A" in etat and "31,80" in etat


def test_un_courriel_trop_vieux_n_est_plus_envoye(app, smtp):
    from app.services.courriels_rh import EXPIRATION_JOURS, expedier_en_attente
    with app.app_context():
        from app.extensions import db
        from app.models import CourrielRh
        vieux = CourrielRh(destinataire="vieux@example.org", evenement="recup_soumise", sujet="s", corps="c",
                           created_at=dt.datetime.now() - dt.timedelta(days=EXPIRATION_JOURS + 3))
        frais = CourrielRh(destinataire="frais@example.org", evenement="recup_soumise", sujet="s", corps="c")
        db.session.add_all([vieux, frais])
        db.session.commit()
        expedier_en_attente()
        assert db.session.get(CourrielRh, vieux.id).envoye_le is None
        assert "Expiré" in db.session.get(CourrielRh, vieux.id).derniere_erreur
        assert db.session.get(CourrielRh, frais.id).envoye_le is not None
    assert [m[0] for m in smtp] == ["frais@example.org"]
