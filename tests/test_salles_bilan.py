"""Taux d'occupation et valorisation des contributions en nature.

Un taux d'occupation faux ne plante pas : il part dans un dossier de
subvention et se retourne contre la structure quand quelqu'un le recoupe.
Chaque chiffre est donc testé contre un calcul vérifiable à la main.
"""
from datetime import date, timedelta
from uuid import uuid4

import pytest

from app.extensions import db
from app.models import (
    CategoriePreneur,
    Espace,
    Occupation,
    Preneur,
    Reservation,
    Site,
    TarifSalle,
)
from app.services.bilan_salles import (
    heures_ouvrables,
    occupation_par_salle,
    phrase_pour_dossier,
    synthese,
    valorisation,
)
from app.services.reservations import appliquer_dates, recalculer, reference_unique
from app.services.temps_ouverture import HORAIRES_DEFAUT, ecrire_horaires

# Une semaine complète de lundi à dimanche, sans férié.
LUNDI = date(2026, 10, 5)
DIMANCHE = date(2026, 10, 11)


@pytest.fixture()
def centre(app):
    with app.app_context():
        suffixe = uuid4().hex[:8]
        site = Site(
            nom="Centre (bilan)", code=f"test-bil-{suffixe}",
            horaires_json=ecrire_horaires(HORAIRES_DEFAUT),   # 9h-18h en semaine, 9h-12h30 samedi
            valeur_locative_annuelle=36500.0,                  # 100 € par jour, calcul facile
            regime_sous_location="autorisee",
        )
        db.session.add(site)
        db.session.flush()
        grande = Espace(site=site, nom="Grande salle", louable=True, ordre=10)
        petite = Espace(site=site, nom="Petite salle", louable=True, ordre=20)
        jamais = Espace(site=site, nom="Salle oubliée", louable=True, ordre=30)
        categorie = CategoriePreneur(code=f"c_{suffixe}", libelle="Association")
        db.session.add_all([grande, petite, jamais, categorie])
        db.session.flush()
        for unite, montant in (("heure", 10.0), ("demi_journee", 30.0)):
            db.session.add(TarifSalle(
                espace_id=grande.id, categorie_id=categorie.id, unite=unite,
                montant=montant, date_debut=date(2020, 1, 1),
            ))
        preneur = Preneur(nom="Asso Test", categorie_id=categorie.id,
                          assurance_rc_fin=date(2030, 12, 31))
        db.session.add(preneur)
        db.session.commit()
        ids = {"site": site.id, "grande": grande.id, "petite": petite.id,
               "jamais": jamais.id, "categorie": categorie.id, "preneur": preneur.id}
        site_id = site.id

    yield ids

    with app.app_context():
        for modele in (Reservation, TarifSalle):
            for objet in modele.query.all():
                db.session.delete(objet)
        db.session.commit()
        for modele, cle in ((Site, site_id), (Preneur, ids["preneur"]),
                            (CategoriePreneur, ids["categorie"])):
            objet = db.session.get(modele, cle)
            if objet is not None:
                db.session.delete(objet)
        db.session.commit()


def _occuper(espace_id, jour, debut_h, fin_h, origine="seance", titre="Atelier"):
    occ = Occupation(
        espace_id=espace_id, date_jour=jour,
        minute_debut=debut_h * 60, minute_fin=fin_h * 60,
        origine=origine, statut="confirme", titre=titre,
    )
    db.session.add(occ)
    db.session.commit()
    return occ


# ---------------------------------------------------------------------------
# Le dénominateur : les heures réellement ouvrables
# ---------------------------------------------------------------------------

def test_le_potentiel_suit_les_horaires_douverture(app, centre):
    """Un taux calculé sur 24 h par jour ne veut rien dire et se retourne
    contre la structure."""
    with app.app_context():
        site = db.session.get(Site, centre["site"])
        # 5 jours × 9 h (9h-18h) + samedi 3,5 h (9h-12h30) + dimanche fermé.
        assert heures_ouvrables(site, LUNDI, DIMANCHE) == 48.5


def test_un_ferie_sort_du_potentiel(app, centre):
    with app.app_context():
        site = db.session.get(Site, centre["site"])
        # Le 1er mai 2026 est un vendredi : la semaine perd ses 9 h.
        semaine_avec = heures_ouvrables(site, date(2026, 4, 27), date(2026, 5, 3))
        semaine_sans = heures_ouvrables(site, date(2026, 4, 20), date(2026, 4, 26))
        assert semaine_sans - semaine_avec == 9.0


def test_sans_horaires_le_potentiel_reste_raisonnable(app, centre):
    """8 h par jour ouvré plutôt que 24 : un dénominateur absurde produit
    un taux ridicule qui décrédibilise tout le bilan."""
    with app.app_context():
        site = db.session.get(Site, centre["site"])
        site.horaires_json = None
        db.session.commit()
        assert heures_ouvrables(site, LUNDI, DIMANCHE) == 40.0


# ---------------------------------------------------------------------------
# Le numérateur : ce qui occupe vraiment
# ---------------------------------------------------------------------------

def test_taux_doccupation_dune_salle(app, centre):
    with app.app_context():
        _occuper(centre["grande"], LUNDI, 9, 18)          # 9 h
        _occuper(centre["grande"], LUNDI + timedelta(days=1), 14, 18)  # 4 h
        lignes = occupation_par_salle(centre["site"], LUNDI, DIMANCHE)
        grande = [l for l in lignes if l["espace"].nom == "Grande salle"][0]

        assert grande["heures"] == 13.0
        assert grande["potentiel"] == 48.5
        assert grande["taux"] == 26.8  # 13 / 48,5


def test_une_salle_jamais_occupee_ressort(app, centre):
    """Soit elle sert sans être déclarée, soit il y a un potentiel à
    mobiliser — dans les deux cas il faut le voir."""
    with app.app_context():
        _occuper(centre["grande"], LUNDI, 9, 12)
        donnees = synthese(centre["site"], LUNDI, DIMANCHE)
        oubliees = {l["espace"].nom for l in donnees["salles_jamais_utilisees"]}
        assert "Salle oubliée" in oubliees
        assert "Grande salle" not in oubliees


def test_les_familles_regroupent_les_origines(app, centre):
    """Qui lit un bilan se moque de savoir si l'activité vient du module
    Activités ou a été tapée à la main."""
    with app.app_context():
        _occuper(centre["grande"], LUNDI, 9, 11, origine="seance")
        _occuper(centre["grande"], LUNDI, 11, 12, origine="interne")
        _occuper(centre["grande"], LUNDI, 14, 15, origine="reunion")
        _occuper(centre["grande"], LUNDI, 15, 17, origine="location")
        lignes = occupation_par_salle(centre["site"], LUNDI, DIMANCHE)
        grande = [l for l in lignes if l["espace"].nom == "Grande salle"][0]

        assert grande["familles"]["activites"] == 3.0   # séance + interne
        assert grande["familles"]["equipe"] == 1.0
        assert grande["familles"]["tiers"] == 2.0
        assert grande["heures"] == 6.0


def test_une_indisponibilite_sort_du_potentiel_au_lieu_de_plomber_le_taux(app, centre):
    """Une salle en travaux n'avait pas à être utilisée : la compter comme
    un échec serait injuste et fausserait le pilotage."""
    with app.app_context():
        _occuper(centre["grande"], LUNDI, 9, 18, origine="blocage", titre="Travaux")
        lignes = occupation_par_salle(centre["site"], LUNDI, DIMANCHE)
        grande = [l for l in lignes if l["espace"].nom == "Grande salle"][0]

        assert grande["heures"] == 0.0, "un blocage n'occupe pas"
        assert grande["potentiel"] == 39.5, "48,5 h moins les 9 h de travaux"
        assert grande["indisponible"] == 9.0


def test_une_occupation_annulee_ne_compte_pas(app, centre):
    with app.app_context():
        occ = _occuper(centre["grande"], LUNDI, 9, 18)
        occ.statut = "annule"
        db.session.commit()
        assert synthese(centre["site"], LUNDI, DIMANCHE)["total_heures"] == 0.0


def test_hors_periode_exclu(app, centre):
    with app.app_context():
        _occuper(centre["grande"], LUNDI - timedelta(days=30), 9, 18)
        assert synthese(centre["site"], LUNDI, DIMANCHE)["total_heures"] == 0.0


def test_loccupation_indirecte_ne_compte_pas_deux_fois(app, centre):
    """Réserver une demi-salle rend la grande indisponible mais ne
    l'occupe pas : compter des deux côtés gonflerait le taux global."""
    with app.app_context():
        demi = Espace(site_id=centre["site"], parent_id=centre["grande"],
                      nom="Demi-salle", louable=True, ordre=15)
        db.session.add(demi)
        db.session.commit()
        _occuper(demi.id, LUNDI, 9, 13)

        donnees = synthese(centre["site"], LUNDI, DIMANCHE)
        assert donnees["total_heures"] == 4.0, "4 h une seule fois, pas 8"


# ---------------------------------------------------------------------------
# Valorisation
# ---------------------------------------------------------------------------

def test_la_contribution_recue_est_proratisee(app, centre):
    with app.app_context():
        # 36 500 € par an = 100 € par jour ; 7 jours = 700 €.
        valeurs = valorisation(centre["site"], LUNDI, DIMANCHE)
        assert valeurs["recu_annuel"] == 36500.0
        assert round(valeurs["recu_periode"]) == 700


def test_la_somme_des_espaces_prime_sur_lestimation_globale(app, centre):
    """Plus fine que l'estimation du site, et les deux ne s'additionnent pas."""
    with app.app_context():
        db.session.get(Espace, centre["grande"]).valeur_locative_annuelle = 10000.0
        db.session.get(Espace, centre["petite"]).valeur_locative_annuelle = 5000.0
        db.session.commit()
        valeurs = valorisation(centre["site"], LUNDI, DIMANCHE)
        assert valeurs["recu_annuel"] == 15000.0
        assert valeurs["source_valeur"] == "espaces"


def test_une_gratuite_chiffre_ce_que_la_structure_donne(app, centre):
    with app.app_context():
        r = Reservation(
            reference=reference_unique(LUNDI), preneur_id=centre["preneur"],
            espace_id=centre["grande"], categorie_id=centre["categorie"],
            titre="Gym", statut="confirmee", gratuite=True,
            motif_gratuite="Partenariat",
        )
        db.session.add(r)
        db.session.flush()
        appliquer_dates(r, [LUNDI], "09:00", "13:00")
        db.session.flush()
        recalculer(r)
        db.session.commit()

        valeurs = valorisation(centre["site"], LUNDI, DIMANCHE)
        assert valeurs["donne"] == 30.0, "4 h → forfait demi-journée à 30 €"
        assert valeurs["heures_gratuites"] == 4.0
        assert valeurs["structures_accueillies"] == 1
        assert valeurs["encaisse"] == 0.0


def test_un_rabais_consenti_compte_aussi_comme_contribution(app, centre):
    with app.app_context():
        r = Reservation(
            reference=reference_unique(LUNDI), preneur_id=centre["preneur"],
            espace_id=centre["grande"], categorie_id=centre["categorie"],
            titre="Réunion", statut="confirmee",
            montant_manuel=10.0, motif_montant_manuel="Tarif négocié",
        )
        db.session.add(r)
        db.session.flush()
        appliquer_dates(r, [LUNDI], "09:00", "13:00")
        db.session.flush()
        recalculer(r)
        db.session.commit()

        valeurs = valorisation(centre["site"], LUNDI, DIMANCHE)
        assert valeurs["encaisse"] == 10.0
        assert valeurs["donne"] == 20.0, "30 € au barème, 10 € demandés"


def test_une_reservation_annulee_nest_pas_valorisee(app, centre):
    with app.app_context():
        r = Reservation(
            reference=reference_unique(LUNDI), preneur_id=centre["preneur"],
            espace_id=centre["grande"], categorie_id=centre["categorie"],
            titre="Annulée", statut="annulee", gratuite=True, motif_gratuite="X",
        )
        db.session.add(r)
        db.session.flush()
        appliquer_dates(r, [LUNDI], "09:00", "13:00")
        db.session.flush()
        recalculer(r)
        db.session.commit()
        assert valorisation(centre["site"], LUNDI, DIMANCHE)["donne"] == 0.0


# ---------------------------------------------------------------------------
# La phrase pour un dossier
# ---------------------------------------------------------------------------

def test_la_phrase_est_directement_recopiable(app, centre):
    with app.app_context():
        _occuper(centre["grande"], LUNDI, 9, 18)
        site = db.session.get(Site, centre["site"])
        donnees = synthese(centre["site"], LUNDI, DIMANCHE)
        phrase = phrase_pour_dossier(site, donnees)

        assert "taux d'occupation" in phrase
        assert "05/10/2026" in phrase and "11/10/2026" in phrase
        assert "9 heures" in phrase
        assert phrase.endswith(".")


# ---------------------------------------------------------------------------
# Écrans
# ---------------------------------------------------------------------------

def test_lecran_de_bilan_repond(admin_client, centre):
    reponse = admin_client.get(f"/salles/bilan?site_id={centre['site']}")
    assert reponse.status_code == 200
    page = reponse.get_data(as_text=True)
    assert "Taux d'occupation" in page
    assert "Contributions volontaires en nature" in page


def test_periode_personnalisee(admin_client, centre):
    reponse = admin_client.get(
        f"/salles/bilan?site_id={centre['site']}&debut={LUNDI.isoformat()}&fin={DIMANCHE.isoformat()}"
    )
    assert reponse.status_code == 200
    assert "05/10/2026" in reponse.get_data(as_text=True)


def test_dates_inversees_corrigees(admin_client, centre):
    reponse = admin_client.get(
        f"/salles/bilan?site_id={centre['site']}&debut={DIMANCHE.isoformat()}&fin={LUNDI.isoformat()}"
    )
    assert reponse.status_code == 200


def test_lexport_tableur_est_un_vrai_xlsx(admin_client, centre):
    from io import BytesIO

    from openpyxl import load_workbook

    reponse = admin_client.get(f"/salles/bilan/export.xlsx?site_id={centre['site']}")
    assert reponse.status_code == 200
    classeur = load_workbook(BytesIO(reponse.data))
    assert classeur.sheetnames == ["Occupation", "Valorisation"]
    valeurs = [c.value for ligne in classeur["Valorisation"].iter_rows() for c in ligne]
    assert any("Contributions volontaires" in str(v) for v in valeurs)
    assert any("taux d'occupation" in str(v) for v in valeurs), "la phrase toute faite doit y être"


# ---------------------------------------------------------------------------
# Périmètre de calcul : le dénominateur décide de tout
# ---------------------------------------------------------------------------

def test_les_bureaux_sont_exclus_du_taux_global_par_defaut(app, centre):
    """Un centre social qui annonce 1,5 % d'occupation parce qu'on a divisé
    par huit bureaux se discrédite tout seul."""
    with app.app_context():
        for n in range(8):
            db.session.add(Espace(site_id=centre["site"], nom=f"Bureau {n + 1}",
                                  type_espace="bureau", ordre=100 + n))
        db.session.commit()
        _occuper(centre["grande"], LUNDI, 9, 18)  # 9 h sur 48,5 h

        defaut = synthese(centre["site"], LUNDI, DIMANCHE)
        tous = synthese(centre["site"], LUNDI, DIMANCHE, "tous")

        assert defaut["nb_salles"] == 3, "les trois salles, sans les bureaux"
        assert tous["nb_salles"] == 11
        assert defaut["taux_global"] > tous["taux_global"] * 3, \
            "inclure les bureaux écrase le taux"


def test_perimetre_louables(app, centre):
    with app.app_context():
        db.session.get(Espace, centre["jamais"]).louable = False
        db.session.commit()
        louables = synthese(centre["site"], LUNDI, DIMANCHE, "louables")
        assert {l["espace"].nom for l in louables["lignes"]} == {"Grande salle", "Petite salle"}


def test_perimetre_inconnu_retombe_sur_le_defaut(app, centre):
    with app.app_context():
        assert synthese(centre["site"], LUNDI, DIMANCHE, "n-importe-quoi")["perimetre"] == "activite"


def test_la_phrase_precise_sur_combien_despaces(app, centre):
    with app.app_context():
        _occuper(centre["grande"], LUNDI, 9, 18)
        site = db.session.get(Site, centre["site"])
        phrase = phrase_pour_dossier(site, synthese(centre["site"], LUNDI, DIMANCHE))
        assert "espace(s) d'activité" in phrase
