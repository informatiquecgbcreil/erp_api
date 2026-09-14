"""Lot 3b : contrats, états des lieux, factures, impayés, matériel partagé.

Un document mal formé ne plante pas à la génération : il plante quand la
secrétaire l'ouvre devant le preneur. Les tests vérifient donc que chaque
fichier est un .docx réellement relisible, et que son contenu dit ce qu'il
doit dire — une convention pour une gratuité, un contrat pour une location.
"""
from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document as LireDocx

from app.extensions import db
from app.models import (
    CategoriePreneur,
    Espace,
    Occupation,
    OccupationRessource,
    Preneur,
    Reservation,
    RessourceMobile,
    Site,
    TarifSalle,
)
from app.services import documents_salles as docs
from app.services.reservations import appliquer_dates, recalculer, reference_unique
from app.services.salles import (
    appliquer_ressources,
    disponibilite_ressource,
    quantite_retenue,
)

JEUDI = date(2026, 10, 15)


@pytest.fixture()
def bail(app):
    with app.app_context():
        suffixe = uuid4().hex[:8]
        site = Site(
            nom="Centre (docs)", code=f"test-doc-{suffixe}",
            bailleur_nom="Centre Social de Creil", bailleur_adresse="1 rue des Tests",
            code_postal="60100", ville="Creil", bailleur_siret="12345678900012",
            bailleur_representant="La présidente", bailleur_qualite="Présidente",
            regime_sous_location="autorisee",
        )
        db.session.add(site)
        db.session.flush()
        salle = Espace(site=site, nom="Grande salle", louable=True,
                       capacite_usage=40, capacite_reglementaire=60, pmr=True)
        categorie = CategoriePreneur(code=f"cat_{suffixe}", libelle="Association")
        db.session.add_all([salle, categorie])
        db.session.flush()
        for unite, montant in (("heure", 12.0), ("demi_journee", 35.0)):
            db.session.add(TarifSalle(
                espace_id=salle.id, categorie_id=categorie.id, unite=unite,
                montant=montant, date_debut=date(2020, 1, 1),
            ))
        preneur = Preneur(
            nom="Gym Volontaire", categorie_id=categorie.id,
            assurance_rc_fin=date(2030, 12, 31), representant="Le président",
            adresse="2 rue du Sport", code_postal="60100", ville="Creil",
            siret="98765432100019", telephone="03 44 00 00 00",
        )
        db.session.add(preneur)
        db.session.commit()
        ids = {"site": site.id, "salle": salle.id, "categorie": categorie.id,
               "preneur": preneur.id}
        site_id = site.id

    yield ids

    with app.app_context():
        for modele in (Reservation, RessourceMobile, TarifSalle):
            for objet in modele.query.all():
                db.session.delete(objet)
        db.session.commit()
        for modele, cle in ((Site, site_id), (Preneur, ids["preneur"]),
                            (CategoriePreneur, ids["categorie"])):
            objet = db.session.get(modele, cle)
            if objet is not None:
                db.session.delete(objet)
        db.session.commit()


def _reservation(bail, **kwargs):
    r = Reservation(
        reference=reference_unique(JEUDI), preneur_id=bail["preneur"],
        espace_id=bail["salle"], categorie_id=bail["categorie"],
        titre="Cours de gym", statut="confirmee", effectif=25, **kwargs,
    )
    db.session.add(r)
    db.session.flush()
    appliquer_dates(r, [JEUDI], "09:00", "13:00")
    db.session.flush()
    recalculer(r)
    db.session.commit()
    return r


def _texte(document) -> str:
    """Tout le texte d'un document, tableaux compris."""
    morceaux = [p.text for p in document.paragraphs]
    for table in document.tables:
        for ligne in table.rows:
            morceaux.extend(c.text for c in ligne.cells)
    return "\n".join(morceaux)


def _relire(document):
    """Écrit puis relit le document : c'est le seul moyen de vérifier qu'il
    s'ouvrira vraiment chez le preneur."""
    tampon = BytesIO()
    document.save(tampon)
    tampon.seek(0)
    return LireDocx(tampon)


# ---------------------------------------------------------------------------
# Contrat et convention
# ---------------------------------------------------------------------------

def test_le_contrat_est_un_docx_relisible(app, bail):
    with app.app_context():
        r = _reservation(bail)
        relu = _relire(docs.contrat(r, "Centre Social"))
        assert len(relu.paragraphs) > 20


def test_le_contrat_nomme_les_deux_parties(app, bail):
    with app.app_context():
        r = _reservation(bail)
        texte = _texte(docs.contrat(r, "Centre Social"))
        assert "Centre Social de Creil" in texte
        assert "Gym Volontaire" in texte
        assert "12345678900012" in texte and "98765432100019" in texte
        assert "La présidente" in texte


def test_le_contrat_porte_le_detail_du_prix(app, bail):
    with app.app_context():
        r = _reservation(bail)
        texte = _texte(docs.contrat(r, ""))
        assert "35.00" in texte, "le prix calculé doit figurer"
        assert "293 B" in texte, "la mention de TVA est obligatoire"


def test_une_gratuite_produit_une_convention_pas_un_contrat(app, bail):
    """On ne signe pas la même chose : un « contrat de location » sur une
    mise à disposition sans contrepartie ferait mauvaise figure devant un
    financeur comme devant un contrôle."""
    with app.app_context():
        r = _reservation(bail, gratuite=True, motif_gratuite="Convention de partenariat")
        texte = _texte(docs.contrat(r, ""))
        assert "CONVENTION DE MISE À DISPOSITION" in texte.upper()
        assert "CONTRAT DE MISE" not in texte.upper()
        assert "titre gratuit" in texte
        assert "Convention de partenariat" in texte
        assert "contribution volontaire en nature" in texte
        assert r.type_document == "convention"


def test_le_contrat_alerte_quand_lassurance_manque(app, bail):
    with app.app_context():
        db.session.get(Preneur, bail["preneur"]).assurance_rc_fin = None
        r = _reservation(bail)
        texte = _texte(docs.contrat(r, ""))
        assert "ATTENTION" in texte
        db.session.get(Preneur, bail["preneur"]).assurance_rc_fin = date(2030, 12, 31)
        db.session.commit()


def test_le_prix_impose_est_justifie_dans_le_contrat(app, bail):
    with app.app_context():
        r = _reservation(bail, montant_manuel=20.0, motif_montant_manuel="Tarif négocié")
        texte = _texte(docs.contrat(r, ""))
        assert "dérogation au barème" in texte
        assert "Tarif négocié" in texte


def test_les_conditions_generales_sont_modifiables_sans_developpeur(app, bail):
    with app.app_context():
        site = db.session.get(Site, bail["site"])
        site.conditions_generales = "Article unique — on range les chaises."
        db.session.commit()
        texte = _texte(docs.contrat(_reservation(bail), ""))
        assert "on range les chaises" in texte
        assert "propriété commerciale" not in texte, "les conditions par défaut sont remplacées"


def test_les_conditions_par_defaut_servent_de_point_de_depart(app, bail):
    with app.app_context():
        texte = _texte(docs.contrat(_reservation(bail), ""))
        assert "précaire et révocable" in texte
        assert "responsabilité civile" in texte


# ---------------------------------------------------------------------------
# État des lieux
# ---------------------------------------------------------------------------

def test_letat_des_lieux_dentree_est_un_formulaire_vierge(app, bail):
    with app.app_context():
        texte = _texte(docs.etat_des_lieux(_reservation(bail), sortie=False))
        assert "entrée" in texte.lower()
        assert "Dégradations constatées" in texte
        assert "Sols et revêtements" in texte
        assert "dépôt de garantie" not in texte.lower(), "le sort de la caution est une affaire de sortie"


def test_letat_des_lieux_de_sortie_statue_sur_la_caution(app, bail):
    with app.app_context():
        r = _reservation(bail, caution_montant=300.0)
        texte = _texte(docs.etat_des_lieux(r, sortie=True))
        assert "sortie" in texte.lower()
        assert "Restitution intégrale" in texte
        assert "300.00" in texte


# ---------------------------------------------------------------------------
# Facture et attestation
# ---------------------------------------------------------------------------

def test_la_facture_porte_son_numero_et_le_detail(app, bail):
    with app.app_context():
        r = _reservation(bail)
        texte = _texte(docs.facture(r, "FAC-2026-0001", "Centre Social"))
        assert "FAC-2026-0001" in texte
        assert "TOTAL À RÉGLER" in texte
        assert "35.00" in texte


def test_la_facture_isole_la_caution_du_montant_a_regler(app, bail):
    """Une caution n'est pas une recette : la confondre avec le prix fausse
    la comptabilité et surprend le preneur."""
    with app.app_context():
        r = _reservation(bail, caution_montant=300.0)
        texte = _texte(docs.facture(r, "FAC-2026-0002", ""))
        assert "n'est pas encaissé au titre de la présente facture" in texte


def test_la_facture_deduit_lacompte(app, bail):
    with app.app_context():
        r = _reservation(bail, acompte_montant=10.0, acompte_regle_le=date(2026, 9, 1))
        texte = _texte(docs.facture(r, "FAC-2026-0003", ""))
        assert "Acompte reçu" in texte
        assert "RESTE À RÉGLER" in texte


def test_lattestation_chiffre_les_heures_et_la_valorisation(app, bail):
    """C'est ce que les associations demandent pour leurs propres dossiers
    de subvention."""
    with app.app_context():
        r = _reservation(bail, gratuite=True, motif_gratuite="Partenariat")
        texte = _texte(docs.attestation(r, "Centre Social"))
        assert "4.0 heures" in texte
        assert "titre gratuit" in texte
        assert "35.00" in texte, "la valorisation doit figurer"


# ---------------------------------------------------------------------------
# Écriture sur disque
# ---------------------------------------------------------------------------

def test_lecriture_produit_un_fichier_ouvrable(app, bail, tmp_path):
    with app.app_context():
        r = _reservation(bail)
        chemin_docx, _ = docs.ecrire(
            docs.contrat(r, ""), str(tmp_path), docs.nom_fichier(r, "contrat"), pdf=False,
        )
        assert chemin_docx.endswith(".docx")
        assert LireDocx(chemin_docx) is not None


def test_le_nom_de_fichier_est_parlant_et_sans_surprise(app, bail):
    with app.app_context():
        r = _reservation(bail)
        nom = docs.nom_fichier(r, "contrat")
        assert nom.startswith("contrat_LOC-")
        assert "Gym-Volontaire" in nom
        assert all(c.isalnum() or c in "-_" for c in nom), "rien qui casse un système de fichiers"


# ---------------------------------------------------------------------------
# Matériel partagé : le vidéoprojecteur unique
# ---------------------------------------------------------------------------

def test_un_seul_videoprojecteur_pour_deux_salles(app, bail):
    """LE cas d'Antoine : deux salles libres, un seul appareil."""
    with app.app_context():
        videoproj = RessourceMobile(nom="Vidéoprojecteur", quantite=1)
        db.session.add(videoproj)
        db.session.flush()

        r = _reservation(bail)
        occ = r.occupations[0]
        avertissements = appliquer_ressources(occ, {videoproj.id: 1})
        db.session.commit()
        assert avertissements == []

        etat = disponibilite_ressource(videoproj, JEUDI, "10:00", "12:00")
        assert etat["retenue"] == 1 and etat["restante"] == 0 and etat["epuisee"]


def test_le_materiel_se_libere_sur_un_autre_creneau(app, bail):
    with app.app_context():
        videoproj = RessourceMobile(nom="Vidéoprojecteur", quantite=1)
        db.session.add(videoproj)
        db.session.flush()
        r = _reservation(bail)
        appliquer_ressources(r.occupations[0], {videoproj.id: 1})
        db.session.commit()

        # 09:00-13:00 est pris ; l'après-midi ne l'est pas.
        assert disponibilite_ressource(videoproj, JEUDI, "14:00", "16:00")["restante"] == 1
        assert disponibilite_ressource(videoproj, JEUDI + timedelta(days=1), "09:00", "13:00")["restante"] == 1


def test_depassement_de_stock_prevenu_sans_etre_interdit(app, bail):
    """Il arrive qu'on emprunte un vidéoprojecteur à côté : l'application
    n'a pas à trancher à la place de l'équipe, mais elle doit le dire."""
    with app.app_context():
        videoproj = RessourceMobile(nom="Vidéoprojecteur", quantite=1)
        db.session.add(videoproj)
        db.session.flush()
        r = _reservation(bail)
        avertissements = appliquer_ressources(r.occupations[0], {videoproj.id: 3})
        db.session.commit()

        assert len(avertissements) == 1
        assert "3 demandé" in avertissements[0]
        assert OccupationRessource.query.count() == 1, "la demande est enregistrée malgré tout"


def test_une_occupation_annulee_rend_le_materiel(app, bail):
    with app.app_context():
        videoproj = RessourceMobile(nom="Vidéoprojecteur", quantite=1)
        db.session.add(videoproj)
        db.session.flush()
        r = _reservation(bail)
        appliquer_ressources(r.occupations[0], {videoproj.id: 1})
        db.session.commit()

        r.occupations[0].statut = "annule"
        db.session.commit()
        assert quantite_retenue(videoproj.id, JEUDI, 540, 780) == 0


def test_supprimer_une_occupation_libere_son_materiel(app, bail):
    with app.app_context():
        videoproj = RessourceMobile(nom="Vidéoprojecteur", quantite=2)
        db.session.add(videoproj)
        db.session.flush()
        r = _reservation(bail)
        appliquer_ressources(r.occupations[0], {videoproj.id: 2})
        db.session.commit()
        assert OccupationRessource.query.count() == 1

        db.session.delete(r)
        db.session.commit()
        assert OccupationRessource.query.count() == 0, "pas de ligne orpheline"


# ---------------------------------------------------------------------------
# Impayés et cautions
# ---------------------------------------------------------------------------

def test_une_occupation_passee_non_reglee_est_impayee(app, bail):
    with app.app_context():
        r = _reservation(bail)
        appliquer_dates(r, [date.today() - timedelta(days=10)], "09:00", "13:00")
        db.session.flush()
        recalculer(r)
        db.session.commit()
        assert r.impayee is True


def test_une_occupation_a_venir_nest_pas_impayee(app, bail):
    with app.app_context():
        r = _reservation(bail)
        appliquer_dates(r, [date.today() + timedelta(days=10)], "09:00", "13:00")
        db.session.flush()
        recalculer(r)
        db.session.commit()
        assert r.impayee is False


def test_une_gratuite_nest_jamais_impayee(app, bail):
    with app.app_context():
        r = _reservation(bail, gratuite=True, motif_gratuite="Partenariat")
        appliquer_dates(r, [date.today() - timedelta(days=10)], "09:00", "13:00")
        db.session.flush()
        recalculer(r)
        db.session.commit()
        assert r.impayee is False


def test_un_solde_regle_solde_limpaye(app, bail):
    with app.app_context():
        r = _reservation(bail, solde_regle_le=date.today())
        appliquer_dates(r, [date.today() - timedelta(days=10)], "09:00", "13:00")
        db.session.flush()
        recalculer(r)
        db.session.commit()
        assert r.reste_du == 0.0
        assert r.impayee is False


def test_une_caution_encaissee_apres_coup_doit_etre_rendue(app, bail):
    with app.app_context():
        r = _reservation(bail, caution_montant=300.0, caution_encaissee=True)
        appliquer_dates(r, [date.today() - timedelta(days=5)], "09:00", "13:00")
        db.session.commit()
        assert r.caution_a_restituer is True

        r.caution_restituee_le = date.today()
        db.session.commit()
        assert r.caution_a_restituer is False


# ---------------------------------------------------------------------------
# Écrans et téléchargements
# ---------------------------------------------------------------------------

def test_les_ecrans_du_lot3b_repondent(admin_client, bail):
    for url in ("/salles/impayes", "/salles/ressources"):
        assert admin_client.get(url).status_code == 200, url


def test_les_documents_se_telechargent(admin_client, app, bail):
    with app.app_context():
        r = _reservation(bail)
        rid = r.id

    for genre in ("contrat", "etat_lieux_entree", "etat_lieux_sortie", "facture", "attestation"):
        reponse = admin_client.get(f"/salles/reservation/{rid}/document/{genre}")
        assert reponse.status_code == 200, genre
        assert len(reponse.data) > 8000, f"{genre} : fichier suspicieusement petit"
        # Un .docx est une archive ZIP : les deux premiers octets le disent.
        if reponse.data[:2] == b"PK":
            assert LireDocx(BytesIO(reponse.data)) is not None, genre


def test_la_facture_prend_un_numero_a_la_premiere_edition(admin_client, app, bail):
    with app.app_context():
        r = _reservation(bail)
        rid = r.id
        assert r.facture_numero is None

    admin_client.get(f"/salles/reservation/{rid}/document/facture")
    with app.app_context():
        r = db.session.get(Reservation, rid)
        premier = r.facture_numero
        assert premier and premier.startswith("FAC-")

    # Rééditer ne doit pas consommer un second numéro.
    admin_client.get(f"/salles/reservation/{rid}/document/facture")
    with app.app_context():
        assert db.session.get(Reservation, rid).facture_numero == premier


def test_une_gratuite_ne_se_facture_pas(admin_client, app, bail):
    with app.app_context():
        r = _reservation(bail, gratuite=True, motif_gratuite="Partenariat")
        rid = r.id
    reponse = admin_client.get(
        f"/salles/reservation/{rid}/document/facture", follow_redirects=True,
    )
    assert "ne se facture pas" in reponse.get_data(as_text=True)


def test_le_contrat_refuse_de_sortir_si_un_controle_bloque(admin_client, app, bail):
    with app.app_context():
        db.session.get(Preneur, bail["preneur"]).assurance_rc_fin = None
        r = _reservation(bail)
        rid = r.id
        db.session.commit()

    reponse = admin_client.get(
        f"/salles/reservation/{rid}/document/contrat", follow_redirects=True,
    )
    assert "non édité" in reponse.get_data(as_text=True)

    with app.app_context():
        db.session.get(Preneur, bail["preneur"]).assurance_rc_fin = date(2030, 12, 31)
        db.session.commit()


def test_genre_de_document_inconnu_renvoie_404(admin_client, app, bail):
    with app.app_context():
        rid = _reservation(bail).id
    assert admin_client.get(f"/salles/reservation/{rid}/document/nimporte-quoi").status_code == 404


def test_une_prestation_facturee_retient_son_materiel(app, admin_client, bail):
    """Le vidéoprojecteur se facture ET n'existe qu'en un exemplaire.
    Le déclarer deux fois serait absurde : retenir la prestation doit
    retenir le matériel, sinon on facture un appareil déjà promis."""
    from app.models import PrestationSalle

    with app.app_context():
        videoproj = RessourceMobile(nom="Vidéoprojecteur", quantite=1)
        db.session.add(videoproj)
        db.session.flush()
        prestation = PrestationSalle(
            libelle="Vidéoprojecteur", montant=15.0, unite="forfait",
            ressource_id=videoproj.id,
        )
        db.session.add(prestation)
        db.session.commit()
        pid, vid = prestation.id, videoproj.id

    admin_client.post(
        "/salles/reservation/nouvelle",
        data={
            "preneur_id": str(bail["preneur"]), "espace_id": str(bail["salle"]),
            "titre": "Conférence", "date_debut": JEUDI.isoformat(),
            "heure_debut": "09:00", "heure_fin": "13:00",
            f"prestation_{pid}": "1",
        },
        follow_redirects=True,
    )

    with app.app_context():
        r = Reservation.query.filter_by(titre="Conférence").one()
        assert r.montant_calcule == 50.0, "35 € de salle + 15 € de prestation"
        # Et le matériel est bien immobilisé sur la date.
        assert OccupationRessource.query.filter_by(ressource_id=vid).count() == 1
        assert disponibilite_ressource(
            db.session.get(RessourceMobile, vid), JEUDI, "10:00", "12:00"
        )["restante"] == 0


def test_le_materiel_se_retient_depuis_lecran_doccupation(admin_client, app, bail):
    with app.app_context():
        videoproj = RessourceMobile(nom="Sono", quantite=2)
        db.session.add(videoproj)
        db.session.commit()
        vid = videoproj.id

    admin_client.post(
        "/salles/occuper",
        data={
            "espace_id": str(bail["salle"]), "titre": "Fête de quartier",
            "origine": "interne", "date_debut": JEUDI.isoformat(),
            "heure_debut": "14:00", "heure_fin": "18:00",
            f"ressource_{vid}": "2",
        },
        follow_redirects=True,
    )

    with app.app_context():
        occ = Occupation.query.filter_by(titre="Fête de quartier").one()
        assert len(occ.ressources) == 1 and occ.ressources[0].quantite == 2
        assert disponibilite_ressource(
            db.session.get(RessourceMobile, vid), JEUDI, "15:00", "17:00"
        )["epuisee"] is True
