"""Territoire du centre : la fin du territoire de Creil codé en dur.

Ce que ces tests garantissent :

1. **Creil ne bouge pas, sauf ses erreurs.** Avec le territoire historique
   (celui que la migration reprend), le bilan SENACS classe chaque habitant
   comme l'ancien code — recopié ici tel qu'il était — sur un jeu de
   quartiers choisis pour être piégeux, à quelques corrections près, listées
   une à une (« Hauts de Creil (hors Rouher) » n'est pas le Rouher).
   Le Magatomatique suit les mêmes règles.
2. **Un autre centre obtient ses propres lignes** : zéro, un ou plusieurs
   QPV suivis, des quartiers distingués ou non, ses « Habitants de … ».
3. **Une base neuve n'hérite pas de Creil**, et une base de Creil le garde.
4. Les écrans (Administration > Identité, assistant de premier démarrage)
   enregistrent le territoire sans écraser ce qu'ils ne montrent pas.
"""
import uuid
from datetime import date

import pytest

from app.services.territoire import TERRITOIRE_HISTORIQUE, Territoire, lire_liste

CREIL = Territoire(
    ville_siege="Creil",
    code_postal="60100",
    qpv_suivis=("Hauts de Creil",),
    quartiers_distingues=("Rouher",),
    latitude=49.2583,
    longitude=2.4750,
)


def _suffixe() -> str:
    return uuid.uuid4().hex[:6]


# ---------------------------------------------------------------------------
# Lecture des réglages
# ---------------------------------------------------------------------------

def test_une_valeur_par_ligne_sans_doublon_ni_vide():
    assert lire_liste("Saint-Jean\n\n  Argentine \nsaint-jean;Notre-Dame") == (
        "Saint-Jean", "Argentine", "Notre-Dame")
    assert lire_liste(None) == ()
    assert lire_liste("   ") == ()


def test_la_ville_siege_se_reconnait_quelle_que_soit_l_ecriture():
    nogent = Territoire(ville_siege="Nogent-sur-Oise")
    assert nogent.habite_ville_siege("NOGENT SUR OISE")
    assert nogent.habite_ville_siege("nogent-sur-oise (60)")
    assert not nogent.habite_ville_siege("Creil")
    assert not nogent.habite_ville_siege(None)
    # Sans ville siège réglée, personne n'y habite.
    assert not Territoire().habite_ville_siege("Creil")
    assert Territoire().libelle_ville == "Ville siège"


# ---------------------------------------------------------------------------
# 1. Creil : identique à l'ancien calcul
# ---------------------------------------------------------------------------

def _ancien_bucket(p, villes_qpv):
    """Le classement SENACS tel qu'il était codé avant le réglage du territoire."""
    from app.services.recherche_texte import sans_accent

    quartier_distingue, qpv_de_reference = "Rouher", "Hauts de Creil"
    quartier = p.quartier
    if quartier is None:
        ville = (sans_accent(p.ville) or "").strip()
        if not ville:
            return "Non renseigné"
        return "Non renseigné" if ville in villes_qpv else "Hors QPV"
    nom = sans_accent(quartier.nom) or ""
    qpv = (getattr(quartier, "qpv", None) or "").strip()
    if sans_accent(quartier_distingue) in nom:
        return quartier_distingue
    if qpv:
        return qpv_de_reference if sans_accent(qpv) == sans_accent(qpv_de_reference) else "Autre QPV"
    if sans_accent(qpv_de_reference) in nom or "hauts-de-creil" in nom:
        return qpv_de_reference
    if getattr(quartier, "is_qpv", False) or "qpv" in nom:
        return "Autre QPV"
    if ((quartier.ville or "") or (p.ville or "")).strip():
        return "Hors QPV"
    return "Non renseigné"


def _ancien_is_qpv(p):
    if not p.quartier:
        return False
    if (p.quartier.qpv or "").strip() or p.quartier.is_qpv:
        return True
    ville = (p.quartier.ville or "").strip().lower()
    nom = (p.quartier.nom or "").strip().lower()
    return "qpv" in nom or (ville == "creil" and ("rouher" in nom or "hauts de creil" in nom))


#: (ville, nom, qpv, is_qpv) — des cas réels et des cas tordus.
QUARTIERS_PIEGES = [
    ("Creil", "Rouher", "Hauts de Creil", True),
    ("Creil", "Rouher", None, False),
    ("Creil", "Rouher (QPV Hauts de Creil)", None, False),
    ("Creil", "Hauts de Creil (hors Rouher)", "Hauts de Creil", True),
    ("Creil", "Hauts de Creil (hors Rouher)", None, True),
    ("Creil", "Cavée de Senlis", "Hauts de Creil", True),
    ("Creil", "Cavée de Senlis", None, False),
    ("Creil", "Hauts de Creil", None, False),
    ("Creil", "hauts-de-creil nord", None, False),
    ("Creil", "Bas de Creil", None, False),
    ("Creil", "Plateau", "HAUTS DE CRÉIL", True),
    ("Creil", "Jaurès Gournay", "Jaurès Gournay", True),
    ("Creil", "Jaurès Gournay", None, True),
    ("Creil", "Quartier QPV sans nom", None, False),
    ("Nogent-sur-Oise", "Les Granges", "Les Granges", True),
    ("Nogent-sur-Oise", "Centre", None, False),
    ("Montataire", "Rouher-bis", None, False),
    ("", "Sans ville", None, False),
]

#: Les seuls écarts avec l'ancien bilan SENACS : des erreurs corrigées.
#: L'ancien code cherchait « rouher » n'importe où dans le nom du quartier.
CORRECTIONS_SENACS = {
    # La ligne Rouher comptait tous les Hauts de Creil.
    ("Hauts de Creil (hors Rouher)", "Hauts de Creil"): "Hauts de Creil",
    ("Hauts de Creil (hors Rouher)", None): "Hauts de Creil",
    # Un quartier d'une autre ville n'est pas le Rouher.
    ("Rouher-bis", None): "Hors QPV",
}

VILLES_A_QPV = {"creil", "nogent-sur-oise", "montataire"}


def _en_memoire(ville, nom=None, qpv=None, is_qpv=False):
    """Un habitant et son quartier, sans base : seuls leurs attributs comptent."""
    from types import SimpleNamespace

    quartier = SimpleNamespace(ville=ville, nom=nom, qpv=qpv, is_qpv=is_qpv) if nom else None
    return SimpleNamespace(quartier=quartier, ville=ville or None)


def test_creil_garde_ses_chiffres_sauf_les_erreurs_corrigees(app, territoire_creil):
    from app.models import Participant
    from app.services.senacs import _bucket_quartier
    from app.services.territoire import territoire

    habitants = [_en_memoire(*cas) for cas in QUARTIERS_PIEGES]
    habitants += [_en_memoire(ville) for ville in ("Creil", "CREIL", "Nogent-sur-Oise", "Laigneville", "")]
    with app.app_context():
        assert territoire() == CREIL
        for p in habitants:
            q = p.quartier
            cle = (q.nom, q.qpv) if q is not None else None
            attendu = CORRECTIONS_SENACS.get(cle) or _ancien_bucket(p, VILLES_A_QPV)
            assert _bucket_quartier(p, VILLES_A_QPV) == attendu, (p.ville, q and q.nom)

            is_qpv = Participant.is_qpv.fget(p)
            if q is not None and q.nom == "hauts-de-creil nord":
                # Écrit avec des tirets et sans QPV renseigné : déjà compté aux
                # Hauts de Creil par SENACS, désormais aussi dans le taux QPV.
                assert _ancien_is_qpv(p) is False and is_qpv is True
                continue
            assert is_qpv == _ancien_is_qpv(p), (p.ville, q and q.nom)


# ---------------------------------------------------------------------------
# 1 bis. Le Magatomatique, sur les quartiers réels du centre
# ---------------------------------------------------------------------------

#: Les quartiers de Creil et de l'agglomération tels que le centre les tient,
#: et la colonne attendue. « Haut Montataire » est un nom inventé : l'ancien
#: tri l'aurait rangé aux Hauts de Creil (le mot « haut » dans le nom).
PROVENANCES_ATTENDUES = [
    (("Creil", "Rouher", "Hauts de Creil", True), "Rouher"),
    (("Creil", "Hauts de Creil (hors Rouher)", "Hauts de Creil", True), "Hauts de Creil"),
    (("Creil", "Bas de Creil", None, False), "Creil hors QPV"),
    (("Creil", "Jaurès Gournay", "Jaurès Gournay", True), "Autre QPV"),
    (("Creil", "Jaurès Gournay", None, True), "Autre QPV"),
    (("Nogent-sur-Oise", "Les Granges", "Les Granges", True), "Autre QPV"),
    (("Nogent-sur-Oise", "Centre", None, False), "Hors Creil"),
    (("Montataire", "Haut Montataire", None, False), "Hors Creil"),
    (("Villers-Saint-Paul",), "Hors Creil"),   # commune sans QPV, pas de quartier
    (("Nogent-sur-Oise",), "Inconnu"),         # commune à QPV, quartier non dit
    (("",), "Inconnu"),
]


def test_les_quartiers_du_centre_dans_le_magatomatique(app, territoire_creil):
    from app.services.senacs import provenance
    from app.services.territoire import territoire

    with app.app_context():
        terr = territoire()
        assert terr.colonnes_provenance == [
            "Rouher", "Hauts de Creil", "Autre QPV", "Creil hors QPV", "Hors Creil", "Inconnu"]
        for cas, attendu in PROVENANCES_ATTENDUES:
            assert provenance(_en_memoire(*cas), VILLES_A_QPV, terr) == attendu, cas


def test_un_qpv_suivi_a_part_a_sa_colonne():
    """Jaurès Gournay ajouté aux QPV suivis : sa colonne à lui. Un quartier
    est rattaché par le QPV inscrit sur sa fiche, ou à défaut par son nom ;
    la case « QPV » seule ne dit pas lequel."""
    from app.services.senacs import provenance

    terr = Territoire(ville_siege="Creil", qpv_suivis=("Hauts de Creil", "Jaurès Gournay"),
                      quartiers_distingues=("Rouher",))
    assert terr.colonnes_provenance == [
        "Rouher", "Hauts de Creil", "Jaurès Gournay", "Autre QPV", "Creil hors QPV", "Hors Creil", "Inconnu"]
    assert provenance(_en_memoire("Creil", "Jaurès Gournay", "Jaurès Gournay", True), VILLES_A_QPV, terr) \
        == "Jaurès Gournay"
    assert provenance(_en_memoire("Creil", "Jaurès Gournay", None, True), VILLES_A_QPV, terr) == "Jaurès Gournay"
    assert provenance(_en_memoire("Creil", "Rue Jean Jaurès", None, True), VILLES_A_QPV, terr) == "Autre QPV"


def test_colonnes_d_un_centre_sans_territoire_regle():
    assert Territoire().colonnes_provenance == ["QPV", "Hors QPV", "Inconnu"]
    beauvais = Territoire(ville_siege="Beauvais", qpv_suivis=("Saint-Jean",))
    assert beauvais.colonnes_provenance == [
        "Saint-Jean", "Autre QPV", "Beauvais hors QPV", "Hors Beauvais", "Inconnu"]


def test_l_export_magatomatique_suit_le_territoire(app, admin_client, territoire_creil):
    """De bout en bout : le classeur Excel porte les colonnes du territoire,
    remplies, et son tableau de synthèse couvre toutes les colonnes."""
    import datetime as dt
    from io import BytesIO

    from openpyxl import load_workbook

    from app.extensions import db
    from app.models import AtelierActivite, PresenceActivite, SessionActivite

    suf = _suffixe()
    secteur = f"Prov{suf}"
    with app.app_context():
        habitants = [
            _habitant(app, "Creil", "Rouher", qpv="Hauts de Creil", is_qpv=True),
            _habitant(app, "Creil", "Hauts de Creil (hors Rouher)", qpv="Hauts de Creil", is_qpv=True),
            _habitant(app, "Creil", "Bas de Creil"),
            _habitant(app, "Creil", "Jaurès Gournay", qpv="Jaurès Gournay", is_qpv=True),
            _habitant(app, f"Villers{suf}"),
        ]
        atelier = AtelierActivite(nom=f"Atelier{suf}", secteur=secteur, type_atelier="COLLECTIF",
                                  capacite_defaut=10, is_active=True)
        db.session.add(atelier)
        db.session.flush()
        seance = SessionActivite(atelier_id=atelier.id, secteur=secteur, session_type="COLLECTIF",
                                 date_session=dt.date(dt.date.today().year, 1, 2), heure_debut="14:00",
                                 heure_fin="16:00", capacite=10, statut="realisee")
        db.session.add(seance)
        db.session.flush()
        db.session.add_all([PresenceActivite(session_id=seance.id, participant_id=p.id) for p in habitants])
        db.session.commit()

    r = admin_client.get(f"/stats-impact/magatomatique.xlsx?export_mode=complete&secteur={secteur}")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.data))["Synthese"]
    entetes = [c.value for c in ws[3]]
    ligne = next(dict(zip(entetes, row)) for row in ws.iter_rows(min_row=4, values_only=True)
                 if row[1] == f"Atelier{suf}")
    assert {k: ligne[k] for k in ("Rouher", "Hauts de Creil", "Autre QPV", "Creil hors QPV",
                                  "Hors Creil", "Inconnu")} == {
        "Rouher": 1, "Hauts de Creil": 1, "Autre QPV": 1, "Creil hors QPV": 1, "Hors Creil": 1, "Inconnu": 0}
    assert "Bas de Creil" not in entetes and "Autres" not in entetes
    # Le tableau Excel de la synthèse s'étend jusqu'à la dernière colonne (genre).
    assert ws.tables["SyntheseAteliers"].ref.split(":")[1].rstrip("0123456789") == \
        ws.cell(row=3, column=len([e for e in entetes if e])).column_letter


# ---------------------------------------------------------------------------
# 2. Un autre centre : ses QPV, ses quartiers, sa ville
# ---------------------------------------------------------------------------

def _habitant(app, ville, nom_quartier=None, qpv=None, is_qpv=False):
    from app.extensions import db
    from app.models import Participant, Quartier

    suf = _suffixe()
    quartier_id = None
    if nom_quartier:
        # Le suffixe (unicité en base) va entre parenthèses : un quartier
        # distingué se reconnaît à son nom, précisions entre parenthèses à part.
        quartier = Quartier(ville=ville, nom=f"{nom_quartier} ({suf})", qpv=qpv, is_qpv=is_qpv)
        db.session.add(quartier)
        db.session.flush()
        quartier_id = quartier.id
    p = Participant(nom=f"Hab{suf}", prenom="X", ville=ville, quartier_id=quartier_id)
    db.session.add(p)
    db.session.flush()
    return p


@pytest.mark.parametrize("qpv_suivis, attendus", [
    # Un centre rural ou sans QPV suivi : tous les QPV sur une seule ligne.
    ((), {"Argentine": "QPV", "Saint-Jean": "QPV", "Notre-Dame": "QPV"}),
    # Un QPV suivi : les autres sont « Autre QPV ».
    (("Saint-Jean",), {"Argentine": "Autre QPV", "Saint-Jean": "Saint-Jean", "Notre-Dame": "Autre QPV"}),
    # Trois QPV suivis : chacun sa ligne.
    (("Saint-Jean", "Argentine", "Notre-Dame"),
     {"Argentine": "Argentine", "Saint-Jean": "Saint-Jean", "Notre-Dame": "Notre-Dame"}),
])
def test_zero_un_ou_plusieurs_qpv_suivis(app, qpv_suivis, attendus):
    from app.services.senacs import _bucket_quartier

    beauvais = Territoire(ville_siege="Beauvais", qpv_suivis=qpv_suivis)
    with app.app_context():
        for qpv, attendu in attendus.items():
            p = _habitant(app, "Beauvais", f"Quartier {qpv}", qpv=qpv, is_qpv=True)
            assert _bucket_quartier(p, set(), beauvais) == attendu
        # Hors QPV et quartier inconnu ne dépendent pas des QPV suivis.
        assert _bucket_quartier(_habitant(app, "Beauvais", "Centre-ville"), set(), beauvais) == "Hors QPV"
        assert _bucket_quartier(_habitant(app, ""), set(), beauvais) == "Non renseigné"
        from app.extensions import db
        db.session.rollback()


def test_rien_de_creil_dans_le_bilan_d_un_autre_centre(app):
    """Un quartier nommé « Rouher » ou « Hauts de Creil » n'a plus de
    ligne à lui chez un centre qui ne l'a pas demandé."""
    from app.services.senacs import _bucket_quartier

    beauvais = Territoire(ville_siege="Beauvais", qpv_suivis=("Saint-Jean",))
    with app.app_context():
        assert _bucket_quartier(_habitant(app, "Beauvais", "Rouher", is_qpv=True), set(), beauvais) == "Autre QPV"
        assert _bucket_quartier(_habitant(app, "Beauvais", "Hauts de Creil"), set(), beauvais) == "Hors QPV"
        from app.extensions import db
        db.session.rollback()


def test_un_quartier_distingue_garde_sa_ligne_avant_son_qpv(app):
    from app.services.senacs import _bucket_quartier

    beauvais = Territoire(ville_siege="Beauvais", qpv_suivis=("Saint-Jean",),
                          quartiers_distingues=("Les Champs Dolent",))
    with app.app_context():
        distingue = _habitant(app, "Beauvais", "Les Champs Dolent", qpv="Saint-Jean", is_qpv=True)
        voisin = _habitant(app, "Beauvais", "Rue de Calais", qpv="Saint-Jean", is_qpv=True)
        assert _bucket_quartier(distingue, set(), beauvais) == "Les Champs Dolent"
        assert _bucket_quartier(voisin, set(), beauvais) == "Saint-Jean"
        from app.extensions import db
        db.session.rollback()


ANNEE_BILAN = 1997  # aucune autre donnée de test : le bilan ne compte que les nôtres


def test_bilan_senacs_et_indicateurs_d_un_centre_de_beauvais(app, regler_territoire):
    """De bout en bout : régler le territoire change le bilan et les indicateurs."""
    from app.extensions import db
    from app.models import AtelierActivite, PresenceActivite, SessionActivite
    from app.services.indicators import _add_demography_metrics
    from app.services.senacs import publics_annee

    regler_territoire(territoire_ville="Beauvais", territoire_qpv_suivis="Saint-Jean\nArgentine",
                     territoire_quartiers_distingues=None)
    with app.app_context():
        habitants = [
            _habitant(app, "Beauvais", "Saint-Jean", qpv="Saint-Jean", is_qpv=True),
            _habitant(app, "BEAUVAIS", "Argentine", qpv="Argentine", is_qpv=True),
            _habitant(app, "Beauvais", "Centre-ville"),
            _habitant(app, "Allonne"),
            _habitant(app, "Creil", "Rouher", qpv="Hauts de Creil", is_qpv=True),
        ]
        atelier = AtelierActivite(nom=f"Atelier Beauvais {_suffixe()}", secteur="Numérique")
        db.session.add(atelier)
        db.session.flush()
        seance = SessionActivite(atelier_id=atelier.id, secteur="Numérique",
                                 date_session=date(ANNEE_BILAN, 3, 10), duree_minutes=60)
        db.session.add(seance)
        db.session.flush()
        db.session.add_all([PresenceActivite(session_id=seance.id, participant_id=p.id) for p in habitants])
        db.session.commit()

        quartiers = publics_annee(ANNEE_BILAN)["quartiers"]
        assert quartiers == {"Argentine": 1, "Autre QPV": 1, "Hors QPV": 2, "Saint-Jean": 1}

        mesures = {}
        _add_demography_metrics(mesures, [p.id for p in habitants], date(ANNEE_BILAN, 12, 31))
        assert mesures["participants_creil"] == 3  # le code reste, il compte Beauvais
        assert mesures["participants_qpv"] == 3


def test_les_habitants_de_creil_restent_comptes_a_creil(app, territoire_creil):
    from app.services.indicators import _add_demography_metrics

    with app.app_context():
        creillois = _habitant(app, "Creil")
        voisin = _habitant(app, "Nogent-sur-Oise")
        mesures = {}
        _add_demography_metrics(mesures, [creillois.id, voisin.id])
        assert mesures["participants_creil"] == 1
        from app.extensions import db
        db.session.rollback()


# ---------------------------------------------------------------------------
# 3. Carte et attestations
# ---------------------------------------------------------------------------

def test_la_carte_s_ouvre_sur_la_ville_siege_sinon_sur_la_france(app, territoire_creil, regler_territoire):
    from app.services.cartographie import CENTRE_FRANCE, centre_par_defaut

    with app.app_context():
        assert centre_par_defaut() == {"lat": 49.2583, "lon": 2.4750, "zoom": 13}
    regler_territoire(territoire_ville="Beauvais")
    with app.app_context():
        assert centre_par_defaut() == CENTRE_FRANCE
    regler_territoire(territoire_ville="Beauvais", territoire_latitude=49.43, territoire_longitude=2.08)
    with app.app_context():
        assert centre_par_defaut() == {"lat": 49.43, "lon": 2.08, "zoom": 13}


def test_la_page_carte_porte_le_centre_regle(admin_client, regler_territoire):
    regler_territoire(territoire_ville="Beauvais", territoire_latitude=49.43, territoire_longitude=2.08)
    page = admin_client.get("/quartiers/carte").get_data(as_text=True)
    assert 'data-centre-lat="49.43"' in page
    assert "49.2583" not in page


def test_les_attestations_sont_faites_a_la_ville_siege(app, territoire_creil, regler_territoire):
    from app.activite.helpers import _attestation_document_fields

    with app.test_request_context("/"):
        assert _attestation_document_fields()["city"] == "Creil"
    regler_territoire(territoire_ville="Beauvais")
    with app.test_request_context("/"):
        assert _attestation_document_fields()["city"] == "Beauvais"
    with app.test_request_context("/?city=Allonne"):
        assert _attestation_document_fields()["city"] == "Allonne"


# ---------------------------------------------------------------------------
# 4. Reprise des bases existantes
# ---------------------------------------------------------------------------

def test_une_base_neuve_n_herite_pas_de_creil(fresh_app):
    from app.services.territoire import territoire

    with fresh_app.app_context():
        # (L'assistant de premier démarrage a pu y régler une ville.)
        terr = territoire()
        assert terr.ville_siege != "Creil"
        assert terr.qpv_suivis == () and terr.quartiers_distingues == ()


def test_la_reprise_pose_creil_sur_une_base_de_creil_sans_rien_ecraser(app, regler_territoire):
    from app.extensions import db
    from app.models import InstanceSettings
    from app.services.territoire import reprendre_territoire_historique

    with app.app_context():
        _habitant(app, "Creil")
        db.session.commit()

    # Déjà réglée (sur Beauvais) : la reprise ne touche à rien.
    regler_territoire(territoire_ville="Beauvais")
    with app.app_context():
        assert reprendre_territoire_historique(db.session.connection()) is False
        db.session.commit()
        assert InstanceSettings.query.first().territoire_ville == "Beauvais"

    # Pas réglée, mais des habitants de Creil : Creil, comme avant.
    regler_territoire()
    with app.app_context():
        assert reprendre_territoire_historique(db.session.connection()) is True
        db.session.commit()
        reglages = InstanceSettings.query.first()
        for colonne, valeur in TERRITOIRE_HISTORIQUE.items():
            assert getattr(reglages, colonne) == valeur
        # Rejouable.
        assert reprendre_territoire_historique(db.session.connection()) is False


# ---------------------------------------------------------------------------
# 5. Les écrans
# ---------------------------------------------------------------------------

def test_administration_enregistre_le_territoire(admin_client, app, regler_territoire, monkeypatch):
    from app.models import InstanceSettings
    from app.services import territoire as service

    appels = []

    def localiser(ville, code_postal):
        appels.append((ville, code_postal))
        return {"lat": 49.43, "lon": 2.08}

    monkeypatch.setattr(service, "localiser_ville_siege", localiser)
    regler_territoire()
    r = admin_client.post("/admin/instance", data={
        "app_name": "", "organization_name": "",
        "territoire_ville": "beauvais", "territoire_code_postal": "60 000",
        "territoire_qpv_suivis": "Saint-Jean\r\nArgentine\r\n",
        "territoire_quartiers_distingues": "",
    })
    assert r.status_code == 302
    assert appels == [("Beauvais", "60000")]
    with app.app_context():
        reglages = InstanceSettings.query.first()
        assert reglages.territoire_ville == "Beauvais"
        assert reglages.territoire_code_postal == "60000"
        assert reglages.territoire_qpv_suivis == "Saint-Jean\nArgentine"
        assert reglages.territoire_quartiers_distingues is None
        assert (reglages.territoire_latitude, reglages.territoire_longitude) == (49.43, 2.08)

    # Réafficher la page : un QPV par ligne dans le champ.
    page = admin_client.get("/admin/instance").get_data(as_text=True)
    assert "Saint-Jean\nArgentine</textarea>" in page

    # Réenregistrer sans changer de ville ne relance pas la recherche.
    admin_client.post("/admin/instance", data={
        "territoire_ville": "Beauvais", "territoire_code_postal": "60000",
        "territoire_qpv_suivis": "Saint-Jean", "territoire_quartiers_distingues": "",
    })
    assert len(appels) == 1


def test_sans_internet_la_ville_est_gardee_et_on_previent(admin_client, app, regler_territoire, monkeypatch):
    from app.models import InstanceSettings
    from app.services import territoire as service

    monkeypatch.setattr(service, "localiser_ville_siege", lambda ville, cp: None)
    regler_territoire()
    r = admin_client.post("/admin/instance", data={"territoire_ville": "Beauvais"}, follow_redirects=True)
    assert "pas pu être placée sur la carte" in r.get_data(as_text=True)
    with app.app_context():
        reglages = InstanceSettings.query.first()
        assert reglages.territoire_ville == "Beauvais"
        assert reglages.territoire_latitude is None


def test_un_formulaire_sans_les_champs_territoire_n_efface_rien(admin_client, app, territoire_creil):
    """Un autre formulaire de la même table (ou un vieil onglet ouvert)
    ne doit pas remettre le territoire à zéro."""
    from app.models import InstanceSettings

    admin_client.post("/admin/instance", data={"app_name": "", "organization_name": ""})
    with app.app_context():
        reglages = InstanceSettings.query.first()
        for colonne, valeur in TERRITOIRE_HISTORIQUE.items():
            assert getattr(reglages, colonne) == valeur
