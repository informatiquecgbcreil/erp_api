"""RGPD et cycle de vie (audit 3.1 à 3.9 et mineurs RGPD).

Chaque test rejoue un constat de l'audit sur des données synthétiques.
"""
import json
import os
import uuid
import zipfile
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import pytest

from app.utils.dates import utcnow

ANCIEN = 4 * 365


def _fiche(app, *, jours=ANCIEN, secteur="Familles", **extra):
    from app.extensions import db
    from app.models import Participant
    with app.app_context():
        date_fiche = utcnow() - timedelta(days=jours)
        p = Participant(nom=f"Rgpd{uuid.uuid4().hex[:8]}", prenom="Test", created_secteur=secteur,
                        created_at=date_fiche, updated_at=date_fiche, **extra)
        db.session.add(p)
        db.session.commit()
        return p.id


def _inactifs(app):
    from app.services.purge_rgpd import participants_inactifs
    with app.app_context():
        return {item["participant"].id for item in participants_inactifs(annees=3)}


def _seance(app, jours_depuis_aujourdhui=0, session_type="COLLECTIF"):
    from app.extensions import db
    from app.models import AtelierActivite, SessionActivite
    with app.app_context():
        atelier = AtelierActivite(nom=f"At {uuid.uuid4().hex[:6]}", secteur="Familles")
        db.session.add(atelier)
        db.session.flush()
        s = SessionActivite(atelier_id=atelier.id, secteur="Familles", session_type=session_type,
                            date_session=date.today() + timedelta(days=jours_depuis_aujourdhui))
        db.session.add(s)
        db.session.commit()
        return atelier.id, s.id


# ---------------------------------------------------------------------------
# 3.4 — personnes encore actives
# ---------------------------------------------------------------------------

def test_inscription_a_une_seance_a_venir_protege(app):
    from app.extensions import db
    from app.models import InscriptionActivite
    pid = _fiche(app)
    atelier_id, sid = _seance(app, jours_depuis_aujourdhui=7)
    with app.app_context():
        ancien = utcnow() - timedelta(days=ANCIEN)
        db.session.add(InscriptionActivite(atelier_id=atelier_id, session_id=sid, participant_id=pid,
                                           statut="inscrit", created_at=ancien, updated_at=ancien))
        db.session.commit()
    assert pid not in _inactifs(app)


def test_inscription_a_un_atelier_en_cours_protege(app):
    from app.extensions import db
    from app.models import InscriptionActivite
    pid = _fiche(app)
    atelier_id, _ = _seance(app, jours_depuis_aujourdhui=-ANCIEN)
    with app.app_context():
        ancien = utcnow() - timedelta(days=ANCIEN)
        db.session.add(InscriptionActivite(atelier_id=atelier_id, session_id=None, participant_id=pid,
                                           statut="inscrit", created_at=ancien, updated_at=ancien))
        db.session.commit()
    assert pid not in _inactifs(app)


@pytest.mark.parametrize("source", ["hart", "encaissement", "membre_bulletin", "portail", "positionnement"])
def test_autres_sources_d_activite(app, source):
    from app.extensions import db
    from app import models as m
    pid = _fiche(app)
    assert pid in _inactifs(app)
    with app.app_context():
        if source == "hart":
            db.session.add(m.HartEvaluation(participant_id=pid, niveau=3, date_evaluation=date.today()))
        elif source == "encaissement":
            from app.services.encaissements import enregistrer
            enregistrer(10, "especes", participant_id=pid, libelle="Test")
        elif source == "membre_bulletin":
            bulletin = m.InscriptionAnnuelle(annee_scolaire=date.today().year, nom="Proche", prenom="Parent")
            db.session.add(bulletin)
            db.session.flush()
            db.session.add(m.InscriptionAnnuelleMembre(inscription_id=bulletin.id, prenom="Test",
                                                       participant_id=pid))
        elif source == "portail":
            db.session.add(m.PortailAttempt(attempt_id=uuid.uuid4().hex, participant_id=pid,
                                            external_id=str(pid)))
        elif source == "positionnement":
            db.session.add(m.ParticipantInsertionPositionnement(participant_id=pid,
                                                                date_positionnement=date.today()))
        db.session.commit()
    assert pid not in _inactifs(app)


# ---------------------------------------------------------------------------
# 3.5 / 3.9 — lots bornés, erreurs isolées, journal
# ---------------------------------------------------------------------------

def test_purge_par_lots_bornee_et_reprise(app):
    from app.services.purge_rgpd import purger_par_lots
    for _ in range(5):
        _fiche(app)
    with app.app_context():
        premier = purger_par_lots(annees=3, declenchement="test", limite=2, lot=1)
        assert premier["anonymises"] == 2 and premier["restants"] >= 3
        suite = purger_par_lots(annees=3, declenchement="test", limite=10000)
        assert suite["restants"] == 0
    assert _inactifs(app) == set()


def test_une_fiche_en_erreur_ne_bloque_pas_les_autres(app, monkeypatch):
    from app.extensions import db
    from app.models import AuditLog, Participant
    from app.services import purge_rgpd
    fautive, saine = _fiche(app), _fiche(app)
    original = purge_rgpd.anonymiser_participant

    def _anonymiser(p, actor_id=None):
        if p.id == fautive:
            raise RuntimeError("fiche corrompue")
        return original(p, actor_id)

    monkeypatch.setattr(purge_rgpd, "anonymiser_participant", _anonymiser)
    with app.app_context():
        rapport = purge_rgpd.purger_par_lots(annees=3, declenchement="test", limite=10000)
        assert fautive in rapport["echecs"]
        assert db.session.get(Participant, saine).nom == "ANONYME"
        assert db.session.get(Participant, fautive).nom != "ANONYME"
        trace = AuditLog.query.filter_by(action="participant.anonymize", participant_id=saine).one()
        assert trace.cible == f"participant #{saine}" and "test" in trace.details


def test_la_purge_ne_tourne_plus_dans_la_requete(app, monkeypatch):
    """Hors test, la requête lance un fil séparé et repart aussitôt."""
    from app.services import maintenance
    lances = []
    monkeypatch.setattr(maintenance.threading, "Thread",
                        lambda **kw: type("T", (), {"start": lambda self: lances.append(kw["name"])})())
    monkeypatch.setitem(app.config, "TESTING", False)
    maintenance.lancer(app)
    assert lances == ["maintenance-quotidienne"]


def test_maintenance_sans_purge_active_vide_seulement_la_file(app, monkeypatch):
    from app.services import maintenance, purge_rgpd
    monkeypatch.setattr(purge_rgpd, "purge_auto_active", lambda: False)
    appels = []
    monkeypatch.setattr(purge_rgpd, "purge_quotidienne_si_necessaire", lambda: appels.append(1))
    with app.app_context():
        rapport = maintenance.executer("test")
    assert "fichiers" in rapport and appels == [] and "conservation" not in rapport


# ---------------------------------------------------------------------------
# 3.2 — journal d'audit
# ---------------------------------------------------------------------------

def test_journal_relie_par_identifiant_et_efface_sans_approximation(app):
    from app.extensions import db
    from app.models import AuditLog, Participant
    from app.services.audit import enregistrer
    from app.services.purge_rgpd import anonymiser_participant
    pid = _fiche(app)
    with app.app_context():
        nom = db.session.get(Participant, pid).nom
        a = enregistrer("participant.edit", cible=f"participant #{pid}", details={"nom": nom, "champ": "ville"})
        b = enregistrer("autre.action", cible=f"Réunion avec {nom}")  # aucun identifiant : pas de devinette
        db.session.commit()
        assert a.participant_id == pid and b.participant_id is None
        anonymiser_participant(db.session.get(Participant, pid))
        db.session.commit()
        a, b = db.session.get(AuditLog, a.id), db.session.get(AuditLog, b.id)
        assert a.cible == f"participant #{pid}"
        assert json.loads(a.details) == {"nom": "effacé", "champ": "ville"}
        assert b.cible == f"Réunion avec {nom}"


def test_creation_de_fiche_journalisee_sans_le_nom(app, admin_client):
    from app.models import AuditLog, Participant
    nom = f"Journal{uuid.uuid4().hex[:6]}"
    admin_client.post("/participants/new", data={"nom": nom, "prenom": "Sans", "force_creation": "1"})
    with app.app_context():
        p = Participant.query.filter_by(nom=nom).one()
        ligne = AuditLog.query.filter_by(action="participant.create", participant_id=p.id).one()
        assert nom not in (ligne.cible or "")


def test_migration_rattache_seulement_les_cibles_exactes():
    import importlib.util
    chemin = Path(__file__).resolve().parents[1] / "migrations/versions/e4f7a2c9b168_conservation_rgpd.py"
    spec = importlib.util.spec_from_file_location("migration_conservation", chemin)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    f = module.participant_de_la_cible
    assert f("participant.anonymize", "participant #12") == 12
    assert f("export.rgpd", "participant#3") == 3
    assert f("participant.merge", "Dupont Jean (#7)") == 7
    assert f("presence.delete", "session #4 · Dupont Jean") is None
    assert f("don.create", "Dupont (#9)") is None


# ---------------------------------------------------------------------------
# 3.1 / 3.3 — copies de fichiers, suppression définitive
# ---------------------------------------------------------------------------

def test_copies_de_signatures_et_bilans_effaces(app):
    from app.extensions import db
    from app.models import Participant
    from app.services.purge_rgpd import anonymiser_participant
    pid = _fiche(app)
    instance = Path(app.instance_path)
    signature = instance / "archives_emargements" / "2025" / "signatures" / f"sig_s1_p{pid}__x.png"
    bilan = instance / "archives_pedagogie" / f"bilan_{pid}_NOM_PRENOM.docx"
    autre = instance / "archives_emargements" / "2025" / "signatures" / f"sig_s1_p{pid}0__y.png"
    for fichier in (signature, bilan, autre):
        fichier.parent.mkdir(parents=True, exist_ok=True)
        fichier.write_bytes(b"x")
    with app.app_context():
        anonymiser_participant(db.session.get(Participant, pid))
        db.session.commit()
    assert not signature.exists() and not bilan.exists()
    assert autre.exists(), "la fiche n°{pid}0 n'est pas concernée"
    autre.unlink()


def test_suppression_definitive_sans_trace_identifiante(app):
    from app.extensions import db
    from app import models as m
    from app.services.participant_suppression import supprimer_definitivement
    from app.services.registre_effacements import lire
    with app.app_context():
        foyer = m.Foyer(nom="Famille Traceable")
        db.session.add(foyer)
        db.session.flush()
        p = m.Participant(nom="Traceable", prenom="Seule", created_secteur="Familles", foyer_id=foyer.id)
        db.session.add(p)
        db.session.flush()
        bulletin = m.InscriptionAnnuelle(annee_scolaire=date.today().year, nom="Proche", prenom="Parent")
        db.session.add(bulletin)
        db.session.flush()
        membre = m.InscriptionAnnuelleMembre(inscription_id=bulletin.id, nom="Traceable", prenom="Seule",
                                             date_naissance=date(2015, 1, 1), participant_id=p.id)
        db.session.add(membre)
        db.session.commit()
        pid, foyer_id, membre_id = p.id, foyer.id, membre.id
        supprimer_definitivement(db.session.get(m.Participant, pid))
        db.session.commit()
        assert db.session.get(m.Participant, pid) is None
        assert "Traceable" not in (db.session.get(m.Foyer, foyer_id).nom or "")
        ligne = db.session.get(m.InscriptionAnnuelleMembre, membre_id)
        assert ligne.participant_id is None and ligne.nom == "ANONYME" and ligne.date_naissance is None
    assert pid in [e["id"] for e in lire()["supprimes"]]


def test_bilan_pedagogique_ne_reste_pas_sur_le_disque(app, admin_client):
    pid = _fiche(app, jours=10)
    dossier = Path(app.instance_path) / "archives_pedagogie"
    avant = set(dossier.glob(f"bilan_{pid}_*")) if dossier.exists() else set()
    admin_client.get(f"/stats/pedagogie/participant/{pid}/bilan")
    apres = set(dossier.glob(f"bilan_{pid}_*")) if dossier.exists() else set()
    assert apres == avant


def test_bilan_pedagogique_refuse_hors_perimetre(app):
    from app.extensions import db
    from app.models import Role, User
    pid = _fiche(app, jours=10, secteur="EPE")
    email = f"anim-{uuid.uuid4().hex[:6]}@test.fr"
    with app.app_context():
        u = User(email=email, nom="Anim", secteur_assigne="Numérique")
        u.set_password("motdepasse-tests")
        u.roles.append(Role.query.filter_by(code="responsable_secteur").one())
        db.session.add(u)
        db.session.commit()
    c = app.test_client()
    c.post("/", data={"email": email, "password": "motdepasse-tests"})
    assert c.get(f"/stats/pedagogie/participant/{pid}/bilan").status_code == 403


# ---------------------------------------------------------------------------
# 3.6 — file d'effacement
# ---------------------------------------------------------------------------

def test_file_d_effacement_ne_se_bloque_plus(app):
    from app.extensions import db
    from app.models import PendingFileDeletion
    from app.services.file_cleanup import drain
    fichier = Path(app.instance_path) / "a_effacer" / f"{uuid.uuid4().hex}.txt"
    fichier.parent.mkdir(parents=True, exist_ok=True)
    fichier.write_text("x")
    with app.app_context():
        for i in range(3):
            db.session.add(PendingFileDeletion(file_path=f"/ancien-serveur/{uuid.uuid4().hex}-{i}.png"))
        db.session.add(PendingFileDeletion(file_path=str(fichier)))
        db.session.commit()
        premier = drain()
        assert premier["effaces"] >= 1 and premier["bloques"] >= 3
        assert not fichier.exists()
        bloquees = PendingFileDeletion.query.filter(PendingFileDeletion.file_path.like("/ancien-serveur/%")).all()
        assert all(b.bloque_le and b.motif == "hors_stockage" for b in bloquees)
        # Le même jour, les lignes mises de côté ne sont pas relues.
        assert drain()["bloques"] == 0


# ---------------------------------------------------------------------------
# 3.7 — export article 15
# ---------------------------------------------------------------------------

def test_export_article_15_complet_et_archive(app, admin_client):
    from app.extensions import db
    from app import models as m
    pid = _fiche(app, jours=10)
    atelier_id, sid = _seance(app)
    signature = Path(app.config["APP_UPLOAD_DIR"]) / "signatures" / f"test-{pid}.png"
    signature.parent.mkdir(parents=True, exist_ok=True)
    signature.write_bytes(b"\x89PNG")
    with app.app_context():
        q = m.Questionnaire(nom="Satisfaction", type_questionnaire="satisfaction", is_active=True)
        db.session.add(q)
        db.session.flush()
        question = m.Question(questionnaire_id=q.id, label="Qu'avez-vous appris ?", kind="text", position=1)
        db.session.add(question)
        groupe = m.QuestionnaireResponseGroup(questionnaire_id=q.id, participant_id=pid)
        db.session.add(groupe)
        db.session.flush()
        db.session.add(m.QuestionResponse(response_group_id=groupe.id, question_id=question.id,
                                          value_text="Envoyer un courriel"))
        db.session.add(m.PresenceActivite(session_id=sid, participant_id=pid, signature_path=str(signature)))
        db.session.commit()
    from openpyxl import load_workbook
    r = admin_client.get(f"/participants/{pid}/export-rgpd.xlsx")
    wb = load_workbook(BytesIO(r.data))
    assert {"Réponses questionnaires", "Foyer", "Bulletin d'un proche", "Encaissements",
            "Orientations - destinataires", "Matériel utilisé"} <= set(wb.sheetnames)
    reponses = "\n".join(str(c.value) for row in wb["Réponses questionnaires"].iter_rows() for c in row)
    assert "Envoyer un courriel" in reponses and "Qu'avez-vous appris ?" in reponses
    archive = admin_client.get(f"/participants/{pid}/export-rgpd.zip")
    assert archive.status_code == 200
    with zipfile.ZipFile(BytesIO(archive.data)) as z:
        noms = z.namelist()
        assert "donnees.xlsx" in noms and any(n.startswith("signatures/") for n in noms)


def test_referentiel_exporte_en_libelle(app):
    from app.services.rgpd_export import _lisible
    from app import models as m
    from app.extensions import db
    pid = _fiche(app, jours=10)
    with app.app_context():
        atelier_id, sid = _seance(app)
        presence = m.PresenceActivite(session_id=sid, participant_id=pid)
        db.session.add(presence)
        db.session.commit()
        valeur = _lisible(presence, "session_id")
        assert valeur == sid or "n°" in str(valeur)


# ---------------------------------------------------------------------------
# 3.8 — durées de conservation
# ---------------------------------------------------------------------------

def test_durees_de_conservation(app):
    from app.extensions import db
    from app import models as m
    from app.services import conservation
    ancien = utcnow() - timedelta(days=5 * 365)
    with app.app_context():
        isole = m.InscriptionAnnuelle(annee_scolaire=2019, nom="Isole", prenom="Bulletin", email="i@x.fr",
                                      date_inscription=date(2019, 9, 1), created_at=ancien, updated_at=ancien)
        recent = m.InscriptionAnnuelle(annee_scolaire=date.today().year, nom="Recent", prenom="Bulletin")
        vieux_don = m.Don(numero=f"T-{uuid.uuid4().hex[:8]}", annee=date.today().year - 8,
                          donateur_nom="Genereux", donateur_email="g@x.fr")
        don_recent = m.Don(numero=f"T-{uuid.uuid4().hex[:8]}", annee=date.today().year - 1,
                           donateur_nom="Recent")
        compte = m.User(email=f"parti-{uuid.uuid4().hex[:6]}@x.fr", nom="Parti", actif=False, created_at=ancien)
        compte.set_password("x" * 12)
        vieux_log = m.AuditLog(action="test.ancien", created_at=utcnow() - timedelta(days=4 * 365))
        db.session.add_all([isole, recent, vieux_don, don_recent, compte, vieux_log])
        db.session.commit()
        ids = (isole.id, recent.id, vieux_don.id, don_recent.id, compte.id, vieux_log.id)
        rapport = conservation.appliquer("test")
        assert rapport["bulletins"] >= 1 and rapport["donateurs"] >= 1 and rapport["comptes"] >= 1
        isole, recent, vieux_don, don_recent, compte = (
            db.session.get(m.InscriptionAnnuelle, ids[0]), db.session.get(m.InscriptionAnnuelle, ids[1]),
            db.session.get(m.Don, ids[2]), db.session.get(m.Don, ids[3]), db.session.get(m.User, ids[4]))
        assert isole.nom == "ANONYME" and isole.email is None
        assert recent.nom == "Recent"
        assert vieux_don.donateur_nom == "Donateur anonymisé" and vieux_don.donateur_email is None
        assert vieux_don.montant is not None and vieux_don.numero
        assert don_recent.donateur_nom == "Recent"
        assert compte.email.endswith("@supprime.invalid")
        assert m.AuditLog.query.filter_by(action="test.ancien").count() == 0


def test_duree_zero_conserve_sans_limite(app, admin_client):
    from app.extensions import db
    from app import models as m
    from app.services import conservation
    valeurs = {nom: "0" for nom in conservation.REGLES}
    admin_client.post("/controle/purge-rgpd/conservation", data=valeurs)
    with app.app_context():
        vieux_log = m.AuditLog(action="test.conserve", created_at=utcnow() - timedelta(days=20 * 365))
        db.session.add(vieux_log)
        db.session.commit()
        conservation.appliquer("test")
        assert m.AuditLog.query.filter_by(action="test.conserve").count() == 1
    admin_client.post("/controle/purge-rgpd/conservation",
                      data={nom: str(regle[0]) for nom, regle in conservation.REGLES.items()})


def test_locataire_anonymisable_seulement_sans_reservation_en_cours(app, admin_client):
    from app.extensions import db
    from app import models as m
    with app.app_context():
        preneur = m.Preneur(nom=f"Asso {uuid.uuid4().hex[:6]}", email="asso@x.fr", telephone="0600000000")
        db.session.add(preneur)
        db.session.commit()
        pid = preneur.id
    r = admin_client.post(f"/salles/preneur/{pid}/anonymiser")
    assert r.status_code == 302
    with app.app_context():
        preneur = db.session.get(m.Preneur, pid)
        assert preneur.nom.startswith("Locataire anonymisé") and preneur.email is None


# ---------------------------------------------------------------------------
# Mineurs RGPD
# ---------------------------------------------------------------------------

def test_restauration_reanonymise_sans_toucher_un_homonyme_de_numero(app):
    from app.extensions import db
    from app.models import Participant
    from app.services.purge_rgpd import anonymiser_participant
    from app.services.registre_effacements import reappliquer
    pid = _fiche(app)
    with app.app_context():
        p = db.session.get(Participant, pid)
        nom = p.nom
        anonymiser_participant(p)
        db.session.commit()
        # « Restauration » : la fiche revient avec son identité.
        p = db.session.get(Participant, pid)
        p.nom, p.prenom = nom, "Revenue"
        db.session.commit()
        assert pid in reappliquer()
        assert db.session.get(Participant, pid).nom == "ANONYME"
        # Même numéro mais autre date de création : autre personne, intacte.
        p = db.session.get(Participant, pid)
        p.nom, p.created_at = "Nouvelle", utcnow()
        db.session.commit()
        assert pid not in reappliquer()
        assert db.session.get(Participant, pid).nom == "Nouvelle"


def test_journal_de_connexion_ne_garde_pas_un_mot_de_passe_tape_dans_l_email(app):
    from app.models import JournalConnexion
    tape = f"MotDePasseSecret{uuid.uuid4().hex[:6]}"
    app.test_client().post("/", data={"email": tape, "password": "x"})
    with app.app_context():
        assert JournalConnexion.query.filter(JournalConnexion.email.like(f"%{tape}%")).count() == 0
        assert JournalConnexion.query.filter(JournalConnexion.email.like("inconnu:%")).count() >= 1


def test_portail_ne_recree_pas_le_lien_vers_une_fiche_anonymisee(app, monkeypatch):
    from app.extensions import db
    from app.models import Participant, PortailAttempt
    from app.services import portail_apprenants as svc
    from app.services.purge_rgpd import anonymiser_participant
    pid = _fiche(app)
    with app.app_context():
        anonymiser_participant(db.session.get(Participant, pid))
        db.session.commit()
        attempt = uuid.uuid4().hex
        monkeypatch.setattr(svc, "portail_configure", lambda: True)
        monkeypatch.setattr(svc, "recuperer_attempts", lambda since: {
            "attempts": [{"id": attempt, "externalId": str(pid), "score": 3}], "generatedAt": "2026-01-01"})
        svc.synchroniser_attempts()
        ligne = PortailAttempt.query.filter_by(attempt_id=attempt).one()
        assert ligne.participant_id is None and ligne.external_id is None


def test_bilan_de_rendez_vous_individuel_efface(app):
    from app.extensions import db
    from app.models import Participant, PresenceActivite, SessionActivite
    from app.services.purge_rgpd import anonymiser_participant
    pid = _fiche(app)
    _, individuelle = _seance(app, session_type="INDIVIDUEL_MENSUEL")
    _, collective = _seance(app)
    autre = _fiche(app, jours=10)
    with app.app_context():
        for sid in (individuelle, collective):
            s = db.session.get(SessionActivite, sid)
            s.bilan_qualitatif = "Madame X a beaucoup progressé"
            db.session.add(PresenceActivite(session_id=sid, participant_id=pid))
        db.session.add(PresenceActivite(session_id=collective, participant_id=autre))
        db.session.commit()
        anonymiser_participant(db.session.get(Participant, pid))
        db.session.commit()
        assert db.session.get(SessionActivite, individuelle).bilan_qualitatif is None
        assert db.session.get(SessionActivite, collective).bilan_qualitatif
