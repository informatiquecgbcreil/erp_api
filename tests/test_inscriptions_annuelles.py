"""Module Inscriptions annuelles : saisie, fiche participant, règlement, export.

Le cycle complet d'un bulletin de rentrée :
- saisie du bulletin (coordonnées, ateliers cochés + champ libre, bénévolat
  avec créneaux jour × demi-journée et « je ne sais pas ») ;
- transformation en fiche participant portant le statut d'attente, avec
  inscription automatique aux ateliers choisis ;
- bascule AUTOMATIQUE du statut à la première présence pointée, quelle que
  soit la porte d'entrée (le garde-fou est posé sur la session SQLAlchemy) ;
- confirmation du règlement, avec création de l'adhésion dans le module
  Adhésions & participation ;
- rattachement à une fiche existante plutôt que création d'un doublon ;
- fiche imprimable, export XLSX, cloisonnement par permission.
"""
import uuid
from datetime import date
from io import BytesIO

import pytest

ANNEE = 2031  # campagne « bac à sable », loin des autres jeux de tests


def _suffixe():
    return uuid.uuid4().hex[:6]


@pytest.fixture()
def atelier(app):
    """Un atelier actif, cochable sur le bulletin."""
    suf = _suffixe()
    with app.app_context():
        from app.extensions import db
        from app.models import AtelierActivite

        a = AtelierActivite(nom=f"Atelier numérique {suf}", secteur="Adultes", capacite_defaut=12)
        db.session.add(a)
        db.session.commit()
        return {"id": a.id, "nom": a.nom}


def _payload_bulletin(suf: str, atelier_id: int, **extra):
    data = {
        "annee": ANNEE,
        "nom": f"Rentree{suf}",
        "prenom": "Soline",
        "adresse": "12 rue des Lilas",
        "code_postal": "60100",
        "ville": "Creil",
        "email": f"soline.{suf}@example.org",
        "telephone": "0344000000",
        "secteur_orienteur": "Adultes",
        "ateliers": [str(atelier_id)],
        "ateliers_libre": "Aimerait aussi de la couture le jeudi",
        "commentaire": "Vient sur conseil d'une voisine.",
        "benevolat_souhaite": "1",
        "benevolat_mission": "Accompagnement scolaire, aide aux devoirs",
        "creneaux": ["mardi|matin", "jeudi|apres_midi"],
        "benevolat_dispo_inconnue": "1",
    }
    data.update(extra)
    return data


def _creer_bulletin(admin_client, app, atelier_id, **extra):
    suf = _suffixe()
    r = admin_client.post(
        "/inscriptions-annuelles/nouvelle",
        data=_payload_bulletin(suf, atelier_id, **extra),
        follow_redirects=False,
    )
    assert r.status_code == 302, r.data[:400]
    with app.app_context():
        from app.models import InscriptionAnnuelle

        inscription = (
            InscriptionAnnuelle.query
            .filter(InscriptionAnnuelle.nom == f"Rentree{suf}")
            .order_by(InscriptionAnnuelle.id.desc())
            .first()
        )
        assert inscription is not None
        return inscription.id, suf


# ---------------------------------------------------------------------------
# Saisie du bulletin
# ---------------------------------------------------------------------------

def test_saisie_bulletin_enregistre_tout_le_formulaire(admin_client, app, atelier):
    """Toutes les données récoltées atterrissent en base, créneaux compris."""
    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])

    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.annee_scolaire == ANNEE
        assert i.libelle_annee == f"{ANNEE}-{ANNEE + 1}"
        assert i.prenom == "Soline"
        assert i.adresse_complete == "12 rue des Lilas, 60100 Creil"
        assert i.email == f"soline.{suf}@example.org"
        assert i.secteur_orienteur == "Adultes"
        assert i.statut == "saisie"
        assert i.reglement_confirme is False

        # Ateliers : cases cochées + champ libre
        assert [a.id for a in i.ateliers] == [atelier["id"]]
        assert "couture" in i.ateliers_libre

        # Bénévolat : mission libre, créneaux cochés, « je ne sais pas »
        assert i.benevolat_souhaite is True
        assert "devoirs" in i.benevolat_mission
        assert i.benevolat_dispo_inconnue is True
        assert i.creneaux_benevolat == [("mardi", "matin"), ("jeudi", "apres_midi")]
        assert i.a_creneau("mardi", "matin") is True
        assert i.a_creneau("mardi", "apres_midi") is False


def test_formulaires_saffichent(admin_client, app, atelier):
    """Le formulaire se rend aussi bien vide qu'avec un bulletin existant."""
    page = admin_client.get(f"/inscriptions-annuelles/nouvelle?annee={ANNEE}").data.decode("utf-8")
    assert atelier["nom"] in page                    # ateliers groupés par secteur
    assert "Je ne sais pas encore" in page           # grille bénévolat dépliable
    assert f"{ANNEE}-{ANNEE + 1}" in page

    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])
    page = admin_client.get(f"/inscriptions-annuelles/{inscription_id}/modifier").data.decode("utf-8")
    assert f"Rentree{suf}" in page                   # champs pré-remplis
    assert "couture" in page
    assert 'name="creneaux" value="mardi|matin" checked' in page


def test_nom_obligatoire_et_email_verifie(admin_client, app, atelier):
    """Deux garde-fous de saisie : identité minimale et e-mail plausible."""
    r = admin_client.post(
        "/inscriptions-annuelles/nouvelle",
        data={"annee": ANNEE, "nom": "", "prenom": "Sans nom"},
        follow_redirects=True,
    )
    assert r.status_code == 200
    assert "obligatoires" in r.data.decode("utf-8")

    suf = _suffixe()
    r = admin_client.post(
        "/inscriptions-annuelles/nouvelle",
        data=_payload_bulletin(suf, atelier["id"], email="pas-un-email"),
        follow_redirects=True,
    )
    assert "invalide" in r.data.decode("utf-8")
    with app.app_context():
        from app.models import InscriptionAnnuelle

        assert InscriptionAnnuelle.query.filter_by(nom=f"Rentree{suf}").first() is None


def test_retirer_lenvie_de_benevolat_efface_les_creneaux(admin_client, app, atelier):
    """Une envie retirée ne doit pas laisser la personne dans la grille."""
    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])

    donnees = _payload_bulletin(suf, atelier["id"])
    donnees.pop("benevolat_souhaite")
    donnees.pop("benevolat_dispo_inconnue")
    donnees["nom"] = f"Rentree{suf}"
    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/modifier", data=donnees, follow_redirects=False
    )
    assert r.status_code == 302

    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.benevolat_souhaite is False
        assert i.disponibilites == []
        assert i.benevolat_mission is None
        assert i.benevolat_dispo_inconnue is False


# ---------------------------------------------------------------------------
# Transformation en fiche participant
# ---------------------------------------------------------------------------

def test_creation_fiche_participant_avec_statut_dattente(admin_client, app, atelier):
    """La fiche est créée « en attente de 1re participation » et la personne
    est inscrite aux ateliers cochés (module Activité)."""
    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])

    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/creer-participant",
        data={"inscrire_ateliers": "1"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app.app_context():
        from app.extensions import db
        from app.models import (
            STATUT_PARTICIPANT_ATTENTE,
            InscriptionActivite,
            InscriptionAnnuelle,
        )

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.statut == "en_attente"
        assert i.participant_id is not None

        p = i.participant
        assert p.statut_inscription == STATUT_PARTICIPANT_ATTENTE
        assert p.attend_premiere_participation is True
        # Les coordonnées du bulletin sont recopiées sur la fiche.
        assert p.ville == "Creil"
        assert p.email == f"soline.{suf}@example.org"
        assert p.created_secteur == "Adultes"
        # Une envie de bénévolat n'est pas un engagement : le drapeau reste à off.
        assert p.est_benevole is False

        # L'atelier coché est devenu une inscription d'activité.
        lignes = InscriptionActivite.query.filter_by(participant_id=p.id).all()
        assert [l.atelier_id for l in lignes] == [atelier["id"]]
        assert lignes[0].statut == "inscrit"


def test_deuxieme_creation_refusee(admin_client, app, atelier):
    """On ne crée pas deux fiches pour le même bulletin."""
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])
    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/creer-participant", data={})
    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/creer-participant", data={}, follow_redirects=True
    )
    assert "déjà une fiche participant" in r.data.decode("utf-8")


def test_rattachement_a_une_fiche_existante(admin_client, app, atelier):
    """Un ancien inscrit se rattache à sa fiche : pas de doublon créé."""
    suf = _suffixe()
    with app.app_context():
        from app.extensions import db
        from app.models import Participant

        ancien = Participant(nom=f"Rentree{suf}", prenom="Soline", telephone="0344999999")
        db.session.add(ancien)
        db.session.commit()
        ancien_id = ancien.id

    r = admin_client.post(
        "/inscriptions-annuelles/nouvelle",
        data=_payload_bulletin(suf, atelier["id"]),
        follow_redirects=False,
    )
    assert r.status_code == 302
    with app.app_context():
        from app.models import InscriptionAnnuelle

        inscription_id = InscriptionAnnuelle.query.filter_by(nom=f"Rentree{suf}").first().id

    # La fiche ressemblante est proposée sur la page détail.
    page = admin_client.get(f"/inscriptions-annuelles/{inscription_id}").data.decode("utf-8")
    assert "fiches ressemblantes" in page

    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/rattacher",
        data={"participant_id": ancien_id, "inscrire_ateliers": "1"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app.app_context():
        from app.extensions import db
        from app.models import STATUT_PARTICIPANT_ATTENTE, InscriptionAnnuelle, Participant

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.participant_id == ancien_id
        assert i.statut == "en_attente"

        p = db.session.get(Participant, ancien_id)
        assert p.statut_inscription == STATUT_PARTICIPANT_ATTENTE
        # Rattachement : on comble les trous sans écraser l'existant.
        assert p.telephone == "0344999999"
        assert p.ville == "Creil"
        assert Participant.query.filter_by(nom=f"Rentree{suf}").count() == 1


# ---------------------------------------------------------------------------
# Première participation : la bascule automatique
# ---------------------------------------------------------------------------

def test_premiere_presence_fait_tomber_le_statut_dattente(admin_client, app, atelier):
    """Pointer la personne présente suffit : aucune action manuelle."""
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])
    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/creer-participant", data={})

    with app.app_context():
        from app.extensions import db
        from app.models import (
            STATUT_PARTICIPANT_ACTIF,
            InscriptionAnnuelle,
            PresenceActivite,
            SessionActivite,
        )

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        participant_id = i.participant_id

        seance = SessionActivite(
            atelier_id=atelier["id"], secteur="Adultes", session_type="COLLECTIF",
            date_session=date(ANNEE, 10, 3), heure_debut="14:00", heure_fin="16:00",
            statut="realisee",
        )
        db.session.add(seance)
        db.session.flush()

        # Émargement brut : le garde-fou est posé sur la session SQLAlchemy,
        # donc il s'applique sans que le module ait été appelé.
        db.session.add(PresenceActivite(session_id=seance.id, participant_id=participant_id))
        db.session.commit()

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.statut == "active"
        assert i.premiere_participation_le == date(ANNEE, 10, 3)
        assert i.participant.statut_inscription == STATUT_PARTICIPANT_ACTIF
        assert i.participant.attend_premiere_participation is False


def test_rafraichir_statuts_rattrape_les_presences_hors_application(app, atelier):
    """Filet de sécurité pour les reprises de données : un bulletin resté en
    attente alors que la personne est déjà venue se recale."""
    with app.app_context():
        from app.extensions import db
        from app.models import (
            STATUT_PARTICIPANT_ACTIF,
            STATUT_PARTICIPANT_ATTENTE,
            InscriptionAnnuelle,
            Participant,
            PresenceActivite,
            SessionActivite,
        )
        from app.services.inscriptions_annuelles import rafraichir_statuts

        suf = _suffixe()
        p = Participant(nom=f"Reprise{suf}", prenom="Alex", statut_inscription=STATUT_PARTICIPANT_ATTENTE)
        seance = SessionActivite(
            atelier_id=atelier["id"], secteur="Adultes", session_type="COLLECTIF",
            date_session=date(ANNEE, 11, 7), statut="realisee",
        )
        db.session.add_all([p, seance])
        db.session.flush()

        i = InscriptionAnnuelle(
            annee_scolaire=ANNEE, date_inscription=date(ANNEE, 9, 1),
            nom=p.nom, prenom=p.prenom, statut="en_attente", participant_id=p.id,
        )
        db.session.add(i)
        db.session.flush()

        # Présence insérée en SQL pur : le garde-fou ORM ne la voit pas.
        db.session.execute(
            PresenceActivite.__table__.insert().values(session_id=seance.id, participant_id=p.id)
        )
        db.session.commit()

        inscription_id, participant_id = i.id, p.id
        assert rafraichir_statuts(ANNEE) >= 1

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.statut == "active"
        assert i.premiere_participation_le == date(ANNEE, 11, 7)
        assert db.session.get(Participant, participant_id).statut_inscription == STATUT_PARTICIPANT_ACTIF


# ---------------------------------------------------------------------------
# Règlement
# ---------------------------------------------------------------------------

def test_confirmation_reglement_sans_adhesion(admin_client, app, atelier):
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])

    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/reglement",
        data={"action": "confirmer", "montant": "12,50", "mode": "cheque",
              "date_reglement": f"{ANNEE}-09-15", "reglement_commentaire": "chèque n°42"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.reglement_confirme is True
        assert i.reglement_montant == 12.5           # la virgule décimale est acceptée
        assert i.reglement_mode == "cheque"
        assert i.reglement_date == date(ANNEE, 9, 15)
        assert i.reglement_commentaire == "chèque n°42"
        assert i.cotisation_id is None

    # Le règlement se dé-confirme (erreur de saisie à l'accueil).
    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/reglement", data={"action": "annuler"})
    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.reglement_confirme is False
        assert i.reglement_date is None


def test_confirmation_reglement_cree_ladhesion_et_le_versement(admin_client, app, atelier):
    """La case « créer l'adhésion » alimente le module Adhésions : source
    unique pour les impayés, la caisse et les bilans."""
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])
    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/creer-participant", data={})

    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/reglement",
        data={"action": "confirmer", "montant": "15", "mode": "especes",
              "date_reglement": f"{ANNEE}-09-10", "creer_adhesion": "1"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app.app_context():
        from app.extensions import db
        from app.models import Cotisation, InscriptionAnnuelle

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.cotisation_id is not None

        cotisation = db.session.get(Cotisation, i.cotisation_id)
        assert cotisation.type_cotisation == "adhesion_individuelle"
        assert cotisation.annee_scolaire == ANNEE
        assert cotisation.participant_id == i.participant_id
        assert cotisation.montant_du == 15.0
        assert cotisation.montant_regle == 15.0
        assert cotisation.solde is True
        assert cotisation.paiements[0].mode == "especes"


def test_adhesion_impossible_sans_fiche_participant(admin_client, app, atelier):
    """Sans fiche, le règlement est quand même confirmé sur le bulletin — on
    ne bloque pas l'accueil — mais l'adhésion attend et on le dit."""
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])

    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/reglement",
        data={"action": "confirmer", "montant": "15", "mode": "especes", "creer_adhesion": "1"},
        follow_redirects=True,
    )
    contenu = r.data.decode("utf-8")
    # Message flashé : Jinja échappe l'apostrophe des variables.
    assert "transformer l&#39;inscription en fiche participant" in contenu

    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.reglement_confirme is True
        assert i.cotisation_id is None


# ---------------------------------------------------------------------------
# Cycle de vie
# ---------------------------------------------------------------------------

def test_annulation_puis_reactivation(admin_client, app, atelier):
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])

    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/statut", data={"action": "annuler"})
    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        assert db.session.get(InscriptionAnnuelle, inscription_id).statut == "annulee"

    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/statut", data={"action": "reactiver"})
    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        assert db.session.get(InscriptionAnnuelle, inscription_id).statut == "saisie"


def test_suppression_refusee_quand_une_fiche_en_depend(admin_client, app, atelier):
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])
    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/creer-participant", data={})

    r = admin_client.post(
        f"/inscriptions-annuelles/{inscription_id}/supprimer", follow_redirects=True
    )
    assert "annule-le plutôt" in r.data.decode("utf-8")
    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        assert db.session.get(InscriptionAnnuelle, inscription_id) is not None


def test_suppression_dun_bulletin_saisi_en_double(admin_client, app, atelier):
    inscription_id, _ = _creer_bulletin(admin_client, app, atelier["id"])
    r = admin_client.post(f"/inscriptions-annuelles/{inscription_id}/supprimer", follow_redirects=True)
    assert r.status_code == 200
    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle

        assert db.session.get(InscriptionAnnuelle, inscription_id) is None


# ---------------------------------------------------------------------------
# Pages, fiche imprimable et export
# ---------------------------------------------------------------------------

def test_liste_et_compteurs(admin_client, app, atelier):
    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])
    page = admin_client.get(f"/inscriptions-annuelles/?annee={ANNEE}").data.decode("utf-8")
    assert f"Rentree{suf}" in page
    assert f"{ANNEE}-{ANNEE + 1}" in page
    assert "Disponibilités bénévolat" in page

    # Le filtre « bénévolat » retrouve la personne, le filtre « réglé » non.
    assert f"Rentree{suf}" in admin_client.get(
        f"/inscriptions-annuelles/?annee={ANNEE}&benevolat=1"
    ).data.decode("utf-8")
    assert f"Rentree{suf}" not in admin_client.get(
        f"/inscriptions-annuelles/?annee={ANNEE}&reglement=regle"
    ).data.decode("utf-8")


def test_fiche_imprimable(admin_client, app, atelier):
    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])
    page = admin_client.get(f"/inscriptions-annuelles/{inscription_id}/fiche").data.decode("utf-8")
    assert f"Rentree{suf}" in page
    assert "Fiche d'inscription" in page
    assert "réservée à l'accueil" in page
    assert atelier["nom"] in page
    assert "window.print()" in page


def test_export_xlsx_contient_toutes_les_donnees(admin_client, app, atelier):
    from openpyxl import load_workbook

    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])
    r = admin_client.get(f"/inscriptions-annuelles/export.xlsx?annee={ANNEE}")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["Content-Type"]
    assert f"inscriptions_{ANNEE}-{ANNEE + 1}.xlsx" in r.headers["Content-Disposition"]

    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Inscriptions", "Synthèse", "Bénévolat"]

    detail = wb["Inscriptions"]
    entetes = [c.value for c in detail[4]]
    for colonne in ("Nom", "Prénom", "Adresse", "E-mail", "Téléphone",
                    "Secteur qui fait venir", "Ateliers choisis",
                    "Autres souhaits (champ libre)", "Bénévolat souhaité",
                    "Bénévolat — pour quoi faire", "Bénévolat — disponibilités",
                    "Règlement confirmé"):
        assert colonne in entetes

    lignes = [[c.value for c in ligne] for ligne in detail.iter_rows(min_row=5)]
    ligne = next(l for l in lignes if l[entetes.index("Nom")] == f"Rentree{suf}")
    assert ligne[entetes.index("Ateliers choisis")] == atelier["nom"]
    assert "couture" in ligne[entetes.index("Autres souhaits (champ libre)")]
    assert ligne[entetes.index("Bénévolat souhaité")] == "Oui"
    assert "Mardi matin" in ligne[entetes.index("Bénévolat — disponibilités")]
    assert ligne[entetes.index("Règlement confirmé")] == "Non"

    # La feuille bénévolat sert la réunion d'équipe : grille + liste nominative.
    benevolat = [[c.value for c in ligne] for ligne in wb["Bénévolat"].iter_rows()]
    assert any(f"Soline Rentree{suf}" == (l[0] or "") for l in benevolat)


# ---------------------------------------------------------------------------
# Droits
# ---------------------------------------------------------------------------

def test_pages_protegees(client):
    """Un visiteur anonyme est renvoyé vers la connexion, jamais servi."""
    for url in ("/inscriptions-annuelles/", "/inscriptions-annuelles/nouvelle",
                "/inscriptions-annuelles/export.xlsx"):
        r = client.get(url)
        assert r.status_code in (302, 401), url


def test_permissions_declarees_et_attribuees(app):
    with app.app_context():
        from app.models import Permission, Role

        codes = {p.code for p in Permission.query.all()}
        attendues = {
            "inscriptions_annuelles:view",
            "inscriptions_annuelles:edit",
            "inscriptions_annuelles:reglement",
            "inscriptions_annuelles:export",
        }
        assert attendues <= codes

        direction = {p.code for p in Role.query.filter_by(code="direction").first().permissions}
        assert attendues <= direction

        responsable = {p.code for p in Role.query.filter_by(code="responsable_secteur").first().permissions}
        assert attendues <= responsable

        # Le compte technique consulte mais ne saisit pas.
        tech = {p.code for p in Role.query.filter_by(code="admin_tech").first().permissions}
        assert "inscriptions_annuelles:view" in tech
        assert "inscriptions_annuelles:edit" not in tech


def test_hub_publics_propose_le_module(admin_client):
    page = admin_client.get("/publics").data.decode("utf-8")
    assert "Inscriptions annuelles" in page
    assert "/inscriptions-annuelles/" in page


# ---------------------------------------------------------------------------
# RGPD : le bulletin recopie des données personnelles, il suit la fiche
# ---------------------------------------------------------------------------

def test_anonymisation_efface_aussi_le_bulletin(admin_client, app, atelier):
    """Anonymiser la fiche sans toucher au bulletin ne serait pas anonymiser."""
    inscription_id, suf = _creer_bulletin(admin_client, app, atelier["id"])
    admin_client.post(f"/inscriptions-annuelles/{inscription_id}/creer-participant", data={})

    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle
        from app.services.purge_rgpd import NOM_ANONYME, anonymiser_participant

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        anonymiser_participant(i.participant)
        db.session.commit()

        i = db.session.get(InscriptionAnnuelle, inscription_id)
        assert i.nom == NOM_ANONYME
        assert f"Rentree{suf}" != i.nom
        assert i.adresse is None and i.code_postal is None and i.ville is None
        assert i.email is None and i.telephone is None
        assert i.commentaire is None
        # Ce qui décrit la demande survit : les bilans de campagne restent justes.
        assert [a.id for a in i.ateliers] == [atelier["id"]]
        assert i.benevolat_souhaite is True
        assert i.creneaux_benevolat == [("mardi", "matin"), ("jeudi", "apres_midi")]


def test_une_inscription_recente_protege_de_la_purge(app, atelier):
    """Inscrit à la rentrée, jamais encore venu : la purge ne doit pas
    l'anonymiser sous prétexte qu'il n'a aucune présence."""
    from datetime import date as _date

    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle, Participant
        from app.services.purge_rgpd import derniere_activite_par_participant

        suf = _suffixe()
        p = Participant(nom=f"Nouveau{suf}", prenom="Sam")
        db.session.add(p)
        db.session.flush()
        db.session.add(InscriptionAnnuelle(
            annee_scolaire=ANNEE, date_inscription=_date.today(),
            nom=p.nom, prenom=p.prenom, participant_id=p.id, statut="en_attente",
        ))
        db.session.commit()

        activites = derniere_activite_par_participant()
        assert p.id in activites, "une inscription annuelle doit valoir activité"


def test_suppression_definitive_emporte_le_bulletin(app, atelier):
    """La suppression totale d'une fiche ne doit rien laisser derrière elle —
    créneaux de bénévolat et ateliers souhaités compris."""
    with app.app_context():
        from app.extensions import db
        from app.models import (
            InscriptionAnnuelle,
            InscriptionAnnuelleDispo,
            Participant,
            inscription_annuelle_atelier,
        )
        from app.services.inscriptions_annuelles import appliquer_disponibilites
        from app.services.participant_suppression import analyser, supprimer_definitivement

        suf = _suffixe()
        p = Participant(nom=f"Efface{suf}", prenom="Dominique")
        db.session.add(p)
        db.session.flush()

        i = InscriptionAnnuelle(
            annee_scolaire=ANNEE, date_inscription=date(ANNEE, 9, 1),
            nom=p.nom, prenom=p.prenom, participant_id=p.id,
            benevolat_souhaite=True,
        )
        i.ateliers = [db.session.get(type(i).ateliers.property.mapper.class_, atelier["id"])]
        db.session.add(i)
        appliquer_disponibilites(i, ["lundi|matin", "vendredi|apres_midi"])
        db.session.commit()
        inscription_id, participant_id = i.id, p.id

        # Le bulletin est annoncé avant la suppression.
        resume = analyser(db.session.get(Participant, participant_id))
        assert any("inscription annuelle" in libelle for libelle, _ in resume["details"])

        supprimer_definitivement(db.session.get(Participant, participant_id))
        db.session.commit()

        assert db.session.get(Participant, participant_id) is None
        assert db.session.get(InscriptionAnnuelle, inscription_id) is None
        assert InscriptionAnnuelleDispo.query.filter_by(inscription_id=inscription_id).count() == 0
        restes = db.session.execute(
            inscription_annuelle_atelier.select().where(
                inscription_annuelle_atelier.c.inscription_id == inscription_id
            )
        ).all()
        assert restes == []


def test_export_rgpd_contient_les_inscriptions(app, atelier):
    """Droit d'accès (article 15) : la personne doit recevoir son bulletin."""
    with app.app_context():
        from app.extensions import db
        from app.models import InscriptionAnnuelle, Participant
        from app.services.rgpd_export import construire_export_rgpd

        suf = _suffixe()
        p = Participant(nom=f"Acces{suf}", prenom="Noa")
        db.session.add(p)
        db.session.flush()
        db.session.add(InscriptionAnnuelle(
            annee_scolaire=ANNEE, date_inscription=date(ANNEE, 9, 4),
            nom=p.nom, prenom=p.prenom, participant_id=p.id,
            telephone="0344123456", secteur_orienteur="Adultes",
            ateliers_libre="Cours de français", benevolat_souhaite=True,
            benevolat_mission="Sorties familles",
        ))
        db.session.commit()

        wb = construire_export_rgpd(db.session.get(Participant, p.id))
        assert "Inscriptions annuelles" in wb.sheetnames
        feuille = wb["Inscriptions annuelles"]
        lignes = [[c.value for c in ligne] for ligne in feuille.iter_rows(min_row=2)]
        assert any("Cours de français" in (c or "") for ligne in lignes for c in ligne)
        assert any("Sorties familles" in (c or "") for ligne in lignes for c in ligne)
