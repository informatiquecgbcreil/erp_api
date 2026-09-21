"""Régressions : droits croisés, liens directs, réactivation et confidentialité."""
import json
import pytest


@pytest.fixture
def module_scope(app):
    from app.extensions import db
    from app.models import InstanceSettings
    with app.app_context():
        row = InstanceSettings.query.first()
        created = row is None
        row = row or InstanceSettings()
        db.session.add(row)
        previous = row.enabled_modules_json
        row.enabled_modules_json = json.dumps(["presences", "statistiques"])
        db.session.commit()
        row_id = row.id
    yield
    with app.app_context():
        row = db.session.get(InstanceSettings, row_id)
        if created:
            db.session.delete(row)
        else:
            row.enabled_modules_json = previous
        db.session.commit()


@pytest.mark.parametrize("path", ["/rh", "/caisse", "/dons", "/repartition-participation", "/inventaire/", "/salles/", "/partenaires/", "/insertion/", "/questionnaires/", "/stats", "/stats-bilans", "/bilans", "/bilans/export.xlsx", "/bilans/financeurs", "/bilans/inventaire"])
def test_disabled_modules_deny_even_direction(admin_client, module_scope, path):
    response = admin_client.get(path, follow_redirects=True)
    assert response.status_code == 404


def test_active_modules_and_admin_remain_available(admin_client, module_scope):
    for path in ["/dashboard", "/activite/", "/admin/modules", "/admin/instance"]:
        assert admin_client.get(path, follow_redirects=True).status_code == 200


def test_simple_home_and_senacs_hide_disabled_sections(admin_client, module_scope):
    with admin_client.session_transaction() as state:
        state["ui_mode"] = "simple"
    page = admin_client.get("/dashboard").get_data(as_text=True)
    assert "Vos outils" in page
    assert "Statistiques de fréquentation" in page
    assert 'href="/stats-bilans"' not in page
    assert 'href="/rh"' not in page
    senacs = admin_client.get("/bilans/senacs")
    assert senacs.status_code == 200
    assert "Emplois / ETP" not in senacs.get_data(as_text=True)
    assert "TOTAL subventions" not in senacs.get_data(as_text=True)
    assert admin_client.post("/bilans/senacs/emplois", data={"annee": "2026", "intitule": "Test"}).status_code == 404
    documents = admin_client.get("/documents")
    assert documents.status_code == 200
    assert 'href="/bilans/export.xlsx' not in documents.get_data(as_text=True)


def test_senacs_export_excludes_disabled_data(admin_client, module_scope):
    from io import BytesIO
    from openpyxl import load_workbook
    result = admin_client.get("/bilans/senacs/export.xlsx")
    assert result.status_code == 200
    workbook = load_workbook(BytesIO(result.data))
    assert "Public global" in workbook.sheetnames
    assert not {"Finances", "Emplois", "Partenariats"} & set(workbook.sheetnames)
    workbook.close()


def test_role_permissions_are_intersected_not_deleted(app, module_scope):
    from app.models import User
    with app.test_request_context():
        user = User.query.filter_by(email="admin@example.org").one()
        assert any(p.code == "rh:view" for r in user.roles for p in r.permissions)
        assert not user.has_perm("rh:view")
        assert user.has_perm("emargement:view")


def test_modules_reenable_without_reinstall(admin_client, module_scope):
    from app.services.modules import CATALOG
    response = admin_client.post("/admin/modules", data={"modules": list(CATALOG)})
    assert response.status_code == 302
    assert admin_client.get("/rh").status_code == 200
    assert admin_client.post("/admin/modules", data={"modules": ["invented"]}).status_code == 400


def test_finance_without_presence_or_statistics(admin_client, module_scope):
    assert admin_client.post("/admin/modules", data={"modules": ["finances"]}).status_code == 302
    for path in ["/dashboard", "/stats", "/stats-bilans", "/bilans", "/bilans/financeurs"]:
        assert admin_client.get(path, follow_redirects=True).status_code == 200, path
    assert admin_client.get("/activite/").status_code == 404
    assert admin_client.get("/bilans/senacs").status_code == 404


def test_embedded_postgres_tools_override_machine_path(app, tmp_path, monkeypatch):
    from app.services.sauvegarde import _trouver_pg_dump, _trouver_psql
    for key, find in [("PG_DUMP_PATH", _trouver_pg_dump), ("PSQL_PATH", _trouver_psql)]:
        bundled = tmp_path / (key + ".exe")
        bundled.touch()
        monkeypatch.setitem(app.config, key, str(bundled))
        monkeypatch.setattr("app.services.sauvegarde.shutil.which", lambda name: "C:/old-postgres/" + name)
        with app.app_context(): assert find() == str(bundled)


def test_dependencies_empty_and_corrupt_config(app, module_scope):
    from app.services.modules import normalize, enabled_modules
    from app.models import InstanceSettings
    from app.extensions import db
    assert normalize(["statistiques"]) == ["presences", "statistiques"]
    assert normalize([]) == []
    with app.app_context():
        InstanceSettings.query.first().enabled_modules_json = "invalid JSON"
        db.session.commit()
    with app.test_request_context():
        assert enabled_modules() == set()


def test_module_change_requires_csrf(admin_client, app, module_scope, monkeypatch):
    monkeypatch.setitem(app.config, "WTF_CSRF_ENABLED", True)
    assert admin_client.post("/admin/modules", data={"modules": ["finances"]}).status_code == 400


def test_ordinary_finance_cannot_change_structure_modules(app, module_scope):
    from app.models import User, Role
    from app.extensions import db
    with app.app_context():
        u = User(email="modules-finance@example.org", nom="Finance")
        u.set_password("Test-module-finances-12")
        u.roles.append(Role.query.filter_by(code="finance").one())
        db.session.add(u); db.session.commit(); uid = u.id
    client = app.test_client()
    with client.session_transaction() as s: s["_user_id"] = str(uid); s["_fresh"] = True
    try:
        assert client.post("/admin/modules", data={"modules": ["finances"]}).status_code == 403
        assert client.get("/admin/users").status_code == 403
    finally:
        with app.app_context(): db.session.delete(db.session.get(User, uid)); db.session.commit()


@pytest.mark.parametrize("path", ["justifs/secret.pdf", "factures/facture.pdf", "projets/compte-rendu.pdf", "bilans_lourds/2026/ALL/photo.jpg"])
def test_media_anonymous_denied(client, path):
    assert client.get("/media/" + path).status_code == 401


def test_private_media_cannot_bypass_business_route(admin_client):
    assert admin_client.get("/media/justifs/secret.pdf").status_code == 404
    assert admin_client.get("/static/uploads/secret.pdf").status_code == 404


def test_media_logo_remains_available_with_restrictive_headers(app, client, tmp_path, monkeypatch):
    from app.models import InstanceSettings
    from app.extensions import db
    (tmp_path / "branding").mkdir()
    (tmp_path / "branding/logo.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"/>')
    monkeypatch.setitem(app.config, "APP_UPLOAD_DIR", str(tmp_path))
    with app.app_context():
        row = InstanceSettings.query.first() or InstanceSettings()
        old = row.app_logo_path
        row.app_logo_path = "branding/logo.svg"; db.session.add(row); db.session.commit()
    try:
        result = client.get("/media/branding/logo.svg")
        assert result.status_code == 200
        assert "sandbox" in result.headers["Content-Security-Policy"]
        assert result.headers["X-Content-Type-Options"] == "nosniff"
        result.close()
    finally:
        with app.app_context(): InstanceSettings.query.first().app_logo_path = old; db.session.commit()


@pytest.mark.parametrize("target", ["https://external.invalid", "//external.invalid", "/\\external.invalid"])
def test_dashboard_next_cannot_redirect_offsite(admin_client, target):
    response = admin_client.post("/ui-mode", data={"mode": "simple", "next": target})
    assert response.headers["Location"] == "/dashboard"


def test_windows_web_setup_disabled(app, client, monkeypatch):
    monkeypatch.setitem(app.config, "SETUP_DISABLED", True)
    assert client.get("/setup/").status_code == 404


def test_postgres_password_not_in_process_arguments():
    from app.services.sauvegarde import _private_pg_connection
    uri, env = _private_pg_connection("postgresql+psycopg://mcs:s3cr%40t@127.0.0.1:5432/test")
    assert "s3cr" not in uri
    assert env["PGPASSWORD"] == "s3cr@t"


def test_smtp_starttls_validates_certificate_without_invalid_timeout(app, monkeypatch):
    import ssl
    import smtplib
    from app.services.instance_settings import envoyer_email_test
    calls = []
    class SMTP:
        def __init__(self, host, port, timeout): assert timeout == 10
        def ehlo(self): pass
        def starttls(self, *, context):
            assert context.verify_mode == ssl.CERT_REQUIRED
            assert context.check_hostname
            calls.append("tls")
        def login(self, user, password): calls.append("login")
        def send_message(self, message): calls.append("send")
        def quit(self): pass
    monkeypatch.setattr(smtplib, "SMTP", SMTP)
    monkeypatch.setattr("app.services.instance_settings.resolve_mail_settings", lambda c: dict(host="smtp.example.test", port=587, use_tls=True, sender="test@example.test", username="test", password="secret"))
    with app.app_context(): envoyer_email_test({}, "recipient@example.test")
    assert calls == ["tls", "login", "send"]
