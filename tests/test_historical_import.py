"""Transactional migration tests on disposable SQLite databases only."""
from copy import deepcopy
from datetime import date, datetime
import hashlib
import json
import os

import pytest
from flask import Flask
from openpyxl import Workbook
from sqlalchemy import event, select

from app.extensions import db
from app.models import (Participant, Quartier, Secteur, AtelierActivite, SessionActivite,
                        PresenceActivite, HistoricalImportBatch, HistoricalImportSource)
from app.ateliers.historical_import import analyze_import, apply_import, verify_batch, PreviewStale, ImportBlocked


@pytest.fixture
def migration_db(tmp_path):
    app = Flask("migration-tests")
    app.config.update(SQLALCHEMY_DATABASE_URI="sqlite:///" + (tmp_path / "migration.db").as_posix(),
                      SQLALCHEMY_TRACK_MODIFICATIONS=False, TESTING=True)
    db.init_app(app)
    with app.app_context():
        db.create_all()
        db.session.add_all([Secteur(code="num", label="Numérique"), Secteur(code="fam", label="Familles")])
        db.session.commit()
        yield app
        db.session.remove()
        db.engine.dispose()


def workbook(tmp_path, sheets=None):
    """Small boundary fixtures complement, never replace, the real XLSX test."""
    wb = Workbook()
    wb.remove(wb.active)
    sheets = sheets or [{"name": "Atelier", "rows": [("Dupont", "Jean", 1985, "0612345678", "ROUHER", [1, 1, 1])]}]
    for specification in sheets:
        ws = wb.create_sheet(specification["name"])
        ws.cell(3, 8, datetime(2026, 1, 1))
        days = specification.get("days", [12, 19, 26])
        slots = specification.get("slots", ["M"] * len(days))
        for col, (day, slot) in enumerate(zip(days, slots), 8):
            ws.cell(4, col, ["L", "M", "ME", "J", "V", "S", "D"][date(2026, 1, day).weekday()])
            ws.cell(5, col, slot)
            ws.cell(6, col, day)
        for col, label in enumerate(["NOM", "PRENOM", "ANNEE", "TELEPHONE", "QUARTIER", "EMAIL", "ACTIVITE"], 1):
            ws.cell(6, col, label)
        total_col = 8 + len(days)
        ws.cell(6, total_col, "TOTAL")
        total = 0
        for row_index, (nom, prenom, birth, phone, district, presence) in enumerate(specification["rows"], 7):
            values = [nom, prenom, birth, phone, district, specification.get("email"), specification.get("activity", specification["name"])]
            for col, value in enumerate(values, 1):
                ws.cell(row_index, col, value)
            for col, value in enumerate(presence, 8):
                ws.cell(row_index, col, value)
            subtotal = sum(v == 1 for v in presence)
            ws.cell(row_index, total_col, subtotal)
            total += subtotal
        ws.cell(7 + len(specification["rows"]), 1, "TOTAL")
        ws.cell(7 + len(specification["rows"]), total_col, total)
    path = tmp_path / "register_2026.xlsx"
    wb.save(path)
    return path


def decisions_for(plan, sectors=None):
    sectors = sectors or {}
    return {"activities": {a["key"]: {"secteur": sectors.get(a["key"], "Numérique")} for a in plan["activities"]}}


def plan_ready(path, decisions=None):
    initial = analyze_import(path)
    decisions = decisions or decisions_for(initial)
    plan = analyze_import(path, decisions)
    assert plan["ready"], plan["blockers"]
    return plan


def apply_plan(path, plan):
    return apply_import(path, decisions=plan["decisions"], expected_digest=plan["digest"])


def counts():
    return {model.__tablename__: db.session.query(model).count() for model in (
        Participant, Quartier, AtelierActivite, SessionActivite, PresenceActivite, HistoricalImportBatch, HistoricalImportSource)}


def test_analyze_first_then_assign_sector_per_activity(migration_db, tmp_path):
    p = ("Dupont", "Jean", 1985, "0612345678", "ROUHER", [1, 1, 1])
    path = workbook(tmp_path, [{"name": "Informatique", "rows": [p]}, {"name": "Cuisine", "rows": [p]}])
    initial = analyze_import(path)
    assert initial["summary"]["activities_unassigned"] == 2
    assert not initial["ready"]
    assert counts()["participant"] == 0
    decisions = decisions_for(initial, {"Cuisine": "Familles"})
    plan = plan_ready(path, decisions)
    result = apply_plan(path, plan)
    assert result["created"]["participants"] == 1
    assert result["created"]["presences"] == 6
    assert {a.nom: a.secteur for a in AtelierActivite.query} == {"Cuisine": "Familles", "Informatique": "Numérique"}
    assert {s.secteur for s in SessionActivite.query} == {"Familles", "Numérique"}


def test_complementary_duplicates_union_and_single_presence(migration_db, tmp_path):
    rows = [("Dupont", "Jean", 1985, "0612345678", "ROUHER", values) for values in ([1, None, 1], [None, 1, None], [1, None, None])]
    path = workbook(tmp_path, [{"name": "Atelier", "rows": rows}])
    plan = plan_ready(path)
    assert plan["summary"]["presences_detected"] == 4
    assert plan["summary"]["internal_duplicate_presences"] == 1
    assert plan["summary"]["internal_duplicate_people"] == 2
    result = apply_plan(path, plan)
    assert result["created"]["participants"] == 1
    assert result["created"]["presences"] == 3
    links = HistoricalImportSource.query.filter_by(kind="presence").all()
    assert len(links) == 4
    assert len({link.presence_id for link in links}) == 3
    assert {link.source_cell for link in links} == {"H7", "J7", "I8", "H9"}
    assert verify_batch(result["batch_id"])["ok"]


@pytest.mark.parametrize("slots", [["M", "AM"], ["M", "M"]])
def test_distinct_same_day_columns_never_merge(migration_db, tmp_path, slots):
    path = workbook(tmp_path, [{"name": "Atelier", "days": [12, 12], "slots": slots,
        "rows": [("Dupont", "Jean", 1985, "0612345678", "ROUHER", [1, 1])]}])
    result = apply_plan(path, plan_ready(path))
    assert result["created"]["sessions"] == 2
    assert result["created"]["presences"] == 2
    assert all(s.heure_debut is None and s.heure_fin is None for s in SessionActivite.query)
    assert sorted(s.creneau_source for s in SessionActivite.query) == sorted(slots)


def test_name_only_homonyms_review_manual_same_and_different(migration_db, tmp_path):
    rows = [("Martin", "Marie", None, None, None, values) for values in ([1, None, None], [None, 1, None])]
    path = workbook(tmp_path, [{"name": "Atelier", "rows": rows}])
    initial = analyze_import(path)
    assert initial["matching"]["counts"]["REVIEW"] == 2
    d = decisions_for(initial)
    d["participants"] = {p["key"]: {"action": "new", "group": "personne-verifiee"} for p in initial["parser"]["people"]}
    plan = plan_ready(path, d)
    assert plan["summary"]["participants_new"] == 1
    d["participants"] = {p["key"]: {"action": "new"} for p in initial["parser"]["people"]}
    result = apply_plan(path, plan_ready(path, d))
    assert result["created"]["participants"] == 2
    assert result["created"]["presences"] == 2


def test_same_file_reimport_is_idempotent_including_saved_decisions(migration_db, tmp_path):
    path = workbook(tmp_path)
    first = apply_plan(path, plan_ready(path))
    before = counts()
    plan = analyze_import(path)
    assert plan["ready"]
    assert plan["summary"]["presences_existing"] == 3
    assert plan["summary"]["participants_exact"] == 1
    second = apply_plan(path, plan)
    assert second["batch_id"] == first["batch_id"]
    assert second["already_imported"]
    assert second["created"] == {}
    assert counts() == before


def test_dry_run_has_no_dml_and_does_not_commit_caller_work(migration_db, tmp_path):
    path = workbook(tmp_path)
    before = counts()
    writes = []
    def watch(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")):
            writes.append(statement)
    event.listen(db.engine, "before_cursor_execute", watch)
    try:
        analyze_import(path)
        assert not writes
        assert counts() == before
    finally:
        event.remove(db.engine, "before_cursor_execute", watch)
    pending = Participant(nom="Travail", prenom="Non enregistré")
    db.session.add(pending)
    analyze_import(path)
    assert pending in db.session.new
    db.session.rollback()
    assert counts() == before


def test_rollback_on_failure_after_real_inserts(migration_db, tmp_path, monkeypatch):
    from app.ateliers import historical_import as service
    path = workbook(tmp_path)
    plan = plan_ready(path)
    before = counts()
    def fail(*args, **kwargs):
        raise RuntimeError("Erreur injectée après création du lot et du participant")
    monkeypatch.setattr(service, "_provenance", fail)
    with pytest.raises(RuntimeError):
        apply_plan(path, plan)
    assert counts() == before


def test_unresolved_import_refused_without_changes(migration_db, tmp_path):
    path = workbook(tmp_path)
    plan = analyze_import(path)
    before = counts()
    with pytest.raises(ImportBlocked):
        apply_plan(path, plan)
    assert counts() == before
    with pytest.raises(ValueError, match="dry-run"):
        apply_import(path)


def test_stale_database_preview_and_changed_assignment_rejected(migration_db, tmp_path):
    path = workbook(tmp_path)
    plan = plan_ready(path)
    d = deepcopy(plan["decisions"])
    d["activities"]["Atelier"]["secteur"] = "Familles"
    with pytest.raises(PreviewStale):
        apply_import(path, d, expected_digest=plan["digest"])
    db.session.add(Participant(nom="Une", prenom="Autre"))
    db.session.commit()
    before = counts()
    with pytest.raises(PreviewStale):
        apply_plan(path, plan)
    assert counts() == before


def test_existing_activity_typography_reused_in_its_sector(migration_db, tmp_path):
    activity = AtelierActivite(nom=" Accès aux droits  ", secteur="Familles")
    db.session.add(activity)
    db.session.commit()
    path = workbook(tmp_path, [{"name": "Onglet", "activity": "ACCES AUX DROITS", "rows": [("Dupont", "Jean", 1985, None, None, [1, 1, 1])]}])
    d = {"activities": {"Onglet": {"secteur": "Familles"}}}
    plan = plan_ready(path, d)
    assert plan["activities"][0]["atelier_id"] == activity.id
    result = apply_plan(path, plan)
    assert result["created"].get("activities", 0) == 0
    assert all(s.secteur == "Familles" for s in SessionActivite.query)


@pytest.mark.parametrize("reverse", [False, True])
def test_sector_assignment_does_not_confirm_similar_source_activities(migration_db, tmp_path, reverse):
    sheets = [{"name": "A", "activity": "Atelier numerique collectif", "rows": [("Dupont", "Jean", 1985, "0612345678", None, [1, 1, 1])]},
              {"name": "B", "activity": "Atelier numerique collectifs", "rows": [("Durand", "Alice", 1990, "0698765432", None, [1, 1, 1])]}]
    path = workbook(tmp_path, list(reversed(sheets)) if reverse else sheets)
    d = decisions_for(analyze_import(path))
    plan = analyze_import(path, d)
    assert not plan["ready"]
    assert {a["key"] for a in plan["activities"] if a["status"] == "REVIEW"} == {"A", "B"}
    assert {b["key"] for b in plan["blockers"] if b["kind"] == "activity"} == {"A", "B"}
    before = counts()
    with pytest.raises(ImportBlocked):
        apply_plan(path, plan)
    assert counts() == before
    # A human can explicitly confirm a shared business name after reviewing both sheets.
    d["activities"]["B"]["name"] = sheets[0]["activity"]
    result = apply_plan(path, plan_ready(path, d))
    assert result["created"]["activities"] == 1
    assert result["created"]["sessions"] == 6
    assert result["created"]["presences"] == 6


@pytest.mark.parametrize("resolution,expected_activities", [("different_sector", 2), ("confirmed_distinct", 2), ("ignore", 1)])
def test_similar_activity_choices_remain_available(migration_db, tmp_path, resolution, expected_activities):
    path = workbook(tmp_path, [
        {"name": "A", "activity": "Atelier numerique collectif", "rows": [("Dupont", "Jean", 1985, "0612345678", None, [1, 1, 1])]},
        {"name": "B", "activity": "Atelier numerique collectifs", "rows": [("Durand", "Alice", 1990, "0698765432", None, [1, 1, 1])]}])
    d = decisions_for(analyze_import(path))
    if resolution == "different_sector":
        d["activities"]["B"]["secteur"] = "Familles"
    elif resolution == "ignore":
        d["activities"]["B"]["action"] = "ignore"
    else:
        for choice in d["activities"].values():
            choice["action"] = "new"
    result = apply_plan(path, plan_ready(path, d))
    assert result["created"]["activities"] == expected_activities


@pytest.mark.parametrize("target_kind", ["missing_date", "monthly", "boolean_id"])
def test_incompatible_existing_session_is_rejected(migration_db, tmp_path, target_kind):
    a = AtelierActivite(nom="Atelier", secteur="Numérique")
    db.session.add(a)
    db.session.flush()
    s = SessionActivite(atelier_id=a.id, secteur="Numérique",
                        session_type="INDIVIDUEL_MENSUEL" if target_kind == "monthly" else "COLLECTIF",
                        date_session=None if target_kind == "missing_date" else date(2026, 1, 12))
    db.session.add(s)
    db.session.commit()
    path = workbook(tmp_path)
    plan = analyze_import(path)
    d = decisions_for(plan)
    d["sessions"] = {plan["sessions"][0]["key"]: {"action": "existing", "session_id": True if target_kind == "boolean_id" else s.id}}
    before = counts()
    with pytest.raises(ValueError, match="Séance cible"):
        analyze_import(path, d)
    assert counts() == before


def test_activity_target_requires_an_integer_identifier(migration_db, tmp_path):
    db.session.add(AtelierActivite(nom="Atelier", secteur="Numérique"))
    db.session.commit()
    path = workbook(tmp_path)
    with pytest.raises(ValueError, match="Activité cible"):
        analyze_import(path, {"activities": {"Atelier": {"action": "existing", "atelier_id": True}}})


def test_existing_session_and_presence_reused_only_after_review(migration_db, tmp_path):
    p = Participant(nom="Dupont", prenom="Jean", date_naissance=date(1985, 4, 3), telephone="0612345678", ville="Creil")
    a = AtelierActivite(nom="Atelier", secteur="Numérique")
    db.session.add_all([p, a])
    db.session.flush()
    s = SessionActivite(atelier_id=a.id, secteur="Numérique", date_session=date(2026, 1, 12))
    db.session.add(s)
    db.session.flush()
    db.session.add(PresenceActivite(session_id=s.id, participant_id=p.id))
    db.session.commit()
    path = workbook(tmp_path, [{"name": "Atelier", "days": [12], "rows": [("Dupont", "Jean", date(1985, 4, 3), "0612345678", None, [1])]}])
    initial = analyze_import(path)
    d = decisions_for(initial)
    review = analyze_import(path, d)
    assert review["summary"]["sessions_review"] == 1
    assert review["matching"]["rows"][0]["classification"] == "EXACT"
    d["sessions"] = {initial["sessions"][0]["key"]: {"action": "existing", "session_id": s.id}}
    plan = plan_ready(path, d)
    assert plan["summary"]["presences_existing"] == 1
    result = apply_plan(path, plan)
    assert result["created"] == {}
    assert counts()["presence_activite"] == 1


@pytest.mark.parametrize("territory,city,district", [("NOGENT SUR OISE", "Nogent-sur-Oise", None), ("ROUHER", "Creil", "Rouher"), (None, None, None)])
def test_territory_and_year_only_storage(migration_db, tmp_path, territory, city, district):
    path = workbook(tmp_path, [{"name": "Atelier", "rows": [("Dupont", "Jean", 1985, None, territory, [1, 1, 1])]}])
    apply_plan(path, plan_ready(path))
    p = Participant.query.one()
    assert p.ville == city
    assert (p.quartier.nom if p.quartier else None) == district
    assert p.date_naissance is None
    assert p.annee_naissance == 1985
    assert p.age_au(date(2026, 12, 31)) == 41
    assert p.age_au(date(2026, 6, 30)) is None


def test_cannot_reassign_existing_activity_implicitly(migration_db, tmp_path):
    a = AtelierActivite(nom="Atelier", secteur="Familles")
    db.session.add(a)
    db.session.commit()
    path = workbook(tmp_path)
    with pytest.raises(ValueError, match="contredit"):
        analyze_import(path, {"activities": {"Atelier": {"action": "existing", "atelier_id": a.id, "secteur": "Numérique"}}})
    assert db.session.get(AtelierActivite, a.id).secteur == "Familles"


def test_activity_ignore_creates_no_people_or_sessions(migration_db, tmp_path):
    path = workbook(tmp_path)
    plan = plan_ready(path, {"activities": {"Atelier": {"action": "ignore"}}})
    result = apply_plan(path, plan)
    assert result["created"] == {}
    assert counts()["participant"] == 0
    assert counts()["session_activite"] == 0
    assert HistoricalImportSource.query.count() > 0


def test_missing_identity_can_be_corrected_without_losing_attendance(migration_db, tmp_path):
    path = workbook(tmp_path, [{"name": "Atelier", "rows": [("Dupont", None, 1985, None, None, [1, 1, 1])]}])
    initial = analyze_import(path)
    d = decisions_for(initial)
    d["participants"] = {initial["parser"]["people"][0]["key"]: {"action": "new", "values": {"prenom": "Jean vérifié"}}}
    d["acknowledged_anomalies"] = [a["id"] for a in initial["anomalies"]]
    result = apply_plan(path, plan_ready(path, d))
    assert result["created"]["presences"] == 3
    assert Participant.query.one().prenom == "Jean vérifié"


def test_invalid_sector_cannot_be_applied(migration_db, tmp_path):
    path = workbook(tmp_path)
    with pytest.raises(ValueError, match="Secteur"):
        analyze_import(path, {"activities": {"Atelier": {"secteur": "FAUX"}}})


def test_anonymization_removes_historical_identity_copies(migration_db, tmp_path):
    from app.services.participant_privacy import anonymize_participant_fields
    path = workbook(tmp_path)
    result = apply_plan(path, plan_ready(path))
    person = Participant.query.one()
    anonymize_participant_fields(person, strict=True)
    db.session.commit()
    assert person.annee_naissance is None
    assert all(json.loads(s.raw_json).get("redacted") for s in HistoricalImportSource.query.filter_by(participant_id=person.id))
    assert verify_batch(result["batch_id"])["ok"]


def test_deleted_targets_are_reported_and_not_recreated(migration_db, tmp_path):
    from app.services.participant_suppression import supprimer_definitivement
    path = workbook(tmp_path)
    result = apply_plan(path, plan_ready(path))
    supprimer_definitivement(Participant.query.one())
    db.session.commit()
    assert not verify_batch(result["batch_id"])["ok"]
    plan = analyze_import(path)
    assert not plan["ready"]
    assert any(b["kind"] == "provenance" for b in plan["blockers"])


def test_real_workbook_full_dry_run_database_unchanged(migration_db):
    path = os.environ.get("HISTORICAL_XLSX_PATH")
    if not path:
        pytest.skip("Set HISTORICAL_XLSX_PATH to the private real workbook")
    before = counts()
    file_before = hashlib.sha256(open(path, "rb").read()).hexdigest()
    plan = analyze_import(path)
    assert plan["summary"]["sheets"] == 69
    assert plan["summary"]["activities_recognized"] == 66
    assert plan["summary"]["sessions_detected"] == 880
    assert plan["summary"]["presences_detected"] == 7031
    assert plan["summary"]["participants_source"] == 1464
    assert not plan["ready"]
    assert counts() == before
    assert hashlib.sha256(open(path, "rb").read()).hexdigest() == file_before
